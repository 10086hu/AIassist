from __future__ import annotations

import csv
import io
import re
import time
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

import pdfplumber
from openpyxl import load_workbook


MAX_DOCX_XML_BYTES = 64 * 1024 * 1024
MAX_TABLES = 200
MAX_TABLE_ROWS = 20_000
MAX_TABLE_CELLS = 120_000
MAX_TEXT_CHARS = 3_000_000
MAX_PARSE_SECONDS = 25.0

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"

HEADER_ALIASES = {
    "name": ("产品名称", "商品名称", "产品小类", "费用名称", "功能点", "功能名称", "项目名称", "设备名称", "服务名称", "数据服务事项", "数据治理服务事项", "接口用途", "名称", "item_name", "name"),
    "brand": ("品牌", "厂商", "brand"),
    "model": ("型号", "产品型号", "model"),
    "spec": ("规格", "规格参数", "配置", "技术参数", "spec"),
    "category": ("类别", "分类", "费用类别", "项目类别", "category"),
    "unit": ("单位", "计量单位", "unit"),
    "quantity": ("数量", "采购数量", "工程量", "qty", "quantity"),
    "unit_price": ("单价", "含税单价", "申报单价", "人月单价", "unit_price"),
    "total_price": ("合价", "总价", "总金额", "申报金额", "金额", "合计", "投资额", "概算", "预算", "费用", "total_price"),
    "effort": ("人月", "工作量", "开发人月", "effort", "person_month"),
}

PRICE_HEADER_KEYWORDS = ("费用", "投资", "预算", "概算", "价格", "报价", "产品", "设备", "软硬件", "功能点", "人月")
PROJECT_TYPE_OPERATION_KEYWORDS = ("运维项目", "运行维护项目", "运营维护项目", "续建运维")
PROJECT_TYPE_CONSTRUCTION_KEYWORDS = ("新建项目", "建设类项目", "建设项目", "改建项目", "升级改造项目")


@dataclass(frozen=True)
class ParsedPriceItem:
    item_name: str
    brand: str = ""
    model: str = ""
    spec: str = ""
    category: str = ""
    unit: str = "项"
    quantity: float | None = None
    unit_price: float | None = None
    total_price: float | None = None
    effort_person_month: float | None = None
    source_section: str = ""
    source_row: int | None = None
    evidence: str = ""


@dataclass(frozen=True)
class PriceProjectContext:
    project_type: str = "unknown"
    total_investment: float | None = None
    direct_construction_cost: float | None = None
    software_development_cost: float | None = None
    hardware_software_purchase_cost: float | None = None
    application_development_ratio: float | None = None
    security_level: int | None = None

    def public_dict(self) -> dict[str, Any]:
        return {
            "project_type": self.project_type,
            "total_investment": self.total_investment,
            "direct_construction_cost": self.direct_construction_cost,
            "software_development_cost": self.software_development_cost,
            "hardware_software_purchase_cost": self.hardware_software_purchase_cost,
            "application_development_ratio": self.application_development_ratio,
            "security_level": self.security_level,
        }


@dataclass(frozen=True)
class ParsedPriceDocument:
    items: list[ParsedPriceItem]
    context: PriceProjectContext
    warnings: list[str] = field(default_factory=list)
    table_count: int = 0
    relevant_table_count: int = 0
    elapsed_seconds: float = 0.0


def parse_price_document(content: bytes, filename: str) -> ParsedPriceDocument:
    started = time.monotonic()
    suffix = Path(filename).suffix.lower()
    warnings: list[str] = []
    full_text = ""
    table_count = 0
    relevant_table_count = 0
    items: list[ParsedPriceItem] = []

    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        try:
            for sheet in workbook.worksheets:
                _check_deadline(started)
                table_count += 1
                rows = [[_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
                parsed = _parse_rows(rows, sheet.title)
                if parsed:
                    relevant_table_count += 1
                    items.extend(parsed)
                full_text += "\n" + "\n".join(" | ".join(row) for row in rows[:500])
        finally:
            workbook.close()
    elif suffix == ".csv":
        rows = [[_text(cell) for cell in row] for row in csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace")))]
        table_count = 1
        items = _parse_rows(rows, "CSV")
        relevant_table_count = 1 if items else 0
        full_text = "\n".join(" | ".join(row) for row in rows)
    elif suffix == ".docx":
        items, full_text, table_count, relevant_table_count, doc_warnings = _parse_docx(content, started)
        warnings.extend(doc_warnings)
    elif suffix == ".pdf":
        items, full_text, table_count, relevant_table_count, pdf_warnings = _parse_pdf(content, started)
        warnings.extend(pdf_warnings)
    else:
        raise ValueError("价格检查支持 .xlsx、.xlsm、.csv、.docx 或 .pdf 文件")

    deduped = _dedupe(items)
    context = _build_project_context(deduped, full_text)
    return ParsedPriceDocument(
        items=deduped,
        context=context,
        warnings=warnings,
        table_count=table_count,
        relevant_table_count=relevant_table_count,
        elapsed_seconds=round(time.monotonic() - started, 3),
    )


def parse_price_items(content: bytes, filename: str) -> list[ParsedPriceItem]:
    return parse_price_document(content, filename).items


def parse_benchmark_records(content: bytes, filename: str) -> list[dict[str, Any]]:
    parsed = parse_price_document(content, filename)
    records = []
    for item in parsed.items:
        if item.unit_price is None:
            continue
        records.append(
            {
                "item_name": item.item_name,
                "brand": item.brand,
                "model": item.model,
                "spec": item.spec,
                "category": item.category,
                "unit": item.unit,
                "unit_price": item.unit_price,
                "source": Path(filename).name,
            }
        )
    return records


def _parse_docx(
    content: bytes,
    started: float,
) -> tuple[list[ParsedPriceItem], str, int, int, list[str]]:
    warnings: list[str] = []
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entry = archive.getinfo("word/document.xml")
            if entry.file_size > MAX_DOCX_XML_BYTES:
                raise ValueError(f"Word 主文档结构超过 {MAX_DOCX_XML_BYTES // 1024 // 1024}MB，已停止价格解析")
            root = ET.fromstring(archive.read(entry))
    except (KeyError, zipfile.BadZipFile, ET.ParseError) as exc:
        raise ValueError(f"无法解析 Word 文档结构：{exc}") from exc

    body = root.find(f"{W}body")
    if body is None:
        return [], "", 0, 0, ["Word 文档没有可读取的正文"]

    paragraphs: list[str] = []
    items: list[ParsedPriceItem] = []
    table_count = 0
    relevant_count = 0
    total_rows = 0
    total_cells = 0

    for child in body:
        _check_deadline(started)
        if child.tag == f"{W}p":
            text = _xml_text(child)
            if text:
                paragraphs.append(text)
            continue
        if child.tag != f"{W}tbl":
            continue

        table_count += 1
        if table_count > MAX_TABLES:
            warnings.append(f"文档表格超过 {MAX_TABLES} 张，后续表格已跳过")
            break
        rows: list[list[str]] = []
        for row in child.findall(f"{W}tr"):
            cells = [_xml_text(cell) for cell in row.findall(f"{W}tc")]
            if any(cells):
                rows.append(cells)
            total_rows += 1
            total_cells += len(cells)
            if total_rows > MAX_TABLE_ROWS or total_cells > MAX_TABLE_CELLS:
                warnings.append("文档表格规模超过价格解析上限，已停止扫描后续表格")
                break
        summary_items = _parse_summary_rows(rows, f"表{table_count}")
        parsed = _parse_rows(rows, f"表{table_count}")
        if summary_items:
            parsed.extend(summary_items)
        if parsed:
            relevant_count += 1
            items.extend(parsed)
            paragraphs.extend(" | ".join(row) for row in rows[:2_000])
        if total_rows > MAX_TABLE_ROWS or total_cells > MAX_TABLE_CELLS:
            break

    full_text = "\n".join(paragraphs)
    if len(full_text) > MAX_TEXT_CHARS:
        full_text = full_text[:MAX_TEXT_CHARS]
        warnings.append("文档文本超过价格解析上限，项目指标抽取仅使用前段内容")
    return items, full_text, table_count, relevant_count, warnings


def _parse_pdf(
    content: bytes,
    started: float,
) -> tuple[list[ParsedPriceItem], str, int, int, list[str]]:
    items: list[ParsedPriceItem] = []
    text_parts: list[str] = []
    warnings: list[str] = []
    table_count = 0
    relevant_count = 0
    with pdfplumber.open(BytesIO(content)) as pdf:
        for page_index, page in enumerate(pdf.pages, 1):
            _check_deadline(started)
            if page_index > 300:
                warnings.append("PDF 超过 300 页，价格解析已停止扫描后续页面")
                break
            text_parts.append(page.extract_text() or "")
            for table_index, table in enumerate(page.extract_tables() or [], 1):
                table_count += 1
                parsed = _parse_rows(table, f"第{page_index}页表{table_index}")
                if parsed:
                    relevant_count += 1
                    items.extend(parsed)
    full_text = "\n".join(text_parts)[:MAX_TEXT_CHARS]
    return items, full_text, table_count, relevant_count, warnings


def _parse_rows(raw_rows: Iterable[Iterable[Any]], source: str) -> list[ParsedPriceItem]:
    rows = [[_text(cell) for cell in row] for row in raw_rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return []
    header, mapping = _find_header(rows)
    if not _is_price_table(rows[header], mapping):
        return []

    header_cells = rows[header]
    items: list[ParsedPriceItem] = []
    for row_number, row in enumerate(rows[header + 1 :], header + 2):
        name = _value(row, mapping.get("name"))
        if not name or _header(name) in {"名称", "产品名称", "功能点", "费用名称", "项目名称"}:
            continue
        evidence = " | ".join(value for value in row if value)
        if not re.search(r"\d", evidence):
            continue
        unit_price_header = _value(header_cells, mapping.get("unit_price"))
        total_price_header = _value(header_cells, mapping.get("total_price"))
        items.append(
            ParsedPriceItem(
                item_name=name,
                brand=_value(row, mapping.get("brand")),
                model=_value(row, mapping.get("model")),
                spec=_value(row, mapping.get("spec")),
                category=_value(row, mapping.get("category")),
                unit=_value(row, mapping.get("unit")) or _money_unit(total_price_header) or "项",
                quantity=_number(_value(row, mapping.get("quantity"))),
                unit_price=_money(_value(row, mapping.get("unit_price")), unit_price_header),
                total_price=_money(_value(row, mapping.get("total_price")), total_price_header),
                effort_person_month=_number(_value(row, mapping.get("effort"))),
                source_section=source,
                source_row=row_number,
                evidence=evidence,
            )
        )
    return items


def _find_header(rows: list[list[str]]) -> tuple[int, dict[str, int]]:
    normalized = {key: {_header(value) for value in values} for key, values in HEADER_ALIASES.items()}
    best_index, best_mapping, best_score = 0, {}, -1
    for index, row in enumerate(rows[:12]):
        mapping: dict[str, int] = {}
        for column, cell in enumerate(row):
            value = _header(cell)
            if len(value) > 40:
                continue
            for field, aliases in normalized.items():
                if value in aliases or any(alias and alias in value for alias in aliases):
                    mapping[field] = column
        if "name" not in mapping and "category" in mapping and ({"unit_price", "total_price"} & mapping.keys()):
            mapping["name"] = mapping["category"]
        score = len(mapping) + (3 if "name" in mapping else 0) + len({"unit_price", "total_price", "effort"} & mapping.keys())
        if score > best_score:
            best_index, best_mapping, best_score = index, mapping, score
    return best_index, best_mapping


def _is_price_table(header_row: list[str], mapping: dict[str, int]) -> bool:
    if "name" not in mapping:
        return False
    financial = bool({"unit_price", "total_price", "effort"} & mapping.keys())
    described_product = "quantity" in mapping and bool({"brand", "model", "spec", "category"} & mapping.keys())
    header_text = " ".join(header_row)
    return financial or (described_product and any(keyword in header_text for keyword in PRICE_HEADER_KEYWORDS))


def _parse_summary_rows(rows: list[list[str]], source: str) -> list[ParsedPriceItem]:
    if not rows:
        return []
    title = " ".join(rows[0])
    if not any(keyword in title for keyword in ("投资估算总表", "投资概算总表", "项目投资汇总表")):
        return []
    header_hint = "万元" if "万元" in title else "元"
    items: list[ParsedPriceItem] = []
    for row_number, row in enumerate(rows[1:], 2):
        values = [_text(value) for value in row if _text(value)]
        if len(values) < 2:
            continue
        label_index = 1 if re.fullmatch(r"[\d.\s]+", values[0]) and len(values) >= 3 else 0
        label = values[label_index]
        amount_index = label_index + 1
        if amount_index >= len(values):
            continue
        amount = _money(values[amount_index], header_hint)
        if not label or amount is None:
            continue
        items.append(
            ParsedPriceItem(
                item_name=label,
                category="投资估算汇总",
                unit=header_hint,
                quantity=1,
                total_price=amount,
                source_section=source,
                source_row=row_number,
                evidence=" | ".join(values),
            )
        )
    return items


def _build_project_context(items: list[ParsedPriceItem], full_text: str) -> PriceProjectContext:
    total = _item_amount(items, ("项目总投资", "总投资合计", "投资总额", "总计")) or _labeled_money(full_text, ("项目总投资", "总投资合计", "投资总额"))
    direct = _item_amount(items, ("直接建设费", "项目直接建设费", "直接建设费用", "系统建设费")) or _labeled_money(full_text, ("直接建设费", "项目直接建设费", "直接建设费用"))
    software = _item_amount(items, ("应用软件开发费", "软件开发费", "应用开发费", "应用软件开发")) or _labeled_money(full_text, ("应用软件开发费", "软件开发费", "应用开发费"))
    purchase = _item_amount(items, ("软硬件购置费", "软硬件采购费", "设备购置费"))
    if purchase is None:
        purchase = _sum_item_amount(items, ("硬件购置", "产品软件", "安全产品"))
    purchase = purchase or _labeled_money(full_text, ("软硬件购置费", "软硬件采购费", "设备购置费"))
    ratio = software / total if software is not None and total and total > 0 else None
    operation_hits = sum(full_text.count(keyword) for keyword in PROJECT_TYPE_OPERATION_KEYWORDS)
    construction_hits = sum(full_text.count(keyword) for keyword in PROJECT_TYPE_CONSTRUCTION_KEYWORDS)
    project_type = "operation" if operation_hits > construction_hits else "construction" if construction_hits else "unknown"
    security_level = _security_level(full_text)
    return PriceProjectContext(
        project_type=project_type,
        total_investment=total,
        direct_construction_cost=direct,
        software_development_cost=software,
        hardware_software_purchase_cost=purchase,
        application_development_ratio=ratio,
        security_level=security_level,
    )


def _item_amount(items: list[ParsedPriceItem], aliases: tuple[str, ...]) -> float | None:
    candidates: list[float] = []
    for item in items:
        name = _norm(item.item_name)
        if not any(_norm(alias) in name for alias in aliases):
            continue
        amount = item.total_price
        if amount is None and item.unit_price is not None:
            amount = item.unit_price * (item.quantity or 1)
        if amount is not None and amount > 0:
            candidates.append(amount)
    return max(candidates) if candidates else None


def _sum_item_amount(items: list[ParsedPriceItem], aliases: tuple[str, ...]) -> float | None:
    amounts: list[float] = []
    for alias in aliases:
        amount = _item_amount(items, (alias,))
        if amount is not None:
            amounts.append(amount)
    return sum(amounts) if amounts else None


def _labeled_money(text: str, aliases: tuple[str, ...]) -> float | None:
    for line in text.splitlines():
        compact = " ".join(line.split())
        for alias in aliases:
            if alias not in compact:
                continue
            tail = compact[compact.find(alias) + len(alias) :]
            match = re.search(r"(?:为|是|：|:|合计)?\s*([\d,]+(?:\.\d+)?)\s*(万元|万|千元|元)", tail[:80])
            if match:
                return _money(match.group(1) + match.group(2))
    return None


def _security_level(text: str) -> int | None:
    match = re.search(r"(?:等保|等级保护|安全保护等级|信息系统安全等级)\s*(?:为|：|:|第)?\s*([二三23])\s*级", text)
    if not match:
        return None
    return {"二": 2, "三": 3, "2": 2, "3": 3}.get(match.group(1))


def _xml_text(element: ET.Element) -> str:
    return " ".join("".join(node.itertext()).strip() for node in element.iter(f"{W}t") if "".join(node.itertext()).strip())


def _check_deadline(started: float) -> None:
    if time.monotonic() - started > MAX_PARSE_SECONDS:
        raise ValueError(f"价格文档解析超过 {MAX_PARSE_SECONDS:g} 秒，已停止以避免阻塞后端")


def _value(row: list[str], index: int | None) -> str:
    return row[index].strip() if index is not None and index < len(row) else ""


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _header(value: Any) -> str:
    return re.sub(r"[\s_\-（）()：:]+", "", _text(value)).lower()


def _norm(value: Any) -> str:
    return re.sub(r"[\s\-_/（）()，,。．.]+", "", _text(value)).lower()


def _number(value: Any) -> float | None:
    normalized = re.sub(r"\s+", "", _text(value)).replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", normalized)
    return float(match.group(0)) if match else None


def _money(value: Any, header_hint: str = "") -> float | None:
    text = _text(value).replace(",", "")
    number = _number(text)
    if number is None:
        return None
    unit_text = f"{text} {header_hint}"
    if "万元" in unit_text or re.search(r"\d\s*万(?:\s|$)", text):
        return number * 10_000
    if "千元" in unit_text:
        return number * 1_000
    return number


def _money_unit(header: str) -> str:
    if "万元" in header:
        return "万元"
    if "千元" in header:
        return "千元"
    if "元" in header:
        return "元"
    return ""


def _dedupe(items: list[ParsedPriceItem]) -> list[ParsedPriceItem]:
    seen: set[tuple[Any, ...]] = set()
    output: list[ParsedPriceItem] = []
    for item in items:
        key = (_norm(item.item_name), _norm(item.brand), _norm(item.model), item.quantity, item.unit_price, item.total_price, item.source_section, item.source_row)
        if key not in seen:
            seen.add(key)
            output.append(item)
    return output
