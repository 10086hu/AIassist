from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from io import BytesIO, StringIO
from typing import Any, Callable, Iterable

from sqlalchemy.orm import Session

from app.modules.price.parser import parse_price_document
from app.modules.price.service import _evaluate_products, _is_detail_item, _is_product, _is_software
from app.modules.resource.rules import ResourceRuleFinding


@dataclass(frozen=True)
class MaintenanceRule:
    rule_id: str
    rule_number: int
    excel_row: int
    rule_name: str
    rule_category: str
    rule_description: str
    judgement_condition: str


@dataclass(frozen=True)
class MaintenanceDocument:
    filename: str
    raw_text: str
    lines: list[str]
    table_rows: list[list[str]]


MAINTENANCE_RULES: tuple[MaintenanceRule, ...] = (
    MaintenanceRule(
        "MAINT_OPS_001",
        1,
        2,
        "申报总额一致性校验规则",
        "运维项目 / 一致性校验规则",
        "运维方案、自查表及系统填写的申报金额应保持一致。",
        "提取运维方案、自查表及系统填报申报金额，并进行一致性校验。",
    ),
    MaintenanceRule(
        "MAINT_OPS_002",
        2,
        3,
        "申报内容审核规则",
        "运维项目 / 内容合规性审查规则",
        "审核运维项目申报内容，对未拆分细项报价的项目进行提示。",
        "申报内容应拆分至各个硬件或应用软件模块单独报价。",
    ),
    MaintenanceRule(
        "MAINT_OPS_003",
        3,
        4,
        "运维方案架构完整性审核规则",
        "运维项目 / 内容合规性审查规则",
        "运维方案内容框架清晰，内容完善。",
        "判断运维方案是否包含项目概述、需求分析、现状分析、项目运维方案、项目实施进度、人员安排及投资估算。",
    ),
    MaintenanceRule(
        "MAINT_OPS_004",
        4,
        5,
        "运维方案敏感词审核规则",
        "运维项目 / 内容合规性审查规则",
        "运维方案中不宜出现厂商名字、运维成员名字等敏感词。",
        "检查全文中是否出现厂商名字、运维人员名字等需复核内容。",
    ),
    MaintenanceRule(
        "MAINT_OPS_005",
        5,
        6,
        "自查表填报完整性审核规则",
        "运维项目 / 内容合规性审查规则",
        "检查项目自查表数据完整性、准确性、规范性情况。",
        "核查自查表是否完成签字盖章、是否存在未填报内容，并列明缺失字段。",
    ),
    MaintenanceRule(
        "MAINT_OPS_006",
        6,
        7,
        "自查表填报内容提取规则",
        "运维项目 / 内容合规性审查规则",
        "对各个运维项目的自查表内容进行提取。",
        "提取数据量、上云情况、人员配备、匹配度、重大事件、故障次数、满意度、响应率、解决率等字段。",
    ),
    MaintenanceRule(
        "MAINT_OPS_007",
        7,
        8,
        "运维台账完整性审核规则",
        "运维项目 / 内容合规性审查规则",
        "核查运维台账的完整性、准确性、规范性情况。",
        "核验运维台账是否签字盖章，业务时间区间是否完整覆盖全年周期。",
    ),
    MaintenanceRule(
        "MAINT_OPS_008",
        8,
        9,
        "运维频率提取规则",
        "运维项目 / 内容合规性审查规则",
        "提取运维台账中的时间，并计算时间节点间隔。",
        "抓取台账内所有业务时间字段，自动计算各时间节点间隔，核算运维工作执行频率。",
    ),
    MaintenanceRule(
        "MAINT_OPS_009",
        9,
        10,
        "运维人员提取规则",
        "运维项目 / 内容合规性审查规则",
        "提取运维台账中负责运维人员的名字及人数。",
        "提取运维负责人姓名，统计不重复运维人员姓名总数量，并与自查表人数数量核对。",
    ),
    MaintenanceRule(
        "MAINT_OPS_010",
        10,
        11,
        "应用软件运维申报项审核规则",
        "运维项目 / 必要性原则",
        "不属于应用软件及产品软件运维申报范围的申报项，核定价格为 0。",
        "页面设计与框架搭建、数据对接迁移、实施联调测试不属于应用软件及产品软件运维项目申报范围。",
    ),
    MaintenanceRule(
        "MAINT_OPS_011",
        11,
        12,
        "硬件运维申报项审核规则",
        "运维项目 / 必要性原则",
        "不属于硬件运维申报范围的申报项，核定价格为 0。",
        "机柜、综合布线管道、立杆支架、集成费、办公电脑、打印机、辅材不属于硬件运维申报范围。",
    ),
    MaintenanceRule(
        "MAINT_OPS_012",
        12,
        13,
        "雪亮工程项目外包服务费申报项审核原则",
        "运维项目 / 计量计价规则",
        "街镇雪亮工程项目外包服务费核定价格统一为 19 万元。",
        "识别雪亮工程项目外包服务费，并核对核定价格是否为 19 万元。",
    ),
    MaintenanceRule(
        "MAINT_OPS_013",
        13,
        14,
        "财政局项目外包服务费申报项审核原则",
        "运维项目 / 计量计价规则",
        "财政局申报项目外包服务费以申报金额为准。",
        "针对财政局申报项目存在外包服务费的情况，核定价格予以支持，以申报金额为准。",
    ),
    MaintenanceRule(
        "MAINT_OPS_014",
        14,
        15,
        "应用软件运维明细核价规则",
        "运维项目 / 计量计价规则",
        "应用软件运维项目按年限、申报金额和购置金额费率核价。",
        "小额三年以上且无新增项目按去年核定；其他项目按购置金额 6%/4% 口径与申报金额比较。",
    ),
    MaintenanceRule(
        "MAINT_OPS_015",
        15,
        16,
        "产品软件运维明细核价规则",
        "运维项目 / 计量计价规则",
        "产品软件运维项目按年限、申报金额和购置金额费率核价。",
        "小额三年以上项目按去年核定；其他项目按购置金额 4% 与申报金额比较。",
    ),
    MaintenanceRule(
        "MAINT_OPS_016",
        16,
        17,
        "硬件运维明细核价规则",
        "运维项目 / 计量计价规则",
        "硬件运维项目按年限、申报金额和购置金额费率核价。",
        "小额三年以上项目按去年核定；其他项目按购置金额 4% 与申报金额比较。",
    ),
    MaintenanceRule(
        "MAINT_OPS_017",
        17,
        18,
        "等保三级测评费核价规则",
        "运维项目 / 计量计价规则",
        "等保三级测评费按项目立项情况和系统金额核价。",
        "10-100 万元系统核定 4 万元，100 万元以上系统核定 9 万元；已立项审核项目按数据局核定金额。",
    ),
    MaintenanceRule(
        "MAINT_OPS_018",
        18,
        19,
        "等保二级测评费核价规则",
        "运维项目 / 计量计价规则",
        "等保二级测评费按项目立项情况和系统金额核价。",
        "10-100 万元系统核定 3 万元，100 万元以上系统核定 6 万元；已立项审核项目按数据局核定金额。",
    ),
    MaintenanceRule(
        "MAINT_OPS_019",
        19,
        20,
        "历史价格比较规则",
        "运维项目 / 计量计价规则",
        "各运维项目核定价格应小于或等于去年核定价格。",
        "将本年度核定价格与去年核定价格进行对比，汇总核定价格高于去年核定价格的项目。",
    ),
    MaintenanceRule(
        "MAINT_OPS_020",
        20,
        21,
        "新增软硬件产品核价规则",
        "运维项目 / 计量计价规则",
        "新增软硬件产品应参考市级和区级产品库价格。",
        "识别新增硬件、产品软件细项，并匹配上海市产品库、松江区产品库参考价。",
    ),
)


RULE_BY_ID = {rule.rule_id: rule for rule in MAINTENANCE_RULES}
NUMBER_TEXT = r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?"
AMOUNT_PATTERN = re.compile(
    rf"(?P<number>{NUMBER_TEXT})\s*(?P<unit>万元|万|元)?",
    re.IGNORECASE,
)
DATE_PATTERN = re.compile(
    r"(20\d{2}[年/-]\d{1,2}[月/-]\d{0,2}|"
    r"\d{1,2}[月/-]\d{1,2}[日]?|"
    r"\d{1,2}\s*月)"
)
CHINESE_NAME_PATTERN = re.compile(r"(?:负责人|运维人员|维护人员|联系人|工程师|姓名)[:：\s]*([\u4e00-\u9fa5]{2,4})")
MAINTENANCE_FILENAME_KEYWORDS = ("运维", "运行维护", "维护服务", "运营维护", "维保", "续保")
MAINTENANCE_EXPLICIT_PROJECT_KEYWORDS = (
    "运维项目",
    "运行维护项目",
    "运营维护项目",
    "维护服务项目",
    "年度运维",
    "日常运维",
    "维保项目",
    "续保项目",
)
MAINTENANCE_EXPLICIT_MATERIAL_KEYWORDS = (
    "运维方案",
    "维护方案",
    "运行维护方案",
    "运维台账",
    "运维工单",
    "运维自查表",
    "运维核价清单",
    "运维合同",
    "维护服务合同",
    "应用软件运维",
    "产品软件运维",
    "硬件运维",
)
MAINTENANCE_STRUCTURED_MATERIAL_KEYWORDS = (
    "核价清单",
    "核定运维金额",
    "申报运维金额",
    "运维金额",
    "运维单价",
    "运维数量",
    "运维台账",
    "运维工单",
)
CONSTRUCTION_DOCUMENT_KEYWORDS = (
    "可行性研究",
    "可研报告",
    "建设项目",
    "建设方案",
    "初步设计",
    "新建项目",
    "改扩建",
    "应用软件开发投资估算",
)
MAINTENANCE_STRONG_KEYWORDS = (
    "运维项目",
    "运行维护项目",
    "运营维护项目",
    "运维方案",
    "维护方案",
    "运维台账",
    "运维服务",
    "运维内容",
    "运维范围",
    "应用软件运维",
    "产品软件运维",
    "硬件运维",
    "外包服务费",
    "核定价格",
)
MAINTENANCE_SUPPORT_KEYWORDS = (
    "自查表",
    "去年核定",
    "上年核定",
    "本年申报",
    "申报金额",
    "购置金额",
    "服务期",
    "巡检",
    "响应率",
    "解决率",
    "满意度",
    "故障次数",
    "等保测评",
    "数据局核定",
    "上海市产品库",
    "松江区产品库",
)
MAINTENANCE_2GRAM_STRONG_THRESHOLD = 0.60
MAINTENANCE_2GRAM_SUPPORT_THRESHOLD = 0.75


def maintenance_rules_to_public_dicts() -> list[dict[str, str]]:
    return [
        {
            "rule_id": rule.rule_id,
            "rule_name": rule.rule_name,
            "rule_category": rule.rule_category,
            "rule_detail": rule.rule_description,
            "source": "local_builtin",
        }
        for rule in MAINTENANCE_RULES
    ]


def is_maintenance_project_document(content: bytes, filename: str) -> bool:
    """Return True only when the uploaded material looks like an O&M project.

    运维规则目前作为“资源申请合理性”下方的专项试用版存在。这里先做轻量门禁：
    只有文件名、标题区或结构化材料能明确指向运维项目/运维材料时，才自动运行 20 条运维规则。
    """

    filename_text = str(filename or "")
    if _has_explicit_maintenance_signal(filename_text):
        return True
    if _contains_any(filename_text, MAINTENANCE_FILENAME_KEYWORDS) and not _looks_like_construction_document(filename_text):
        return True
    if _looks_like_construction_document(filename_text):
        return False

    try:
        document = parse_maintenance_document(content, filename)
    except Exception:
        return False

    text = f"{filename}\n{document.raw_text}"
    title_scope = text[:5000]
    if _has_explicit_maintenance_signal(title_scope):
        return True

    looks_like_construction = _looks_like_construction_document(title_scope)
    operation_term_count = text.count("运维") + text.count("运行维护") + text.count("维护服务")
    if _contains_any(title_scope, MAINTENANCE_STRUCTURED_MATERIAL_KEYWORDS) and operation_term_count >= 2:
        return True
    if looks_like_construction:
        return False

    strong_hits = {keyword for keyword in MAINTENANCE_STRONG_KEYWORDS if keyword in title_scope}
    support_hits = {keyword for keyword in MAINTENANCE_SUPPORT_KEYWORDS if keyword in title_scope}
    fuzzy_strong_hits = _fuzzy_2gram_keyword_hits(
        title_scope,
        MAINTENANCE_STRONG_KEYWORDS,
        threshold=MAINTENANCE_2GRAM_STRONG_THRESHOLD,
    )
    fuzzy_support_hits = _fuzzy_2gram_keyword_hits(
        title_scope,
        MAINTENANCE_SUPPORT_KEYWORDS,
        threshold=MAINTENANCE_2GRAM_SUPPORT_THRESHOLD,
    )

    if len(strong_hits) >= 2 and operation_term_count >= 3:
        return True
    if strong_hits and operation_term_count >= 5 and len(support_hits | fuzzy_support_hits) >= 2:
        return True
    if strong_hits and fuzzy_strong_hits and operation_term_count >= 8:
        return True
    return False


def _has_explicit_maintenance_signal(value: str) -> bool:
    return _contains_any(value, MAINTENANCE_EXPLICIT_PROJECT_KEYWORDS) or _contains_any(value, MAINTENANCE_EXPLICIT_MATERIAL_KEYWORDS)


def _looks_like_construction_document(value: str) -> bool:
    return _contains_any(value, CONSTRUCTION_DOCUMENT_KEYWORDS)


def evaluate_maintenance_rules(
    content: bytes,
    filename: str,
    selected_rule_ids: Iterable[str] | None = None,
    db: Session | None = None,
) -> list[ResourceRuleFinding]:
    """Run the Songjiang maintenance-project rules under the resource module.

    当前没有运维项目样例，因此本模块采用保守启发式：
    能从文本或表格中明确识别到金额、字段、关键词时给出通过/风险；
    识别不到充分依据时给出“需人工确认”，避免伪造结论。
    """

    selected = {str(item).strip() for item in selected_rule_ids or [] if str(item).strip()}
    document = parse_maintenance_document(content, filename)
    evaluators: dict[str, Callable[[MaintenanceDocument, MaintenanceRule], ResourceRuleFinding]] = {
        "MAINT_OPS_001": _evaluate_total_amount_consistency,
        "MAINT_OPS_002": _evaluate_itemized_quote,
        "MAINT_OPS_003": _evaluate_plan_structure,
        "MAINT_OPS_004": _evaluate_sensitive_terms,
        "MAINT_OPS_005": _evaluate_self_check_completeness,
        "MAINT_OPS_006": _evaluate_self_check_extraction_fields,
        "MAINT_OPS_007": _evaluate_operation_ledger_completeness,
        "MAINT_OPS_008": _evaluate_operation_frequency,
        "MAINT_OPS_009": _evaluate_operation_staff,
        "MAINT_OPS_010": _evaluate_application_scope_exclusions,
        "MAINT_OPS_011": _evaluate_hardware_scope_exclusions,
        "MAINT_OPS_012": _evaluate_xueliang_outsourcing_fee,
        "MAINT_OPS_013": _evaluate_finance_bureau_outsourcing_fee,
        "MAINT_OPS_014": lambda doc, rule: _evaluate_maintenance_price_rule(doc, rule, ("应用软件", "软件开发", "应用系统"), 0.06, 0.04),
        "MAINT_OPS_015": lambda doc, rule: _evaluate_maintenance_price_rule(doc, rule, ("产品软件", "成品软件", "软件产品"), 0.04, 0.04),
        "MAINT_OPS_016": lambda doc, rule: _evaluate_maintenance_price_rule(doc, rule, ("硬件", "服务器", "网络设备", "存储", "终端"), 0.04, 0.04),
        "MAINT_OPS_017": lambda doc, rule: _evaluate_security_assessment_fee(doc, rule, level="三级", low_fee=4.0, high_fee=9.0),
        "MAINT_OPS_018": lambda doc, rule: _evaluate_security_assessment_fee(doc, rule, level="二级", low_fee=3.0, high_fee=6.0),
        "MAINT_OPS_019": _evaluate_history_price_compare,
        "MAINT_OPS_020": lambda doc, rule: _evaluate_new_product_price_reference(doc, rule, db=db, content=content, filename=filename),
    }

    findings: list[ResourceRuleFinding] = []
    for rule in MAINTENANCE_RULES:
        if not _rule_selected(rule, selected):
            continue
        findings.append(evaluators[rule.rule_id](document, rule))
    return findings


def parse_maintenance_document(content: bytes, filename: str) -> MaintenanceDocument:
    lower = filename.lower()
    if lower.endswith(".docx"):
        raw_text, rows = _parse_docx(content)
    elif lower.endswith(".xlsx"):
        raw_text, rows = _parse_xlsx(content)
    elif lower.endswith(".csv"):
        raw_text, rows = _parse_csv(content)
    else:
        raise ValueError("运维项目规则支持 .docx、.xlsx、.csv 文件")

    lines = [line.strip() for line in re.split(r"[\r\n]+", raw_text) if line.strip()]
    return MaintenanceDocument(filename=filename, raw_text=raw_text, lines=lines, table_rows=rows)


def _parse_docx(content: bytes) -> tuple[str, list[list[str]]]:
    from docx import Document

    document = Document(BytesIO(content))
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    rows: list[list[str]] = []
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells if cell.text and cell.text.strip()]
            if values:
                rows.append(values)
                lines.append(" | ".join(values))
    return "\n".join(lines), rows


def _parse_xlsx(content: bytes) -> tuple[str, list[list[str]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), data_only=True)
    rows: list[list[str]] = []
    lines: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values = [_cell_text(cell) for cell in row]
            values = [value for value in values if value]
            if values:
                rows.append(values)
                lines.append(f"{sheet.title} | " + " | ".join(values))
    return "\n".join(lines), rows


def _parse_csv(content: bytes) -> tuple[str, list[list[str]]]:
    text = content.decode("utf-8-sig", errors="ignore")
    rows = [[cell.strip() for cell in row if cell.strip()] for row in csv.reader(StringIO(text))]
    rows = [row for row in rows if row]
    return "\n".join(" | ".join(row) for row in rows), rows


def _evaluate_total_amount_consistency(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    source_keywords = {
        "运维方案申报金额": ("运维方案", "方案", "申报金额", "申报总额"),
        "自查表申报金额": ("自查表", "申报金额", "申报总额"),
        "系统填报申报金额": ("系统填报", "系统填写", "申报金额", "申报总额"),
    }
    values: dict[str, float] = {}
    evidence: list[str] = []
    for source_name, keywords in source_keywords.items():
        for line in doc.lines:
            if _contains_any(line, keywords):
                amount = _largest_amount(line)
                if amount is not None:
                    values[source_name] = amount
                    evidence.append(line)
                    break

    if len(values) < 2:
        return _finding(
            rule,
            "申报总额",
            "需人工确认",
            "资料不足",
            "未同时识别到运维方案、自查表及系统填报中的申报金额，无法完成三方一致性校验。",
            "请在材料中明确运维方案、自查表、系统填报三个来源的申报金额，便于自动比对。",
            values,
            evidence or _evidence(doc, ("申报金额", "申报总额", "预算", "金额")),
        )

    distinct = {round(value, 4) for value in values.values()}
    if len(distinct) == 1:
        return _pass(rule, "申报总额", f"已识别到 {len(values)} 个来源的申报金额，金额一致。", values, evidence)

    return _finding(
        rule,
        "申报总额",
        "高",
        "金额不一致",
        f"识别到不同来源申报金额不一致：{_format_quantities(values)}。",
        "请核对运维方案、自查表和系统填报金额，统一申报总额后重新提交。",
        values,
        evidence,
    )


def _evaluate_itemized_quote(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    detail_rows = [
        " | ".join(row)
        for row in doc.table_rows
        if _contains_any(" ".join(row), ("硬件", "软件", "模块", "明细", "服务项", "申报项", "报价", "金额"))
        and _largest_amount(" ".join(row)) is not None
    ]
    total_amounts = _amounts_near_keywords(doc.lines, ("申报金额", "申报总额", "合计", "总计"))
    if detail_rows:
        return _pass(
            rule,
            "申报内容明细报价",
            f"已识别到 {len(detail_rows)} 条疑似硬件、软件或模块级明细报价。",
            {"明细报价条数": float(len(detail_rows))},
            detail_rows[:5],
        )
    if total_amounts:
        return _finding(
            rule,
            "申报内容明细报价",
            "中",
            "疑似未拆分报价",
            "识别到项目总金额，但未识别到硬件、软件模块或服务项层面的明细报价。",
            "请将申报内容拆分至硬件、应用软件模块或服务项，并分别列示金额。",
            {"总额记录数": float(len(total_amounts))},
            _evidence(doc, ("申报金额", "合计", "总计", "报价")),
        )
    return _manual(rule, "申报内容明细报价", "未识别到申报金额或明细报价表，无法判断是否已拆分细项报价。", "请提供包含申报金额和明细报价的运维方案或报价表。", doc)


def _evaluate_plan_structure(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    required = ("项目概述", "需求分析", "现状分析", "项目运维方案", "项目实施进度", "人员安排", "投资估算")
    missing = [item for item in required if item not in doc.raw_text]
    if not missing:
        return _pass(rule, "运维方案章节架构", "已识别到运维方案要求的主要章节。", {"章节数量": float(len(required))}, _evidence(doc, required))
    return _finding(
        rule,
        "运维方案章节架构",
        "中",
        "章节不完整",
        f"未识别到以下必备章节：{'、'.join(missing)}。",
        "请补充缺失章节，保持运维方案框架完整。",
        {"缺失章节数": float(len(missing))},
        _evidence(doc, required),
    )


def _evaluate_sensitive_terms(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    vendor_terms = (
        "Oracle", "Microsoft", "微软", "IBM", "Cisco", "思科", "VMware", "Windows Server",
        "SQL Server", "Red Hat", "CentOS", "华为", "新华三", "H3C", "浪潮", "联想",
    )
    name_lines = [line for line in doc.lines if CHINESE_NAME_PATTERN.search(line)]
    vendor_lines = _evidence(doc, vendor_terms, limit=8)
    evidence = vendor_lines + name_lines[:5]
    if evidence:
        return _finding(
            rule,
            "运维敏感词",
            "需人工确认",
            "命中敏感表达",
            "识别到厂商名称或疑似运维人员姓名，需要结合上下文判断是否允许出现。",
            "请确认相关厂商、人名是否为必要现状描述；如非必要，建议删除或脱敏。",
            {"命中片段数": float(len(evidence))},
            evidence,
        )
    return _pass(rule, "运维敏感词", "未识别到内置厂商名称或疑似运维人员姓名。", {}, [])


def _evaluate_self_check_completeness(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    evidence = _evidence(doc, ("自查表", "签字", "盖章", "未填", "空白"))
    if "自查表" not in doc.raw_text:
        return _manual(rule, "自查表完整性", "未识别到自查表内容，无法判断签字盖章和空白项。", "请上传或补充项目自查表。", doc)
    missing_flags = [term for term in ("未填", "空白", "缺失", "未盖章", "未签字") if term in doc.raw_text]
    if missing_flags:
        return _finding(
            rule,
            "自查表完整性",
            "中",
            "自查表疑似不完整",
            f"自查表中出现疑似不完整标记：{'、'.join(missing_flags)}。",
            "请补齐自查表空白项，并确认签字盖章齐全。",
            {"缺失标记数": float(len(missing_flags))},
            evidence,
        )
    if "签字" in doc.raw_text and "盖章" in doc.raw_text:
        return _pass(rule, "自查表完整性", "已识别到自查表及签字、盖章相关表述。", {}, evidence)
    return _manual(rule, "自查表完整性", "已识别到自查表，但未同时识别到签字和盖章表述。", "请人工核对自查表是否完成签字盖章。", doc, evidence)


def _evaluate_self_check_extraction_fields(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    fields = (
        "今年产生数据量", "上云情况", "人员配备", "匹配度", "重大事件", "故障", "满意度",
        "单次最长故障时间", "服务响应及时率", "问题解决及时率", "问题平均解决时间",
    )
    found = [field for field in fields if field in doc.raw_text]
    evidence = _evidence(doc, fields)
    if len(found) >= 7:
        return _pass(rule, "自查表字段提取", f"已识别到 {len(found)} 个自查表关键填报字段。", {"识别字段数": float(len(found))}, evidence)
    return _finding(
        rule,
        "自查表字段提取",
        "需人工确认",
        "字段不足",
        f"仅识别到 {len(found)} 个自查表关键字段，可能无法形成完整汇总表。",
        "请确认自查表是否包含数据量、上云、人员、故障、满意度、响应率和解决率等字段。",
        {"识别字段数": float(len(found))},
        evidence,
    )


def _evaluate_operation_ledger_completeness(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    if "台账" not in doc.raw_text:
        return _manual(rule, "运维台账完整性", "未识别到运维台账内容，无法判断台账完整性。", "请上传运维台账或在材料中保留台账表格。", doc)
    evidence = _evidence(doc, ("台账", "签字", "盖章", "全年", "1月", "12月", "202"))
    has_signature = "签字" in doc.raw_text and "盖章" in doc.raw_text
    has_full_year = _has_full_year_coverage(doc.raw_text)
    if has_signature and has_full_year:
        return _pass(rule, "运维台账完整性", "已识别到台账、签字盖章和全年周期相关依据。", {}, evidence)
    missing = []
    if not has_signature:
        missing.append("签字盖章")
    if not has_full_year:
        missing.append("全年周期覆盖")
    return _finding(
        rule,
        "运维台账完整性",
        "需人工确认",
        "台账依据不足",
        f"运维台账未充分体现：{'、'.join(missing)}。",
        "请确认台账签字盖章齐全，并补充覆盖全年周期的业务时间记录。",
        {"缺失项数": float(len(missing))},
        evidence,
    )


def _evaluate_operation_frequency(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    ledger_lines = [line for line in doc.lines if "台账" in line or "运维" in line or "业务时间" in line]
    dates = []
    for line in ledger_lines or doc.lines:
        dates.extend(match.group(0) for match in DATE_PATTERN.finditer(line))
    evidence = (ledger_lines[:5] if ledger_lines else _evidence(doc, ("业务时间", "运维", "台账")))[:8]
    if len(dates) >= 2:
        return _pass(
            rule,
            "运维频率",
            f"已识别到 {len(set(dates))} 个业务时间节点，可用于人工复核运维频率。",
            {"时间节点数": float(len(set(dates)))},
            evidence,
        )
    return _manual(rule, "运维频率", "未识别到足够的台账业务时间节点，无法计算运维频率。", "请提供包含多次业务时间记录的运维台账。", doc, evidence)


def _evaluate_operation_staff(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    names = sorted({match.group(1) for line in doc.lines for match in CHINESE_NAME_PATTERN.finditer(line)})
    evidence = [line for line in doc.lines if CHINESE_NAME_PATTERN.search(line)][:8]
    declared_counts = _numbers_near_keywords(doc.lines, ("人员配备", "运维人员", "人数", "专职人员"))
    quantities: dict[str, float] = {"识别人员数": float(len(names))}
    if declared_counts:
        quantities["自查表疑似人数"] = declared_counts[0]
    if names and declared_counts and round(float(len(names)), 4) != round(declared_counts[0], 4):
        return _finding(
            rule,
            "运维人员",
            "中",
            "人员数量不一致",
            f"台账识别到 {len(names)} 名运维人员，自查表疑似填报 {declared_counts[0]:g} 人。",
            "请核对运维台账负责人姓名与自查表人员数量是否一致。",
            quantities,
            evidence,
        )
    if names:
        return _pass(rule, "运维人员", f"已识别到 {len(names)} 名疑似运维人员。", quantities, evidence)
    return _manual(rule, "运维人员", "未识别到运维负责人姓名，无法统计不重复运维人员数量。", "请提供包含运维负责人或维护人员姓名的台账。", doc)


def _evaluate_application_scope_exclusions(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    return _evaluate_exclusion_terms(
        doc,
        rule,
        "应用软件运维申报范围",
        ("页面设计与框架搭建", "页面设计", "框架搭建", "数据对接迁移", "数据迁移", "实施联调测试", "联调测试"),
        "识别到不属于应用软件及产品软件运维申报范围的申报项，按规则应核定为 0。",
    )


def _evaluate_hardware_scope_exclusions(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    return _evaluate_exclusion_terms(
        doc,
        rule,
        "硬件运维申报范围",
        ("机柜", "综合布线管道", "综合布线", "立杆支架", "集成费", "办公电脑", "打印机", "辅材"),
        "识别到不属于硬件运维申报范围的申报项，按规则应核定为 0。",
    )


def _evaluate_xueliang_outsourcing_fee(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    if "雪亮工程" not in doc.raw_text:
        return _pass(rule, "雪亮工程外包服务费", "未识别到雪亮工程外包服务费场景。", {}, [])
    evidence = _evidence(doc, ("雪亮工程", "外包服务费", "19万", "19 万"))
    amounts = _amounts_near_keywords(doc.lines, ("雪亮工程", "外包服务费"))
    if any(_approx_equal(amount, 19.0, 0.05) for amount in amounts):
        return _pass(rule, "雪亮工程外包服务费", "识别到雪亮工程外包服务费金额为 19 万元。", {"识别金额": 19.0}, evidence)
    if amounts:
        return _finding(
            rule,
            "雪亮工程外包服务费",
            "高",
            "核价不一致",
            f"雪亮工程外包服务费应为 19 万元，识别金额为：{', '.join(f'{item:g}' for item in amounts[:5])} 万元。",
            "请按规则将雪亮工程外包服务费核定为 19 万元，或补充特殊依据。",
            {"规则金额": 19.0, "识别金额": amounts[0]},
            evidence,
        )
    return _manual(rule, "雪亮工程外包服务费", "识别到雪亮工程场景，但未提取到外包服务费金额。", "请补充外包服务费金额并核对是否为 19 万元。", doc, evidence)


def _evaluate_finance_bureau_outsourcing_fee(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    if not _contains_any(doc.raw_text, ("财政局", "财政")):
        return _pass(rule, "财政局外包服务费", "未识别到财政局外包服务费场景。", {}, [])
    evidence = _evidence(doc, ("财政局", "外包服务费", "申报金额", "核定价格"))
    if "外包服务费" not in doc.raw_text:
        return _manual(rule, "财政局外包服务费", "识别到财政局项目，但未识别到外包服务费申报项。", "如存在外包服务费，请明确申报金额；如不存在，可忽略本规则。", doc, evidence)
    amounts = _amounts_near_keywords(doc.lines, ("财政局", "外包服务费", "申报金额"))
    if amounts:
        return _pass(rule, "财政局外包服务费", "识别到财政局外包服务费申报金额，规则口径为以申报金额为准。", {"申报金额": amounts[0]}, evidence)
    return _manual(rule, "财政局外包服务费", "识别到财政局外包服务费，但未提取到申报金额。", "请补充财政局外包服务费申报金额。", doc, evidence)


def _evaluate_maintenance_price_rule(
    doc: MaintenanceDocument,
    rule: MaintenanceRule,
    applicability_keywords: tuple[str, ...],
    compare_rate: float,
    approved_rate: float,
) -> ResourceRuleFinding:
    if not _contains_any(doc.raw_text, applicability_keywords):
        return _pass(rule, rule.rule_name.replace("规则", ""), "未识别到该类运维明细场景。", {}, [])

    evidence = _evidence(doc, applicability_keywords + ("申报金额", "购置金额", "去年核定", "三年", "新增"))
    declared = _first_labeled_amount_near(doc.lines, ("申报金额", "申报价", "报价"))
    purchase = _first_labeled_amount_near(doc.lines, ("购置金额", "购买金额", "原值", "资产原值"))
    last_year = _first_labeled_amount_near(doc.lines, ("去年核定", "上年核定", "历史核定"))
    is_small_legacy = declared is not None and declared < 5 and _contains_any(doc.raw_text, ("三年以上", "3年以上", "三年以上")) and not _contains_any(doc.raw_text, ("新增", "新建"))

    quantities = {
        key: value
        for key, value in {
            "申报金额": declared,
            "购置金额": purchase,
            "去年核定金额": last_year,
        }.items()
        if value is not None
    }

    if is_small_legacy:
        if last_year is None:
            return _manual(rule, rule.rule_name.replace("规则", ""), "识别到小额三年以上项目，但未识别到去年核定金额。", "请补充去年核定金额，按去年核定口径确认本年核定金额。", doc, evidence, quantities)
        return _pass(rule, rule.rule_name.replace("规则", ""), "识别到小额三年以上项目，已提取去年核定金额，可按去年核定口径复核。", quantities, evidence)

    if declared is None or purchase is None:
        return _manual(
            rule,
            rule.rule_name.replace("规则", ""),
            "未同时识别到申报金额和购置金额，无法按费率自动核价。",
            "请在明细中明确申报金额、购置金额、项目年限及是否新增。",
            doc,
            evidence,
            quantities,
        )

    threshold_amount = purchase * compare_rate
    approved_amount = declared if threshold_amount > declared else purchase * approved_rate
    quantities.update({"费率比较金额": threshold_amount, "参考核定金额": approved_amount})
    if declared <= approved_amount + 0.05:
        return _pass(rule, rule.rule_name.replace("规则", ""), "申报金额未超过按规则测算的参考核定金额。", quantities, evidence)

    return _finding(
        rule,
        rule.rule_name.replace("规则", ""),
        "中",
        "核价疑似偏高",
        f"申报金额 {declared:g} 万元，高于按规则测算的参考核定金额 {approved_amount:g} 万元。",
        "请核对购置金额、费率和项目年限；如确需高于参考核定金额，请补充依据。",
        quantities,
        evidence,
    )


def _evaluate_security_assessment_fee(
    doc: MaintenanceDocument,
    rule: MaintenanceRule,
    level: str,
    low_fee: float,
    high_fee: float,
) -> ResourceRuleFinding:
    level_keywords = (f"等保{level}", f"{level}等保", f"{level}等级保护", f"等级保护{level}")
    if not _contains_any(doc.raw_text, level_keywords):
        return _pass(rule, f"等保{level}测评费", f"未识别到等保{level}测评费场景。", {}, [])

    evidence = _evidence(doc, level_keywords + ("测评费", "数据局", "立项", "核定", "系统金额"))
    if _contains_any(doc.raw_text, ("数据局已立项", "已立项审核", "数据局核定")):
        return _manual(rule, f"等保{level}测评费", "识别到数据局已立项或核定口径，应按数据局核定金额复核。", "请录入或核对数据局核定金额。", doc, evidence)

    system_amount = _first_labeled_amount_near(doc.lines, ("系统金额", "系统", "申报金额", "建设金额", "项目金额", "购置金额"))
    fee_amount = _first_labeled_amount_near(doc.lines, ("测评费", "等保测评", "等级保护测评"))
    quantities = {key: value for key, value in {"系统金额": system_amount, "测评费": fee_amount}.items() if value is not None}
    if system_amount is None or fee_amount is None:
        return _manual(rule, f"等保{level}测评费", "未同时识别到系统金额和等保测评费金额，无法自动核价。", "请明确系统金额和等保测评费金额。", doc, evidence, quantities)

    expected = high_fee if system_amount >= 100 else low_fee if system_amount >= 10 else 0.0
    quantities["规则核定金额"] = expected
    if expected <= 0:
        return _manual(rule, f"等保{level}测评费", "系统金额低于 10 万元，规则表未给出明确测评费标准。", "请人工确认该系统是否应申报等保测评费。", doc, evidence, quantities)
    if _approx_equal(fee_amount, expected, 0.05):
        return _pass(rule, f"等保{level}测评费", f"测评费与规则核定金额 {expected:g} 万元一致。", quantities, evidence)
    return _finding(
        rule,
        f"等保{level}测评费",
        "中",
        "测评费不一致",
        f"系统金额 {system_amount:g} 万元，对应等保{level}测评费应为 {expected:g} 万元，识别金额为 {fee_amount:g} 万元。",
        "请按系统金额区间调整等保测评费，或补充数据局核定依据。",
        quantities,
        evidence,
    )


def _evaluate_history_price_compare(doc: MaintenanceDocument, rule: MaintenanceRule) -> ResourceRuleFinding:
    current = _first_labeled_amount_near(doc.lines, ("本年度核定", "今年核定", "本年核定", "核定价格"))
    previous = _first_labeled_amount_near(doc.lines, ("去年核定", "上年核定", "历史价格", "去年价格"))
    evidence = _evidence(doc, ("本年度核定", "今年核定", "去年核定", "上年核定", "历史价格", "核定价格"))
    quantities = {key: value for key, value in {"本年度核定价格": current, "去年核定价格": previous}.items() if value is not None}
    if current is None or previous is None:
        return _manual(rule, "历史价格比较", "未同时识别到本年度核定价格和去年核定价格，无法完成历史价格比较。", "请提供本年度核定价格和去年核定价格。", doc, evidence, quantities)
    if current <= previous + 0.05:
        return _pass(rule, "历史价格比较", "本年度核定价格未高于去年核定价格。", quantities, evidence)
    return _finding(
        rule,
        "历史价格比较",
        "中",
        "高于去年核定价格",
        f"本年度核定价格 {current:g} 万元，高于去年核定价格 {previous:g} 万元。",
        "请核对涨价原因；如确需上调，请补充价格依据和审批说明。",
        quantities,
        evidence,
    )


def _evaluate_new_product_price_reference(
    doc: MaintenanceDocument,
    rule: MaintenanceRule,
    db: Session | None = None,
    content: bytes | None = None,
    filename: str | None = None,
) -> ResourceRuleFinding:
    if not (_contains_any(doc.raw_text, ("新增", "新购", "新增项")) and _contains_any(doc.raw_text, ("硬件", "产品软件", "成品软件", "软件产品"))):
        return _pass(rule, "新增软硬件产品核价", "未识别到新增软硬件产品场景。", {}, [])

    price_library_result = _evaluate_new_product_with_price_library(rule, db, content, filename)
    if price_library_result is not None:
        return price_library_result

    evidence = _evidence(doc, ("新增", "新购", "产品库", "参考价", "上海市产品库", "松江区产品库"))
    missing_reference = _contains_any(doc.raw_text, ("未提供产品库", "未提供参考价", "无产品库", "无参考价", "未匹配", "缺少参考价", "未查询"))
    if missing_reference:
        return _finding(
            rule,
            "新增软硬件产品核价",
            "中",
            "缺少参考价依据",
            "识别到新增软硬件产品，且文本中出现未提供或缺少产品库参考价的表述。",
            "请补充市级或区级产品库参考价，并据此给出参考核定价格。",
            {"新增产品依据数": float(len(evidence))},
            evidence,
        )
    has_reference = _contains_any(doc.raw_text, ("产品库", "参考价", "上海市产品库", "松江区产品库", "政采"))
    if has_reference:
        return _pass(rule, "新增软硬件产品核价", "识别到新增产品及产品库/参考价相关依据。", {}, evidence)
    return _finding(
        rule,
        "新增软硬件产品核价",
        "中",
        "缺少参考价依据",
        "识别到新增软硬件产品，但未识别到上海市产品库、松江区产品库或参考价依据。",
        "请补充市级或区级产品库参考价，并据此给出参考核定价格。",
        {"新增产品依据数": float(len(evidence))},
        evidence,
    )


def _evaluate_new_product_with_price_library(
    rule: MaintenanceRule,
    db: Session | None,
    content: bytes | None,
    filename: str | None,
) -> ResourceRuleFinding | None:
    if db is None or content is None or not filename:
        return None

    try:
        parsed = parse_price_document(content, filename)
    except Exception:
        return None

    candidates = [
        item
        for item in parsed.items
        if _is_detail_item(item) and _is_product(item) and not _is_software(item)
    ]
    if not candidates:
        return None

    result = _evaluate_products(db, parsed.items)
    quantities = {
        "新增软硬件产品数": float(len(candidates)),
        "价格参考库记录数": float(result.get("benchmark_count") or 0),
        "价格库匹配数": float(result.get("matched_count") or 0),
    }
    evidence = [_trim(item.evidence or item.item_name) for item in candidates if item.evidence or item.item_name]

    notices = [notice for notice in result.get("notices") or [] if isinstance(notice, dict)]
    if result.get("library_status") == "empty":
        notice = notices[0] if notices else {}
        return _finding(
            rule,
            "新增软硬件产品核价",
            "需人工确认",
            "价格参考库为空",
            str(notice.get("message") or "已识别到新增软硬件产品，但当前价格参考库为空，无法执行产品库价格比对。"),
            str(notice.get("action") or "请导入上海市产品库、松江区产品库或政府采购/询价参考价后重新检测。"),
            quantities,
            evidence,
        )

    issues = [issue for issue in result.get("issues") or [] if isinstance(issue, dict)]
    if issues:
        quantities["价格异常或待确认项数"] = float(len(issues))
        severity = _highest_severity(str(issue.get("risk_level") or issue.get("severity") or "需人工确认") for issue in issues)
        issue_reasons = [
            f"{issue.get('item_name') or '产品'}：{issue.get('reason') or issue.get('message') or '需复核'}"
            for issue in issues[:5]
        ]
        suggestions = [str(issue.get("suggestion") or "").strip() for issue in issues if str(issue.get("suggestion") or "").strip()]
        issue_evidence = [
            _trim(str(issue.get("evidence") or issue.get("source_section") or issue.get("item_name") or ""))
            for issue in issues
            if issue.get("evidence") or issue.get("source_section") or issue.get("item_name")
        ]
        return _finding(
            rule,
            "新增软硬件产品核价",
            severity,
            "价格参考库比对异常",
            f"按软硬件产品价格参考库比对发现 {len(issues)} 项需复核：" + "；".join(issue_reasons),
            "；".join(dict.fromkeys(suggestions)) or "请核对品牌、型号、规格、申报单价和产品库参考价，必要时补充政府采购/询价依据。",
            quantities,
            issue_evidence or evidence,
        )

    return _pass(
        rule,
        "新增软硬件产品核价",
        "已按软硬件产品价格参考库完成新增产品匹配，未发现申报单价高于参考价的异常。",
        quantities,
        evidence,
    )


def _evaluate_exclusion_terms(
    doc: MaintenanceDocument,
    rule: MaintenanceRule,
    resource_name: str,
    terms: tuple[str, ...],
    reason: str,
) -> ResourceRuleFinding:
    evidence = _evidence(doc, terms, limit=10)
    if not evidence:
        return _pass(rule, resource_name, "未识别到规则列明的核定为 0 的申报项。", {}, [])
    return _finding(
        rule,
        resource_name,
        "高",
        "应核定为0",
        reason,
        "请将上述申报项核定价格调整为 0，或补充其属于运维范围的正式依据。",
        {"命中申报项数": float(len(evidence))},
        evidence,
    )


def _finding(
    rule: MaintenanceRule,
    resource_name: str,
    severity: str,
    result_label: str,
    reason: str,
    suggestion: str | None,
    source_quantities: dict[str, float] | None,
    evidence_examples: list[str] | None,
) -> ResourceRuleFinding:
    return ResourceRuleFinding(
        rule_code=rule.rule_id,
        rule_name=rule.rule_name,
        resource_name=resource_name,
        severity=severity,
        result_label=result_label,
        reason=reason,
        suggestion=suggestion,
        source_quantities=source_quantities or {},
        row_indexes=[],
        evidence_examples=(evidence_examples or [])[:8],
        source_section="资源申请合理性 / 运维项目规则",
    )


def _pass(
    rule: MaintenanceRule,
    resource_name: str,
    reason: str,
    source_quantities: dict[str, float] | None,
    evidence_examples: list[str] | None,
) -> ResourceRuleFinding:
    return _finding(rule, resource_name, "通过", "通过", reason, None, source_quantities, evidence_examples)


def _manual(
    rule: MaintenanceRule,
    resource_name: str,
    reason: str,
    suggestion: str,
    doc: MaintenanceDocument,
    evidence_examples: list[str] | None = None,
    source_quantities: dict[str, float] | None = None,
) -> ResourceRuleFinding:
    evidence = evidence_examples if evidence_examples is not None else _evidence(doc, tuple(rule.rule_name.split()))
    return _finding(rule, resource_name, "需人工确认", "资料不足", reason, suggestion, source_quantities, evidence)


def _rule_selected(rule: MaintenanceRule, selected: set[str]) -> bool:
    if not selected:
        return True
    if "MAINT_OPS_ALL" in selected:
        return True
    values = {
        rule.rule_id,
        str(rule.rule_number),
        str(rule.excel_row),
        rule.rule_name,
        f"名称：{rule.rule_name}",
    }
    return bool(values & selected)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _contains_any(value: str, keywords: Iterable[str]) -> bool:
    normalized = value.lower()
    return any(str(keyword).lower() in normalized for keyword in keywords if str(keyword))


def _fuzzy_2gram_keyword_hits(value: str, keywords: Iterable[str], threshold: float) -> set[str]:
    """Return keywords whose character 2-grams are mostly covered by value."""

    segment_grams = [_char_ngrams(segment) for segment in _ngram_segments(value)]
    segment_grams = [grams for grams in segment_grams if grams]
    if not segment_grams:
        return set()

    hits: set[str] = set()
    for keyword in keywords:
        keyword_text = str(keyword).strip()
        keyword_grams = _char_ngrams(keyword_text)
        if not keyword_grams:
            continue
        if any(len(grams & keyword_grams) / len(keyword_grams) >= threshold for grams in segment_grams):
            hits.add(keyword_text)
    return hits


def _ngram_segments(value: str, max_segment_length: int = 160) -> list[str]:
    raw_segments = re.split(r"[\r\n。；;，,、.!?！？]+", str(value))
    segments: list[str] = []
    for segment in raw_segments:
        normalized = segment.strip()
        if not normalized:
            continue
        if len(normalized) <= max_segment_length:
            segments.append(normalized)
            continue
        for index in range(0, len(normalized), max_segment_length):
            chunk = normalized[index : index + max_segment_length].strip()
            if chunk:
                segments.append(chunk)
    return segments


def _char_ngrams(value: str, n: int = 2) -> set[str]:
    normalized = re.sub(r"[^0-9a-zA-Z\u4e00-\u9fff]+", "", str(value).lower())
    if not normalized:
        return set()
    if len(normalized) <= n:
        return {normalized}
    return {normalized[index : index + n] for index in range(len(normalized) - n + 1)}


def _evidence(doc: MaintenanceDocument, keywords: Iterable[str], limit: int = 6) -> list[str]:
    output: list[str] = []
    seen: set[str] = set()
    for line in doc.lines:
        if not _contains_any(line, keywords):
            continue
        text = _trim(line)
        if text in seen:
            continue
        seen.add(text)
        output.append(text)
        if len(output) >= limit:
            break
    return output


def _trim(value: str, limit: int = 240) -> str:
    text = re.sub(r"\s+", " ", value).strip()
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _format_quantities(values: dict[str, float]) -> str:
    return "；".join(f"{key}={value:g}万元" for key, value in values.items())


def _largest_amount(line: str) -> float | None:
    amounts = _extract_amounts(line)
    return max(amounts) if amounts else None


def _extract_amounts(line: str) -> list[float]:
    values: list[float] = []
    for match in AMOUNT_PATTERN.finditer(line):
        value = _amount_from_match(line, match)
        if value is not None:
            values.append(value)
    return values


def _amount_from_match(line: str, match: re.Match[str]) -> float | None:
    raw = match.group("number")
    unit = match.group("unit") or ""
    try:
        number = float(raw.replace(",", ""))
    except ValueError:
        return None
    if number <= 0:
        return None
    local_context = line[max(0, match.start() - 8) : min(len(line), match.end() + 8)]
    if unit in {"万元", "万"}:
        return number
    if unit == "元":
        return number / 10000
    if _contains_any(local_context, ("金额", "价格", "报价", "费用", "申报", "核定", "预算", "总计", "合计", "购置")):
        return number
    return None


def _amounts_near_keywords(lines: list[str], keywords: Iterable[str]) -> list[float]:
    amounts: list[float] = []
    for line in lines:
        if _contains_any(line, keywords):
            amounts.extend(_extract_amounts(line))
    return amounts


def _first_amount_near(lines: list[str], keywords: Iterable[str]) -> float | None:
    amounts = _amounts_near_keywords(lines, keywords)
    return amounts[0] if amounts else None


def _first_labeled_amount_near(lines: list[str], keywords: Iterable[str]) -> float | None:
    labels = [str(keyword).strip() for keyword in keywords if str(keyword).strip()]
    for line in lines:
        lower_line = line.lower()
        label_ends: list[int] = []
        for label in labels:
            start = lower_line.find(label.lower())
            while start >= 0:
                label_ends.append(start + len(label))
                start = lower_line.find(label.lower(), start + len(label))
        if not label_ends:
            continue
        amount_matches = list(AMOUNT_PATTERN.finditer(line))
        for label_end in sorted(label_ends):
            for match in amount_matches:
                if match.start() < label_end:
                    continue
                if match.start() - label_end > 32:
                    break
                value = _amount_from_match(line, match)
                if value is not None:
                    return value
    return _first_amount_near(lines, labels)


def _numbers_near_keywords(lines: list[str], keywords: Iterable[str]) -> list[float]:
    values: list[float] = []
    pattern = re.compile(r"(?<![\d.])(\d+(?:\.\d+)?)(?![\d.])")
    for line in lines:
        if not _contains_any(line, keywords):
            continue
        for match in pattern.finditer(line):
            try:
                value = float(match.group(1))
            except ValueError:
                continue
            if 0 <= value <= 10000:
                values.append(value)
    return values


def _has_full_year_coverage(text: str) -> bool:
    if "全年" in text or "全年度" in text:
        return True
    compact = re.sub(r"\s+", "", text)
    return bool(re.search(r"(1|01)月.*(12|十二)月", compact) or re.search(r"20\d{2}[-年/]0?1[-月/].*20\d{2}[-年/]12[-月/]", compact))


def _approx_equal(left: float, right: float, tolerance: float = 0.05) -> bool:
    return abs(left - right) <= tolerance


def _highest_severity(values: Iterable[str]) -> str:
    order = {"高": 4, "中": 3, "低": 2, "需人工确认": 1, "通过": 0}
    best = "需人工确认"
    best_score = order[best]
    for value in values:
        score = order.get(value, 1)
        if score > best_score:
            best = value
            best_score = score
    return best
