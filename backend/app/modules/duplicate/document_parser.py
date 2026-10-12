from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
import re
from typing import List, Optional
from docx import Document
from docx.oxml.ns import qn


@dataclass(frozen=True)
class DocumentSection:
    title: str
    content: str
    start_line: int = 0
    end_line: int = 0
    page_no: Optional[int] = None
    heading_path: tuple[str, ...] = ()


@dataclass(frozen=True)
class DocumentSourceBlock:
    text: str
    paragraph: int
    line: int
    section: str = ""
    heading_path: tuple[str, ...] = ()
    heading_level: int = 0
    table_index: int = 0
    table_row: int = 0
    cells: tuple[str, ...] = ()


@dataclass(frozen=True)
class DocumentContent:
    format: str  # "docx"
    raw_text: str
    sections: List[DocumentSection]
    filename: str
    source_blocks: List[DocumentSourceBlock] = field(default_factory=list)


def parse_document(content: bytes, filename: str) -> DocumentContent:
    """根据文件扩展名选择合适的解析器"""
    if filename.lower().endswith(".docx"):
        return _parse_docx(content, filename)
    else:
        raise ValueError(f"不支持的文件格式: {filename}")


def _heading_level(paragraph, text: str) -> int:
    style_name = str(getattr(paragraph.style, "name", "") or "").lower()
    if "heading" in style_name or "标题" in style_name:
        match = re.search(r"(\d+)", style_name)
        return int(match.group(1)) if match else 1
    if len(text) <= 100 and not any(mark in text for mark in "。；，"):
        if re.match(r"^第[一二三四五六七八九十百\d]+[章节]\s*\S", text):
            return 1
        match = re.match(r"^(\d+(?:\.\d+){1,7})\.?\s*(?=[\u4e00-\u9fffA-Za-z“])", text)
        if match:
            if int(match.group(1).split(".")[0]) > 99 or re.match(r"^(万元|元|人月|年|月|日|%)", text[match.end():]):
                return 0
            return min(9, match.group(1).count(".") + 1)
    # Imported Word reports often use custom styles with an outline level.
    elements = [paragraph._p.pPr]
    style = paragraph.style
    for _ in range(4):
        if style is None:
            break
        elements.append(style.element.find(qn("w:pPr")))
        style = style.base_style
    for properties in elements:
        outline = properties.find(qn("w:outlineLvl")) if properties is not None else None
        if outline is not None:
            value = int(outline.get(qn("w:val"), "9"))
            return value + 1 if value < 9 else 0
    return 0


def _parse_docx(content: bytes, filename: str) -> DocumentContent:
    """解析 Word 文档"""
    try:
        doc = Document(BytesIO(content))
    except Exception as e:
        raise ValueError(f"Word 文档解析失败: {e}") from e

    sections: List[DocumentSection] = []
    # One body-order pass: tables keep their owning section and merged cells
    # are resolved without python-docx's repeated full-table row.cells scan.
    blocks = _extract_docx_source_blocks(doc)
    current_title, current_path, current_lines = "", (), []
    start_line = 0
    for block in blocks:
        if block.heading_level:
            if current_lines:
                sections.append(DocumentSection(current_title, "\n".join(current_lines),
                                                start_line, block.line - 1, heading_path=current_path))
            current_title, current_path, current_lines = block.text, block.heading_path, []
            start_line = block.line
        else:
            current_lines.append(" | ".join(block.cells) if block.cells else block.text)
    if current_lines:
        sections.append(DocumentSection(current_title, "\n".join(current_lines),
                                        start_line, blocks[-1].line, heading_path=current_path))
    raw_text = "\n".join(" | ".join(b.cells) if b.cells else b.text for b in blocks)

    return DocumentContent(
        format="docx",
        raw_text=raw_text,
        sections=sections,
        filename=filename,
        source_blocks=blocks,
    )


def _extract_docx_source_blocks(doc: Document) -> list[DocumentSourceBlock]:
    """Build locations with the same paragraph/line numbering as the preview."""
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph

    blocks: list[DocumentSourceBlock] = []
    paragraph_index = 0
    line_index = 0
    current_section = ""
    heading_stack: list[tuple[int, str]] = []
    table_index = 0

    def register(text: str) -> tuple[int, int]:
        nonlocal paragraph_index, line_index
        paragraph_index += 1
        if text.strip():
            line_index += 1
        return paragraph_index, line_index if text.strip() else 0

    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, doc)
            text = paragraph.text.strip()
            location = register(text)
            level = _heading_level(paragraph, text) if text else 0
            if level:
                current_section = text
                while heading_stack and heading_stack[-1][0] >= level:
                    heading_stack.pop()
                heading_stack.append((level, text))
            if text:
                blocks.append(DocumentSourceBlock(text, location[0], location[1], current_section,
                                                  tuple(t for _, t in heading_stack), level))
        elif isinstance(child, CT_Tbl):
            table_index += 1
            table = Table(child, doc)
            merged: dict[int, str] = {}
            for row_index, row_xml in enumerate(child.tr_lst, 1):
                row_texts: list[str] = []
                row_locations: list[tuple[int, int]] = []
                cells: list[str] = []
                column = 0
                for cell_xml in row_xml.tc_lst:
                    cell = _Cell(cell_xml, table)
                    cell_parts: list[str] = []
                    for paragraph in cell.paragraphs:
                        text = paragraph.text.strip()
                        row_locations.append(register(text))
                        if text:
                            cell_parts.append(text)
                    if cell_parts:
                        row_texts.append(" ".join(cell_parts))
                    value = " ".join(cell_parts)
                    props = cell_xml.tcPr
                    vmerge = props.vMerge if props is not None else None
                    if vmerge is not None and vmerge.val == "continue":
                        value = merged.get(column, value)
                    span = cell_xml.grid_span
                    for offset in range(span):
                        cells.append(value)
                        if vmerge is not None:
                            merged[column + offset] = value
                        else:
                            merged.pop(column + offset, None)
                    column += span
                if row_texts:
                    first = next((item for item in row_locations if item[1] > 0), row_locations[0])
                    blocks.append(DocumentSourceBlock(" ".join(row_texts), first[0], first[1], current_section,
                                                      tuple(t for _, t in heading_stack), 0,
                                                      table_index, row_index, tuple(cells)))
    return blocks


def _extract_docx_table_text(doc: Document) -> list[str]:
    rows: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            values = [
                cell.text.strip().replace("\n", " ")
                for cell in row.cells
                if cell.text and cell.text.strip()
            ]
            if values:
                rows.append(" | ".join(values))
    return rows
