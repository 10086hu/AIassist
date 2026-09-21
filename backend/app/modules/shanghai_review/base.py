# -*- coding: utf-8 -*-
from __future__ import annotations

import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional


class ErrorCode:
    ICC_TOTAL_MISMATCH = "ICC-001"
    ICC_SECTION_MISSING = "ICC-002"
    ICC_FEE_CALC_ERROR = "ICC-003"
    ICC_SECOND_CLASS_CAP_EXCEEDED = "ICC-004"
    ICC_INTEGRATION_RATE_WRONG = "ICC-005"
    IDC_MISSING_IN_4_2 = "IDC-003"
    IDC_MISSING_IN_2_6 = "IDC-004"
    SCC_LEVEL_MISMATCH = "SCC-001"
    SCC_RISK_NOT_COVERED = "SCC-002"
    SCC_SECTION_MISSING = "SCC-003"
    MBC_MODULE_MISSING_IN_BUDGET = "MBC-001"
    MBC_BUDGET_ITEM_NO_MODULE = "MBC-002"
    MBC_NUMBER_MISALIGN = "MBC-003"
    PTC_GENERAL_BENEFIT_INSUFFICIENT = "PTC-001"
    PTC_AI_EFFECTIVENESS_INSUFFICIENT = "PTC-003"
    PTC_IRRELEVANT_INDICATOR = "PTC-004"
    DBC_CONTENT_MISSING = "DBC-001"
    DBC_DATA_SERVICE_MISSING = "DBC-002"
    DBC_UPLINK_FEE_VIOLATION = "DBC-003"


@dataclass(frozen=True)
class ValidationError:
    code: str
    message: str
    section: str
    severity: str = "warning"
    suggestion: Optional[str] = None


@dataclass(frozen=True)
class ValidationResult:
    is_valid: bool
    errors: List[ValidationError] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    metadata: Dict[str, Any] = field(default_factory=dict)

    @classmethod
    def from_errors(cls, errors, **meta):
        risks = [e for e in errors if e.severity == "risk"]
        warns = [e.message for e in errors if e.severity == "warning"]
        return cls(is_valid=len(risks) == 0, errors=errors, warnings=warns, metadata=meta)


class BaseValidator(ABC):
    validator_id = "base"
    validator_name = "base"

    @abstractmethod
    def validate(self, document):
        ...

    @staticmethod
    def parse_document(file_bytes, filename):
        import io
        hr = re.compile(r"^\s*(\d+(?:\.\d+){1,2})[\.\s．]+(.+)", re.UNICODE | re.MULTILINE)
        imp = ["2.6", "2.7", "4.2", "4.7", "5.1", "6.1", "6.4", "6.6", "8.2"]
        ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
        nl = chr(10)
        tab = chr(9)
        ft = ""
        if ext in ("docx", "doc"):
            import docx as _d
            doc_obj = _d.Document(io.BytesIO(file_bytes))
            parts = [p.text for p in doc_obj.paragraphs if p.text.strip()]
            for ti, table in enumerate(doc_obj.tables):
                parts.append("===TABLE_%d===" % ti)
                for row in table.rows:
                    parts.append(tab.join(cell.text.strip() for cell in row.cells))
                parts.append("===/TABLE_%d===" % ti)
            ft = nl.join(parts)
        else:
            ft = file_bytes.decode("utf-8", errors="replace")
        doc = {"_full_text": ft}
        matches = list(hr.finditer(ft))
        if matches:
            for j, m in enumerate(matches):
                num = m.group(1)
                start = m.start()
                end = matches[j + 1].start() if j + 1 < len(matches) else len(ft)
                doc[num] = ft[start:end]
        for key in imp:
            if key not in doc:
                doc[key] = ft
        return doc

    @staticmethod
    def register_api(app, prefix="/api/evaluate", tags=None, validators=None):
        import traceback as _tb
        from fastapi import APIRouter, File, Form, HTTPException, UploadFile
        router = APIRouter()
        vals = validators or []
        tgs = tags or ["shanghai"]

        @router.post("/shanghai")
        async def _shanghai_review(file: UploadFile = File(...), pname: str = Form(default="test")):
            fname = file.filename or ""
            content = await file.read()
            if not content:
                raise HTTPException(status_code=400, detail="empty file")
            extt = fname.lower().rsplit(".", 1)[-1] if "." in fname else ""
            if extt not in ("docx", "doc", "txt"):
                raise HTTPException(status_code=400, detail="unsupported: ." + extt)
            try:
                document = BaseValidator.parse_document(content, fname)
            except Exception as exc:
                raise HTTPException(status_code=500, detail="parse error: " + str(exc))
            results = []
            tr = 0
            tw = 0
            ap = True
            for v in vals:
                try:
                    r = v.validate(document)
                except Exception as exc:
                    r = ValidationResult.from_errors([ValidationError("EXCEPTION", str(exc), "", "risk")])
                risks = sum(1 for e in r.errors if e.severity == "risk")
                warns = sum(1 for e in r.errors if e.severity == "warning")
                tr += risks
                tw += warns
                if not r.is_valid:
                    ap = False
                results.append({
                    "validator_id": v.validator_id, "validator_name": v.validator_name,
                    "is_valid": r.is_valid,
                    "errors": [{"code": e.code, "message": e.message, "section": e.section,
                                 "severity": e.severity, "suggestion": e.suggestion} for e in r.errors],
                    "warnings": r.warnings,
                })
            sf = sorted(k for k in document if not k.startswith("_") and len(k) < 30)
            return {"project_name": pname, "filename": fname, "sections_found": sf,
                    "all_pass": ap, "total_risks": tr, "total_warnings": tw, "results": results}

        app.include_router(router, prefix=prefix, tags=tgs)
