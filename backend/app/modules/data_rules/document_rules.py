from __future__ import annotations


import re


import unicodedata


from dataclasses import asdict, dataclass, field, replace


from io import BytesIO


from typing import Any, Iterable


from app.modules.docx_grid import table_grid


from app.modules.evidence_text import positive_mentions


from app.modules.duplicate.document_parser import DocumentContent, parse_document


from app.modules.data_rules.data_governance_service import (
    POLICY,
    assess_data_governance_service_applicability,
    validate_data_governance_service_design,
)


from app.modules.shanghai_review import BaseValidator, DataReportingValidator, ValidationError


@dataclass(frozen=True)
class RuleIssue:
    message: str
    evidence: str | None = None
    section: str | None = None
    severity: str = 'risk'


@dataclass(frozen=True)
class RuleResult:
    rule_excel_row: int
    rule_name: str
    rule_category: str
    rule_description: str
    judgement_condition: str
    passed: bool
    status: str
    severity: str
    summary: str
    metrics: dict[str, Any]
    issues: list[RuleIssue]
    suggestions: list[str]
    evidence_locations: list[dict] = field(default_factory=list)


LEGACY_RULE_IDS = {
    18: "DATA_REASON_001",
    21: "DATA_REASON_002",
}


DATA_REPORTING_RULES = ({
        "rule_id": "DATA_REASON_024",
        "excel_row": 24,
        "rule_name": "数据治理服务内容设计合规性审查规则",
        "rule_category": "内容合规性审查规则",
        "rule_description": "涉及数据治理服务的项目，数据服务事项内容参照《市级数字化项目数据治理服务配置指引（试行）》开展编制。",
        "judgement_condition": "第6.3节数据治理内容附表应符合指引的4类6项服务范围、计量单位和负面清单要求。",
        "error_prefixes": ("DGS-",),
        "default_suggestion": "请按指引附表补充或调整6.3数据治理服务内容，确保服务事项属于4类6项，计量单位正确，且不包含负面清单事项。",
    },)
NUMBER_TEXT = r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?"


MONEY_PATTERN = re.compile(
    rf"(?<![A-Za-z0-9_.])(?P<currency>人民币|RMB|¥|￥)?\s*"
    rf"(?P<number>{NUMBER_TEXT})\s*"
    r"(?:(?:[（(]\s*)?(?P<unit>万元|万\s*元|元|万)(?:\s*[）)])?)?"
    r"(?!\s*(?:\+|个|台|套|块|条|项|人|次|年|月|日|天|小时|页|张|路|颗|份|类|核|票|字节|"
    r"GB|MB|TB|KB|G|M|GHz|MHz|%|％|[A-Za-z]*B))",
    re.IGNORECASE,
)


NUMBER_PATTERN = re.compile(rf"(?<![A-Za-z0-9_.])(?P<number>{NUMBER_TEXT})(?![A-Za-z0-9_.])")


NON_MONEY_UNIT_RE = re.compile(
    r"^\s*(?:个|台|套|块|条|项|人|次|年|月|日|天|小时|页|张|路|颗|份|类|核|"
    r"GB|MB|TB|KB|G|M|GHz|MHz|%|％)",
    re.IGNORECASE,
)


FINANCIAL_CONTEXT_KEYWORDS = (
    "投资",
    "预算",
    "概算",
    "估算",
    "金额",
    "费用",
    "经费",
    "资金",
    "成本",
    "造价",
    "单价",
    "总价",
    "报价",
    "购置费",
    "开发费",
    "服务费",
    "建设费",
    "总计",
    "合计",
    "小计",
)


AMOUNT_COLUMN_KEYWORDS = (
    "金额",
    "投资",
    "预算",
    "概算",
    "估算",
    "费用",
    "经费",
    "资金",
    "成本",
    "造价",
    "单价",
    "总价",
    "合价",
    "报价",
    "小计",
    "合计",
    "总计",
    "万元",
    "元",
)


PROJECT_TOTAL_KEYWORDS = (
    "项目总投资",
    "项目投资总额",
    "项目投资总预算",
    "总投资",
    "投资总额",
    "预算总额",
    "总预算",
)


PROJECT_TOTAL_ALIASES = (
    *PROJECT_TOTAL_KEYWORDS,
    "总投资额",
    "项目总预算",
    "工程总投资",
)


DETAIL_TOTAL_KEYWORDS = ("总计", "合计")


NON_AMOUNT_COLUMN_KEYWORDS = (
    "数量",
    "单位",
    "规格",
    "配置",
    "容量",
    "内存",
    "硬盘",
    "CPU",
    "工作量",
    "人月",
    "服务期",
    "建设周期",
    "周期",
    "月份",
    "条数",
    "页数",
)


INTELLIGENT_KEYWORDS = (
    "人工智能",
    "大模型",
    "智能化应用",
    "智能识别",
    "智能分析",
    "智能搜索",
    "智能推荐",
    "智能问答",
    "知识图谱",
    "机器学习",
    "深度学习",
    "算法模型",
    "OCR",
    "语音识别",
)


BUDGET_CONTEXT_ALIASES = (
    *FINANCIAL_CONTEXT_KEYWORDS,
    "申报金额",
    "申报总额",
    "总金额",
    "金额小计",
    "投资估算额",
    "概算金额",
)


BUDGET_AMOUNT_COLUMN_ALIASES = (
    *AMOUNT_COLUMN_KEYWORDS,
    "申报金额",
    "申报总额",
    "总金额",
    "合计金额",
)


DETAIL_TOTAL_ALIASES = (
    *DETAIL_TOTAL_KEYWORDS,
    "汇总",
    "共计",
    "总额",
    "总和",
    "合计金额",
    "总计金额",
)


INDICATOR_SCOPE_ALIASES = {
    "common": ("通用指标", "通用", "共性指标", "通用绩效指标"),
    "business": ("业务指标", "业务", "行业指标", "业务绩效指标"),
}


def _evaluate_budget_consistency_without_locations(
    document: DocumentContent,
    content: bytes | None = None,
    filename: str | None = None,
) -> RuleResult:
    lines = _meaningful_lines(document.raw_text)
    amount_lines = [
        (index, line, _extract_amounts_from_line(lines, index, document.raw_text))
        for index, line in enumerate(lines)
    ]
    amount_lines = [
        (index, line, amounts)
        for index, line, amounts in amount_lines
        if amounts
    ]

    project_total_candidates = []
    for _index, line, amounts in amount_lines:
        if not _contains_project_total_keyword(line):
            continue
        anchored_amounts = _extract_project_total_amounts(line, amounts)
        if not anchored_amounts:
            continue
        project_total_candidates.append(
            {"line": line, "amount_wan": anchored_amounts[0]}
        )
    project_totals = _distinct_amounts(
        item["amount_wan"] for item in project_total_candidates
    )

    detail_total_candidates = [
        {
            "line": line,
            "amount_wan": _pick_representative_amount(amounts),
            "source": "文本行",
            "kind": "text_total",
        }
        for index, line, amounts in amount_lines
        if _is_detail_total_context(lines, index, line)
        and not _contains_project_total_keyword(line)
        and not _is_non_project_budget_line(line)
    ]

    arithmetic_issues = _find_total_row_arithmetic_issues(lines, amount_lines)
    text_arithmetic_rows_checked = _count_total_row_arithmetic_checks(lines, amount_lines)
    table_budget = (
        _extract_docx_budget_table_checks(content, filename)
        if content and (filename or document.filename).lower().endswith(".docx")
        else _empty_budget_table_checks()
    )
    if table_budget["detail_total_candidates"]:
        detail_total_candidates = _distinct_budget_candidates(
            [*detail_total_candidates, *table_budget["detail_total_candidates"]]
        )
    arithmetic_issues.extend(table_budget["arithmetic_issues"])
    issues: list[RuleIssue] = []
    suggestions: list[str] = []

    if content and (filename or document.filename).lower().endswith('.docx'):
        from .budget_components import find_component_differences
        issues.extend(RuleIssue(**x) for x in find_component_differences(content))

    # Name-level partial matches are diagnostic evidence, not arithmetic errors.
    # Keep them in budget_table_relations without failing an otherwise reconciled budget.
    if len(project_totals) > 1:
        issue_text = "、".join(_format_amount(item) for item in project_totals)
        issues.append(
            RuleIssue(
                message=f"识别到多个不一致的项目总投资金额：{issue_text}。",
                evidence=" | ".join(item["line"] for item in project_total_candidates[:5]),
                section="项目总投资相关表述",
            )
        )
        suggestions.append("请统一全文中的项目总投资、投资概况和预算编制说明金额。")

    project_detail_total_checked = False
    if project_totals and detail_total_candidates:
        project_total = project_totals[0]
        matching_detail_totals = [
            item
            for item in detail_total_candidates
            if _amounts_close(item["amount_wan"], project_total)
        ]
        component_total_candidates = [
            item
            for item in detail_total_candidates
            if _budget_candidate_is_component_total(item)
        ]
        component_total_sum = _sum_candidate_amounts(component_total_candidates)
        if matching_detail_totals:
            project_detail_total_checked = True
        elif any(
            _project_total_breakdown_matches(line, project_total, amounts)
            for _index, line, amounts in amount_lines
            if _contains_project_total_keyword(line)
        ):
            project_detail_total_checked = True
            # A prose breakdown can reconcile internally while independently
            # extracted cost tables use another scope or contain overlaps.
            # Keep the rule conservative when several detail-table totals do
            # not reconcile to the same project total.
            detail_sources = {
                str(item.get("source") or "").strip()
                for item in detail_total_candidates
                if str(item.get("source") or "").strip()
            }
            if (
                component_total_sum is not None
                and not _amounts_close(component_total_sum, project_total)
                and len(detail_sources) >= 2
            ):
                issues.append(
                    RuleIssue(
                        message=(
                            "资料不足：项目总投资的正文拆分可以闭合，但识别到的多张预算明细表合计"
                            f"为 {_format_amount(component_total_sum)}，与项目总投资 {_format_amount(project_total)} 不一致，"
                            "无法确认是否属于同一预算口径或存在重复计列。"
                        ),
                        evidence=" | ".join(item["line"] for item in detail_total_candidates[:5]),
                        section="预算明细表 / 投资估算总表",
                    )
                )
        elif component_total_sum is not None and _amounts_close(component_total_sum, project_total):
            project_detail_total_checked = True
        else:
            hard_mismatch_candidates = [
                item
                for item in detail_total_candidates
                if _budget_candidate_is_grand_total(item)
                or _amount_exceeds(item["amount_wan"], project_total)
            ]
            # Resource lists, quotations and procurement summaries can overlap.
            # Do not hard-fail by blindly adding independent detail-table totals.
            if hard_mismatch_candidates:
                project_detail_total_checked = True
                detail_preview = "、".join(
                    _format_amount(item["amount_wan"])
                    for item in hard_mismatch_candidates[:5]
                )
                issues.append(
                    RuleIssue(
                        message=(
                            "预算合计/总计金额未与项目总投资金额匹配。"
                            f"项目总投资为 {_format_amount(project_total)}，"
                            f"识别到的可比合计/总计包括：{detail_preview}。"
                        ),
                        evidence=" | ".join(item["line"] for item in hard_mismatch_candidates[:5]),
                        section="预算明细表 / 投资估算总表",
                    )
                )
                suggestions.append("请核对投资估算总表、分项明细表合计与项目投资总预算是否一致。")
            else:
                issues.append(
                    RuleIssue(
                        message="资料不足：已识别项目总投资和若干明细表合计，但无法确认二者属于同一预算口径。",
                        evidence=" | ".join(item["line"] for item in detail_total_candidates[:3]),
                        section="预算明细表 / 投资估算总表",
                    )
                )

    calculation_checks_performed = (
        (1 if len(project_totals) > 1 else 0)
        + (1 if project_detail_total_checked else 0)
        + text_arithmetic_rows_checked
        + int(table_budget.get("table_internal_checks") or 0)
        + int(table_budget.get("summary_internal_checks") or 0)
    )

    if not issues and calculation_checks_performed == 0:
        if project_total_candidates and not detail_total_candidates:
            issues.append(
                RuleIssue(
                    message=(
                        "资料不足：仅识别到项目总投资金额，未识别到可用于计算校验的预算明细表合计、"
                        "投资估算总表总计或可加总分项金额。"
                    ),
                    evidence=" | ".join(item["line"] for item in project_total_candidates[:3]),
                    section="预算明细表 / 投资估算总表",
                )
            )
        elif project_total_candidates and detail_total_candidates:
            issues.append(
                RuleIssue(
                    message=(
                        "资料不足：已识别项目总投资和预算合计，但二者无法建立可靠的同口径对应关系，"
                        "可能混有运维费用、年度费用或不同计量单位。"
                    ),
                    evidence=" | ".join(
                        [item["line"] for item in project_total_candidates[:2]]
                        + [item["line"] for item in detail_total_candidates[:2]]
                    ),
                    section="预算明细表/投资估算总表",
                )
            )
        elif detail_total_candidates and not project_total_candidates:
            issues.append(
                RuleIssue(
                    message=(
                        "资料不足：已识别到预算分项或合计金额，但未识别到项目总投资/总预算，"
                        "无法完成表间总和一致性校验。"
                    ),
                    evidence=" | ".join(item["line"] for item in detail_total_candidates[:3]),
                    section="预算明细表 / 投资估算总表",
                )
            )
        elif amount_lines:
            detail_preview = "、".join(
                line
                for _index, line, _amounts in amount_lines[:3]
            )
            issues.append(
                RuleIssue(
                    message=(
                        "资料不足：已识别到金额信息，但缺少项目总投资、合计/总计行或可加总明细，"
                        "无法判断预算计算是否正确。"
                    ),
                    evidence=detail_preview,
                    section="预算明细表 / 投资估算总表",
                )
            )
        else:
            issues.append(
                RuleIssue(
                    message="资料不足：未识别到可用于第18条计算校验的预算金额。",
                    evidence=None,
                    section="预算明细表 / 投资估算总表",
                )
            )
        suggestions.append("请补充项目总投资、预算明细表合计/总计行或可加总分项金额，便于完成第18条计算校验。")

    for issue in arithmetic_issues:
        issues.append(issue)
    if arithmetic_issues:
        suggestions.append("请复核合计/总计行内部加总关系，确保表内计算无误。")

    issues = [replace(issue, severity='warning') if issue.message.startswith(('资料不足', '舍入待核实', '单位待核实')) else issue for issue in issues]
    if issues:
        insufficient_only = all(issue.message.startswith("资料不足") for issue in issues)
        hard_fail = any(issue.severity == 'risk' for issue in issues)
        status = "资料不足" if insufficient_only else "failed" if hard_fail else "warning"
        severity = "risk" if hard_fail else "warning"
        passed = False
        summary = "第18条资料不足，无法完成预算计算校验。" if insufficient_only else "项目预算一致性存在需复核事项。"
    else:
        status = "passed"
        severity = "pass"
        passed = True
        summary = "已完成可识别预算金额关系校验，未发现项目总投资、预算合计和可识别合计行之间的不一致。"

    metrics = {
        "project_total_candidates": project_total_candidates[:10],
        "distinct_project_totals_wan": project_totals,
        "detail_total_candidates": detail_total_candidates[:10],
        "budget_tables_checked": table_budget["tables_checked"],
        "budget_table_relations": table_budget.get("table_relations", []),
        "budget_evidence_examples": table_budget["evidence_examples"][:10],
        "text_arithmetic_rows_checked": text_arithmetic_rows_checked,
        "docx_table_internal_checks": table_budget.get("table_internal_checks") or 0,
        "docx_summary_internal_checks": table_budget.get("summary_internal_checks") or 0,
        "calculation_checks_performed": calculation_checks_performed,
        "component_detail_total_sum_wan": _sum_candidate_amounts(
            [item for item in detail_total_candidates if _budget_candidate_is_component_total(item)]
        ),
    }

    return RuleResult(
        rule_excel_row=18,
        rule_name="项目预算一致性校验规则",
        rule_category="一致性校验规则",
        rule_description="项目预算数据计算无误",
        judgement_condition=(
            "各分项的投资估算明细表内部计算无错误、表间的计算总和无错误，"
            "所有分项明细表总计与项目投资总预算一致"
        ),
        passed=passed,
        status=status,
        severity=severity,
        summary=summary,
        metrics=metrics,
        issues=issues,
        suggestions=suggestions or ["保持全文项目投资金额、各分项明细表和总预算一致。"],
    )


def _evaluate_indicator_quantity_without_locations(
    document: DocumentContent,
    content: bytes | None = None,
    filename: str | None = None,
    project_name: str | None = None,
) -> RuleResult:
    applicability = _rule_21_project_applicability(document, filename=filename, project_name=project_name)
    if applicability.get("indeterminate"):
        return RuleResult(
            rule_excel_row=21,
            rule_name="项目成效考核目标数量设置合规性审查规则",
            rule_category="内容合规性审查规则",
            rule_description="指标设置数量应满足要求",
            judgement_condition=(
                "通用指标中至少设定3个效益指标，业务指标中至少设定4个产出指标或效益指标；"
                "涉及智能化应用的需明确不少于2个成效指标"
            ),
            passed=False,
            status="资料不足",
            severity="warning",
            summary="未能从报告中可靠确定项目区域或项目级别，无法判断第21条是否适用。",
            metrics={"project_scope_applicability": applicability, "indicator_scope_source": "project_scope_unknown"},
            issues=[RuleIssue(message="资料不足：缺少明确的项目区域或项目级别，不能直接套用市级项目指标数量规则。", evidence=None, section="项目基本信息")],
            suggestions=["请补充项目区域和项目级别（市级/区级/其他）后重新校验第21条。"],
        )
    if not applicability["applies"]:
        return RuleResult(
            rule_excel_row=21,
            rule_name="项目成效考核目标数量设置合规性审查规则",
            rule_category="内容合规性审查规则",
            rule_description="指标设置数量应满足要求",
            judgement_condition=(
                "通用指标中至少设定3个效益指标，业务指标中至少设定4个产出指标或效益指标；"
                "涉及智能化应用的需明确不少于2个成效指标"
            ),
            passed=True,
            status="不适用",
            severity="通过",
            summary="当前项目级别不属于第21条市级项目范围，未执行指标数量校验。",
            metrics={
                "common_benefit_indicator_count": 0,
                "business_output_or_benefit_indicator_count": 0,
                "intelligent_project_detected": _contains_intelligent_keyword(document.raw_text),
                "achievement_indicator_count": 0,
                "indicator_scope_source": "project_scope_not_applicable",
                "indicator_lines_sample": [],
                "project_scope_applicability": applicability,
                "skipped_reason": applicability["reason"],
            },
            issues=[],
            suggestions=["请按项目实际预算级别和主管单位要求审查。"],
        )

    section_lines, table_rows, section_source = _extract_indicator_quantity_source(
        document,
        content=content,
        filename=filename,
    )
    if section_source == "2_6_section_not_found":
        return RuleResult(
            rule_excel_row=21,
            rule_name="项目成效考核目标数量设置合规性审查规则",
            rule_category="内容合规性审查规则",
            rule_description="指标设置数量应满足要求",
            judgement_condition=(
                "通用指标中至少设定3个效益指标，业务指标中至少设定4个产出指标或效益指标；"
                "涉及智能化应用的需明确不少于2个成效指标"
            ),
            passed=False,
            status="资料不足",
            severity="warning",
            summary="未识别到可核验的项目成效指标章节或指标表，无法确定指标数量；不要求固定章节编号。",
            metrics={
                "common_benefit_indicator_count": 0,
                "business_output_or_benefit_indicator_count": 0,
                "intelligent_project_detected": _contains_intelligent_keyword(document.raw_text),
                "achievement_indicator_count": 0,
                "indicator_scope_source": section_source,
                "indicator_lines_sample": [],
                "project_scope_applicability": applicability,
                "skipped_reason": "未识别到可核验的成效指标章节或指标表",
            },
            issues=[RuleIssue(
                message="资料不足：未识别到可核验的成效指标章节或指标表，不能自动判定数量是否达标。",
                evidence=None,
                section="2.6 项目成效考核目标（规划指标）",
            )],
            suggestions=["请指出现有成效指标的位置，或补充可验收的目标值及分类；不要求使用固定章节编号。"],
        )

    lines = section_lines
    indicator_lines = _collect_indicator_lines(lines)

    if table_rows:
        (
            common_benefit_count,
            business_output_or_benefit_count,
            achievement_count,
            counted_indicator_lines,
        ) = _count_indicator_table_rows(table_rows)
        if counted_indicator_lines:
            indicator_lines = counted_indicator_lines
    else:
        common_benefit_count = _count_indicator_rows(
            indicator_lines,
            scope_keywords=("通用指标", "通用"),
            type_keywords=("效益指标", "效益"),
        )
        business_output_or_benefit_count = _count_indicator_rows(
            indicator_lines,
            scope_keywords=("业务指标", "业务"),
            type_keywords=("产出指标", "效益指标", "产出", "效益"),
        )
        achievement_count = _count_achievement_indicator_rows(indicator_lines)

    intelligent_project = _has_intelligent_implementation(document.raw_text)

    issues: list[RuleIssue] = []
    suggestions: list[str] = []

    if common_benefit_count < 3:
        issues.append(
            RuleIssue(
                message=f"通用指标中的效益指标识别到 {common_benefit_count} 个，少于规则要求的 3 个。",
                evidence=_join_evidence([line for line in lines if _detect_indicator_scope(line) == 'common'] or lines[:8]),
            )
        )
        suggestions.append("请在通用指标中补足至少 3 个效益指标。")

    if business_output_or_benefit_count < 4:
        issues.append(
            RuleIssue(
                message=(
                    "业务指标中的产出指标或效益指标识别到 "
                    f"{business_output_or_benefit_count} 个，少于规则要求的 4 个。"
                ),
                evidence=_join_evidence([line for line in lines if _detect_indicator_scope(line) == 'business'] or indicator_lines),
            )
        )
        suggestions.append("请在业务指标中补足至少 4 个产出指标或效益指标。")

    if intelligent_project and achievement_count < 2:
        issues.append(
            RuleIssue(
                message=(
                    "文档疑似涉及智能化应用，但成效指标识别到 "
                    f"{achievement_count} 个，少于规则要求的 2 个。"
                ),
                evidence=_join_evidence(
                    [line for line in lines if _contains_intelligent_keyword(line)]
                ),
            )
        )
        suggestions.append("涉及智能化应用时，请明确不少于 2 个成效指标。")

    if issues:
        status = "发现问题"
        severity = "高"
        passed = False
        summary = "项目成效考核目标数量未完全满足规则要求。"
    else:
        status = "通过"
        severity = "通过"
        passed = True
        summary = "项目成效考核目标数量满足规则要求。"

    explicit_scope_labels = tuple(x for values in INDICATOR_SCOPE_ALIASES.values() for x in values if x not in {'通用', '业务'})
    known_scopes = {scope for scope, aliases in INDICATOR_SCOPE_ALIASES.items()
                    if any(_contains_any(line, tuple(a for a in aliases if a not in {'通用', '业务'})) for line in lines)
                    or any(cell.strip() in aliases for row in table_rows for cell in row)}
    classification_known = known_scopes == {'common', 'business'}
    if not classification_known:
        passed, status, severity = False, '资料不足', 'warning'
        summary = '已找到项目评价或成效描述，但未可靠识别通用/业务分类及可验收目标，不能将未分类条目按零项缺失判错。'
        issues = [RuleIssue(message=summary, evidence=_join_evidence(lines[:12]), section='项目评价/成效指标')]
        suggestions = ['请提供现有指标与通用、业务类别的对应说明及目标值；无需机械重写为固定模板。']
    if applicability.get('mode') == 'reference_check':
        passed, status, severity = False, '待复核', 'warning'
        summary = '已执行指标参考核验；项目预算范围存在共同出资或冲突，需确认后再认定数量结论。'
        suggestions = [applicability['reason']] + suggestions

    metrics = {
        "classification_identified": classification_known,
        "common_benefit_indicator_count": common_benefit_count,
        "business_output_or_benefit_indicator_count": business_output_or_benefit_count,
        "intelligent_project_detected": intelligent_project,
        "achievement_indicator_count": achievement_count,
        "achievement_indicator_evidence": [line for line in indicator_lines if _is_intelligent_outcome_indicator(line)],
        "indicator_scope_source": section_source,
        "indicator_lines_sample": indicator_lines[:20],
        "project_scope_applicability": applicability,
    }

    return RuleResult(
        rule_excel_row=21,
        rule_name="项目成效考核目标数量设置合规性审查规则",
        rule_category="内容合规性审查规则",
        rule_description="指标设置数量应满足要求",
        judgement_condition=(
            "通用指标中至少设定3个效益指标，业务指标中至少设定4个产出指标或效益指标；"
            "涉及智能化应用的需明确不少于2个成效指标"
        ),
        passed=passed,
        status=status,
        severity=severity,
        summary=summary,
        metrics=metrics,
        issues=issues,
        suggestions=suggestions or ["保持通用指标、业务指标和智能化应用成效指标数量满足规则要求。"],
    )


def _build_data_reporting_results_without_locations(
    content: bytes,
    filename: str,
    project_name: str = "",
    selected_rule_ids: list[str] | None = None,
    parsed_document: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    selected = {str(item).strip() for item in selected_rule_ids or [] if str(item).strip()}
    document = parsed_document or _parse_data_reporting_document(content, filename, selected)
    validation_errors: list[ValidationError] = []
    legacy_reporting_selected = not selected or any(
        _rule_selected(rule, selected) and str(rule["rule_id"]) != "DATA_REASON_024"
        for rule in DATA_REPORTING_RULES
    )
    if legacy_reporting_selected:
        validation = DataReportingValidator().validate(document)
        validation_errors.extend(validation.errors)
    governance_document = {
        **document,
        "_content": content,
        "_filename": filename,
        "_project_name": project_name,
    }
    governance_applicability = assess_data_governance_service_applicability(governance_document)
    governance_rule_selected = not selected or _rule_selected(
        next(rule for rule in DATA_REPORTING_RULES if str(rule["rule_id"]) == "DATA_REASON_024"),
        selected,
    )
    if governance_rule_selected and governance_applicability["applies"]:
        validation_errors.extend(validate_data_governance_service_design(governance_document))
    grouped_errors = _group_validation_errors_by_rule(validation_errors)

    results: list[dict[str, Any]] = []
    for rule in DATA_REPORTING_RULES:
        if selected and not _rule_selected(rule, selected):
            continue
        if str(rule["rule_id"]) == "DATA_REASON_024" and not governance_applicability["applies"]:
            if governance_applicability.get("indeterminate"):
                results.append(
                    {
                        "rule_id": rule["rule_id"],
                        "rule_excel_row": rule["excel_row"],
                        "rule_name": rule["rule_name"],
                        "rule_category": rule["rule_category"],
                        "rule_description": rule["rule_description"],
                        "judgement_condition": rule["judgement_condition"],
                        "passed": False,
                        "status": "资料不足",
                        "severity": "warning",
                        "summary": "未能可靠确定项目是否属于上海市市级预算单位，暂不执行第24条合规判定。",
                        "metrics": {
                            "error_count": 0,
                            "source": "class_file_data_reporting",
                            "project_scope_applicability": governance_applicability,
                            "skipped_reason": governance_applicability["reason"],
                        },
                        "issues": [
                            {
                                "message": "资料不足：缺少明确的项目地区、项目级别或预算单位归属信息。",
                                "section": "项目基本信息",
                                "evidence": "项目基本信息",
                                "severity": "warning",
                                "suggestion": "请补充项目地区、项目级别及预算单位归属后重新校验第24条。",
                            }
                        ],
                        "suggestions": ["请补充项目地区、项目级别及预算单位归属后重新校验第24条。"],
                    }
                )
            else:
                results.append(
                    {
                        "rule_id": rule["rule_id"],
                        "rule_excel_row": rule["excel_row"],
                        "rule_name": rule["rule_name"],
                        "rule_category": rule["rule_category"],
                        "rule_description": rule["rule_description"],
                        "judgement_condition": rule["judgement_condition"],
                        "passed": True,
                        "status": "不适用",
                        "severity": "pass",
                        "summary": "当前材料不属于上海市市级预算单位适用范围，未执行第24条内容校验。",
                        "metrics": {
                            "error_count": 0,
                            "source": "class_file_data_reporting",
                            "project_scope_applicability": governance_applicability,
                            "skipped_reason": governance_applicability["reason"],
                        },
                        "issues": [],
                        "suggestions": ["第24条仅适用于上海市市级预算单位；当前材料无需按本条整改。"],
                    }
                )
            continue
        issues = grouped_errors.get(str(rule["rule_id"]), [])
        suggestions = _unique_texts(
            [str(issue.get("suggestion") or "").strip() for issue in issues]
            or [str(rule["default_suggestion"])]
        )
        if not suggestions:
            suggestions = [str(rule["default_suggestion"])]

        metrics = {"error_count": len(issues), "source": "class_file_data_reporting"}
        if str(rule["rule_id"]) == "DATA_REASON_024":
            metrics["project_scope_applicability"] = governance_applicability
            metrics["policy_source"] = {k: POLICY[k] for k in ('version', 'source_file', 'sha256', 'services_pages', 'negative_list_page', 'acceptance_pages')}

        results.append(
            {
                "rule_id": rule["rule_id"],
                "rule_excel_row": rule["excel_row"],
                "rule_name": rule["rule_name"],
                "rule_category": rule["rule_category"],
                "rule_description": rule["rule_description"],
                "judgement_condition": rule["judgement_condition"],
                "passed": not issues,
                "status": "通过" if not issues else "发现问题",
                "severity": "pass" if not issues else _highest_validation_severity(
                    [str(issue.get("severity") or "warning") for issue in issues]
                ),
                "summary": "未发现问题。" if not issues else f"发现 {len(issues)} 项需复核内容。",
                "metrics": metrics,
                "issues": issues,
                "suggestions": suggestions,
            }
        )
    for result in results:
        if result['rule_excel_row'] == 24 and result['issues'] and all(i.get('severity') != 'risk' for i in result['issues']):
            result['status'] = '待复核'
            result['summary'] = f"有{len(result['issues'])}项信息需结合正文核实，尚不认定为确定不符合。"
        # The supplied 2026 guide is a benchmark for older reports, not an
        # automatically retroactive compliance obligation. Use report metadata
        # only, never years inside quoted regulations or hardware specs.
        years = re.findall(r'(?<!\d)(20\d{2})(?:年|年度|\d{4})', filename)
        if result['rule_excel_row'] == 24 and years and int(years[0]) < 2026 and result['status'] != '不适用':
            result['metrics']['policy_temporal_mode'] = 'reference_check'
            result['passed'], result['status'], result['severity'] = False, '待复核', 'warning'
            result['summary'] = f"报告年份为{years[0]}，按用户指定2026年指引对标，有{len(result['issues'])}项需核实；不据此作追溯不合规结论。"
            result['suggestions'] = ['请结合实际申报年度、政策生效及过渡安排确认适用标准。'] + result['suggestions']
            for issue in result['issues']:
                issue['severity'] = 'warning'
        scope = result.get('metrics', {}).get('project_scope_applicability', {})
        if scope.get('mode') == 'reference_check':
            result['metrics']['content_check_passed'] = result['passed']
            result['passed'] = False
            result['status'] = '待复核'
            result['severity'] = 'warning'
            result['summary'] = f"已执行内容对标，识别到{len(result['issues'])}项待核实内容；适用范围需结合预算申报主体确认，不据此认定违规。"
            result['suggestions'] = [scope['reason']] + result['suggestions']
            for issue in result['issues']:
                issue['severity'] = 'warning'
    return results


def _group_validation_errors_by_rule(errors: list[ValidationError]) -> dict[str, list[dict[str, Any]]]:
    grouped = {str(rule["rule_id"]): [] for rule in DATA_REPORTING_RULES}
    for error in errors:
        rule = _rule_for_error(str(error.code))
        if rule is None:
            continue
        grouped[str(rule["rule_id"])].append(
            {
                "message": str(error.message),
                "section": str(error.section),
                "evidence": str(error.section),
                "severity": str(error.severity),
                "suggestion": error.suggestion,
            }
        )
    return grouped


def _rule_for_error(code: str) -> dict[str, Any] | None:
    for rule in DATA_REPORTING_RULES:
        if any(code.startswith(prefix) for prefix in rule["error_prefixes"]):
            return rule
    return None


def _rule_selected(rule: dict[str, Any], selected: set[str]) -> bool:
    values = {
        str(rule.get("rule_id") or ""),
        str(rule.get("excel_row") or ""),
        str(rule.get("rule_name") or ""),
    }
    return bool(values & selected)


def _selected_is_governance_rule_only(selected: set[str]) -> bool:
    if not selected:
        return False
    governance_rule = next(
        rule for rule in DATA_REPORTING_RULES if str(rule["rule_id"]) == "DATA_REASON_024"
    )
    if not _rule_selected(governance_rule, selected):
        return False
    allowed_values = {
        str(governance_rule["rule_id"]),
        str(governance_rule["excel_row"]),
        str(governance_rule["rule_name"]),
    }
    return selected <= allowed_values


def _parse_data_reporting_document(
    content: bytes,
    filename: str,
    selected: set[str],
) -> dict[str, Any]:
    if _selected_is_governance_rule_only(selected) and filename.lower().endswith(".docx"):
        paragraphs = _parse_docx_paragraphs_only(content, filename)
        from docx import Document
        doc = Document(BytesIO(content))
        tables = '\n'.join(' | '.join(row) for t in doc.tables for row in table_grid(t))
        return {"_full_text": paragraphs.raw_text + '\n' + tables}
    return BaseValidator.parse_document(content, filename)


def _parse_docx_paragraphs_only(content: bytes, filename: str) -> DocumentContent:
    try:
        from docx import Document
    except Exception as exc:
        raise ValueError(f"Word 文档解析失败: {exc}") from exc

    try:
        doc = Document(BytesIO(content))
    except Exception as exc:
        raise ValueError(f"Word 文档解析失败: {exc}") from exc

    lines = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text and paragraph.text.strip()]
    return DocumentContent(format="docx", raw_text="\n".join(lines), sections=[], filename=filename)


def _rule_21_project_applicability(document, filename=None, project_name=None):
    from .project_scope import assess_project_applicability
    return assess_project_applicability({'_full_text': document.raw_text,
        '_filename': filename, '_project_name': project_name}, 21)


def _highest_validation_severity(values: list[str]) -> str:
    if any(value == "risk" for value in values):
        return "risk"
    if any(value == "warning" for value in values):
        return "warning"
    return values[0] if values else "warning"


def _unique_texts(values: list[str]) -> list[str]:
    output: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = str(value or "").strip()
        if not text or text in seen:
            continue
        seen.add(text)
        output.append(text)
    return output


def _result_to_dict(result: RuleResult) -> dict[str, Any]:
    data = asdict(result)
    data["rule_id"] = LEGACY_RULE_IDS.get(result.rule_excel_row, str(result.rule_excel_row))
    data["issues"] = [asdict(issue) for issue in result.issues]
    data["review_required"] = not result.passed
    data["evidence_scope"] = "可提取正文及表格；扫描图、附件及未明确的申报口径须另行核验"
    return data


def _meaningful_lines(text: str) -> list[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    return [line.strip() for line in normalized.split("\n") if line.strip()]


def _contains_any(text: str, keywords: Iterable[str]) -> bool:
    return any(keyword in text for keyword in keywords)


def _contains_fuzzy_keyword(text: str, keywords: Iterable[str], threshold: float = 0.75) -> bool:
    """Return True when text exactly contains a keyword or covers most keyword 2-grams.

    This is intentionally a recall helper for noun/header recognition only. Rule pass/fail
    decisions still use parsed amounts, quantities and indicator counts.
    """

    keyword_tuple = tuple(str(keyword or "") for keyword in keywords if str(keyword or ""))
    if _contains_any(text, keyword_tuple):
        return True
    return _best_keyword_bigram_recall(text, keyword_tuple) >= threshold


def _contains_project_total_keyword(text: str) -> bool:
    return _contains_fuzzy_keyword(text, PROJECT_TOTAL_ALIASES, threshold=0.78)


def _contains_detail_total_keyword(text: str) -> bool:
    return _contains_fuzzy_keyword(text, DETAIL_TOTAL_ALIASES, threshold=0.72)


def _contains_intelligent_keyword(text: str) -> bool:
    return bool(list(positive_mentions(text, INTELLIGENT_KEYWORDS)))


def _has_intelligent_implementation(text: str) -> bool:
    for clause in positive_mentions(text, INTELLIGENT_KEYWORDS):
        if any(t in clause for t in ("指导意见", "发展趋势", "远期", "不属于本期")):
            continue
        if any(t in clause for t in ("本项目", "本期", "建设", "开发", "部署", "实现", "提供", "推理", "准确率")):
            return True
    return False


def _best_keyword_bigram_recall(text: str, keywords: Iterable[str]) -> float:
    text_grams = _char_bigrams(text)
    if not text_grams:
        return 0.0
    best = 0.0
    for keyword in keywords:
        keyword_grams = _char_bigrams(keyword)
        if not keyword_grams:
            continue
        best = max(best, len(text_grams & keyword_grams) / len(keyword_grams))
    return best


def _char_bigrams(value: str) -> set[str]:
    compact = re.sub(r"[\s,，.。;；:：|/\\()（）\[\]【】《》<>_\-－—、]+", "", str(value or "")).lower()
    if len(compact) < 2:
        return set()
    return {compact[index : index + 2] for index in range(len(compact) - 1)}


def _extract_amounts_from_line(lines: list[str], index: int, full_text: str) -> list[float]:
    line = lines[index]
    value_text = _strip_table_prefix(line)
    explicit_matches = [
        match
        for match in MONEY_PATTERN.finditer(value_text)
        if match.group("currency") or match.group("unit")
    ]

    context_text = _line_context(lines, index)
    budget_context = _is_budget_amount_context(value_text, context_text)

    amounts: list[float] = []
    for match in explicit_matches:
        if not _money_match_is_budget_amount(value_text, match, budget_context):
            continue
        unit = (match.group("unit") or "").replace(" ", "")
        if unit == "万" and not _bare_wan_looks_like_money(value_text, match, budget_context):
            continue
        amount = _money_match_to_wan(match)
        if amount is not None and _amount_is_reasonable(amount):
            amounts.append(amount)

    if budget_context:
        amounts.extend(_extract_unitless_budget_amounts(lines, index, value_text, context_text))

    return _dedupe_amounts(amounts)


def _extract_project_total_amounts(line: str, amounts: list[float]) -> list[float]:
    """Use the amount immediately following the project-total label.

    Project-total sentences often include component amounts and a parenthetical
    calculation base.  Selecting the maximum amount from the sentence can turn
    the calculation base into the project total.
    """

    positions = [
        line.find(keyword)
        for keyword in PROJECT_TOTAL_ALIASES
        if line.find(keyword) >= 0
    ]
    if not positions:
        return []
    label_end = min(positions)
    suffix = line[label_end : min(len(line), label_end + 80)]
    explicit = [
        _money_match_to_wan(match)
        for match in MONEY_PATTERN.finditer(suffix)
        if (match.group("currency") or match.group("unit"))
    ]
    explicit = [value for value in explicit if value is not None and _amount_is_reasonable(value)]
    if explicit:
        return _dedupe_amounts(explicit[:1])
    return amounts[:1]


def _strip_table_prefix(line: str) -> str:
    return re.sub(r"^(?:表\d+行\d+|第\d+页表\d+行\d+)\t", "", line)


def _line_context(lines: list[str], index: int, window: int = 4) -> str:
    start = max(0, index - window)
    return "\n".join(_strip_table_prefix(line) for line in lines[start : index + 1])


def _is_budget_amount_context(value_text: str, context_text: str) -> bool:
    compact_context = re.sub(r"\s+", "", context_text)
    has_financial_context = _contains_fuzzy_keyword(compact_context, BUDGET_CONTEXT_ALIASES, threshold=0.70)
    has_local_money_unit = _has_amount_unit_text(compact_context)
    has_amount_column = _contains_fuzzy_keyword(compact_context, BUDGET_AMOUNT_COLUMN_ALIASES, threshold=0.70)

    # 单纯出现在资源清单、配置清单中的数字，即使附近有“单位”，也不能按预算金额处理。
    if _looks_like_resource_quantity_context(value_text) and not has_financial_context:
        return False

    return has_financial_context and (has_local_money_unit or has_amount_column)


def _has_amount_unit_text(text: str) -> bool:
    return bool(
        re.search(
            r"(单位|金额|投资|预算|概算|估算|费用|经费|资金|成本|造价|总价|合计|总计).*?"
            r"(万元|元(?![\u4e00-\u9fffA-Za-z]))",
            text,
        )
    )


def _looks_like_resource_quantity_context(text: str) -> bool:
    return bool(
        re.search(
            r"(?:数量|配置|规格|容量|内存|硬盘|CPU|服务器|数据库|操作系统|交换机|防火墙|存储|"
            r"GB|MB|TB|KB|GHz|MHz|个|台|套|块|条|核)",
            text,
            flags=re.IGNORECASE,
        )
    )


def _extract_unitless_budget_amounts(
    lines: list[str],
    index: int,
    value_text: str,
    context_text: str,
) -> list[float]:
    if not _has_local_wan_unit(context_text):
        return []

    amounts: list[float] = []
    table_cells = value_text.split("\t") if "\t" in value_text else []
    if table_cells:
        header_cells = _nearest_table_header_cells(lines, index, len(table_cells))
        if not header_cells:
            return []
        for cell_index, cell in enumerate(table_cells):
            if _cell_has_explicit_money(cell):
                continue
            if not _table_cell_is_amount_column(cell_index, cell, header_cells, value_text):
                continue
            unit = _amount_unit_for_cell(cell_index, header_cells, context_text)
            for match in NUMBER_PATTERN.finditer(cell):
                if _number_has_non_money_unit(cell, match):
                    continue
                if _looks_like_ordinal_or_year(cell, match, cell_index):
                    continue
                value = float(match.group("number").replace(",", ""))
                if unit == "元":
                    value = value / 10000
                if _amount_is_reasonable(value):
                    amounts.append(round(value, 4))
        return amounts

    return []


def _has_local_wan_unit(context_text: str) -> bool:
    compact_context = re.sub(r"\s+", "", context_text)
    return _has_amount_unit_text(compact_context)


def _nearest_table_header_cells(lines: list[str], index: int, cell_count: int) -> list[str]:
    current_table = _table_id(lines[index])
    if not current_table:
        return []

    for previous in range(index - 1, max(-1, index - 10), -1):
        if _table_id(lines[previous]) != current_table:
            continue
        cells = _strip_table_prefix(lines[previous]).split("\t")
        if len(cells) != cell_count:
            continue
        if _contains_fuzzy_keyword("".join(cells), (*BUDGET_AMOUNT_COLUMN_ALIASES, "数量", "单位", "规格"), threshold=0.70):
            return cells
    return []


def _table_id(line: str) -> str | None:
    match = re.match(r"^(表\d+|第\d+页表\d+)行\d+\t", line)
    return match.group(1) if match else None


def _cell_has_explicit_money(cell: str) -> bool:
    return any(match.group("currency") or match.group("unit") for match in MONEY_PATTERN.finditer(cell))


def _money_match_is_budget_amount(text: str, match: re.Match[str], budget_context: bool) -> bool:
    if _number_has_non_money_unit(text, match):
        return False

    prefix = text[max(0, match.start() - 12) : match.start()]
    suffix = text[match.end() : match.end() + 12]
    local_context = prefix + suffix
    if _looks_like_resource_quantity_context(local_context) and not _contains_fuzzy_keyword(text, BUDGET_CONTEXT_ALIASES, threshold=0.72):
        return False

    currency = bool(match.group("currency"))
    unit = (match.group("unit") or "").replace(" ", "")
    if unit == "万":
        return _bare_wan_looks_like_money(text, match, budget_context)
    if unit in {"万元", "元"}:
        return currency or budget_context or _contains_fuzzy_keyword(text, BUDGET_CONTEXT_ALIASES, threshold=0.72)
    return currency and not _looks_like_resource_quantity_context(local_context)


def _table_cell_is_amount_column(
    cell_index: int,
    cell: str,
    header_cells: list[str],
    row_text: str,
) -> bool:
    if cell_index == 0 and re.fullmatch(r"\d+(?:\.\d+)?", cell.strip()):
        return False

    if header_cells and cell_index < len(header_cells):
        header = header_cells[cell_index]
        if _contains_any(header, ("数量", "规格", "配置", "单位", "容量")):
            return False
        if _contains_any(header, NON_AMOUNT_COLUMN_KEYWORDS):
            return False
        if _contains_fuzzy_keyword(header, BUDGET_AMOUNT_COLUMN_ALIASES, threshold=0.70):
            return True
        return False

    return False


def _amount_unit_for_cell(cell_index: int, header_cells: list[str], context_text: str) -> str:
    header = header_cells[cell_index] if header_cells and cell_index < len(header_cells) else ""
    compact_header = re.sub(r"\s+", "", header)
    if "万元" in compact_header:
        return "万元"
    if "元" in compact_header:
        return "元"

    compact_context = re.sub(r"\s+", "", context_text)
    if "单位:万元" in compact_context or "单位：万元" in compact_context or "（万元）" in compact_context:
        return "万元"
    if "单位:元" in compact_context or "单位：元" in compact_context or "（元）" in compact_context:
        return "元"
    return "万元"


def _money_match_to_wan(match: re.Match[str]) -> float | None:
    raw_number = match.group("number")
    unit = (match.group("unit") or "").replace(" ", "")
    value = float(raw_number.replace(",", ""))

    if unit == "元":
        value = value / 10000
    return round(value, 4)


def _bare_wan_looks_like_money(text: str, match: re.Match[str], budget_context: bool) -> bool:
    suffix = text[match.end() : match.end() + 8]
    if re.match(r"\s*(?:\+|票|条|个|台|套|块|项|人|次|年|月|日|样本|记录|数据)", suffix):
        return False
    if _looks_like_resource_quantity_context(text):
        return False
    return budget_context and _contains_fuzzy_keyword(text, ("金额", "费用", "经费", "投资", "预算", "概算", "估算", "单价", "总价", "采购", "服务费"), threshold=0.72)


def _number_has_non_money_unit(text: str, match: re.Match[str]) -> bool:
    suffix = text[match.end() : match.end() + 8]
    return bool(NON_MONEY_UNIT_RE.match(suffix))


def _looks_like_ordinal_or_year(text: str, match: re.Match[str], cell_index: int) -> bool:
    number_text = match.group("number").replace(",", "")
    try:
        value = float(number_text)
    except ValueError:
        return True

    if cell_index == 0 and value < 1000 and re.fullmatch(r"\s*" + re.escape(match.group("number")) + r"\s*", text):
        return True
    if value < 1:
        return True
    if 1900 <= value <= 2100 and re.search(r"年|年度", text):
        return True
    if cell_index == 0 and value < 10 and not _contains_any(text, ("合计", "总计", "小计", "金额", "费用", "投资", "预算", "概算", "估算")):
        return True
    return False


def _amount_is_reasonable(amount_wan: float) -> bool:
    return amount_wan >= 0.0001


def _dedupe_amounts(amounts: list[float]) -> list[float]:
    deduped: list[float] = []
    for amount in amounts:
        if not any(_amounts_close(amount, existing) for existing in deduped):
            deduped.append(amount)
    return deduped


def _pick_representative_amount(amounts: list[float]) -> float:
    return max(amounts)


def _distinct_amounts(amounts: Iterable[float]) -> list[float]:
    distinct: list[float] = []
    for amount in amounts:
        if not any(_amounts_close(amount, existing) for existing in distinct):
            distinct.append(amount)
    return distinct


def _amounts_close(left: float, right: float) -> bool:
    tolerance = 0.05  # 万元；仅用于概算摘要四舍五入，不随总投资放大
    return abs(left - right) <= tolerance


def _format_amount(amount: float) -> str:
    formatted = f"{amount:.4f}".rstrip("0").rstrip(".")
    return f"{formatted} 万元"


def _is_detail_total_context(lines: list[str], index: int, line: str) -> bool:
    if _contains_detail_total_keyword(line):
        return True
    if index > 0 and "\t" in line and _contains_detail_total_keyword(lines[index - 1]):
        return True
    return False


def _is_non_project_budget_line(line: str) -> bool:
    """Exclude recurring operating charges from one-time project budget totals."""

    compact = re.sub(r"\s+", "", str(line or "")).lower()
    return any(token in compact for token in ("\u6bcf\u5e74", "\u5e74\u8d39", "\u6708\u8d39", "\u4e91\u79df\u8d39", "\u8fd0\u7ef4\u8d39", "\u79df\u8d39"))


def _project_total_breakdown_matches(line: str, project_total: float, amounts: list[float]) -> bool:
    """Accept a prose total only when its immediate two-part breakdown reconciles."""

    if len(amounts) != 3:
        return False
    return _amounts_close(amounts[0], sum(amounts[1:])) and _amounts_close(amounts[0], project_total)


def _find_total_row_arithmetic_issues(
    lines: list[str],
    amount_lines: list[tuple[int, str, list[float]]],
) -> list[RuleIssue]:
    issues: list[RuleIssue] = []
    for index, line, amounts in amount_lines:
        if not _is_detail_total_context(lines, index, line) or len(amounts) < 3:
            continue
        # Prose total lines may contain a parenthetical calculation base or
        # several nested totals.  They are not a flat arithmetic row.
        if ("（" in line or "(" in line) and len(amounts) >= 4:
            continue
        expected_total = sum(amounts[:-1])
        actual_total = amounts[-1]
        if not _amounts_close(expected_total, actual_total):
            issues.append(
                RuleIssue(
                    message=(
                        "识别到合计/总计行内部加总不一致："
                        f"前序金额合计为 {_format_amount(expected_total)}，"
                        f"行尾合计为 {_format_amount(actual_total)}。"
                    ),
                    evidence=line,
                )
            )
    return issues


def _count_total_row_arithmetic_checks(
    lines: list[str],
    amount_lines: list[tuple[int, str, list[float]]],
) -> int:
    return sum(
        1
        for index, _line, amounts in amount_lines
        if _is_detail_total_context(lines, index, lines[index]) and len(amounts) >= 3
        and not _contains_project_total_keyword(lines[index])
        and not (("（" in lines[index] or "(" in lines[index]) and len(amounts) >= 4)
    )


def _empty_budget_table_checks() -> dict[str, Any]:
    return {
        "detail_total_candidates": [],
        "arithmetic_issues": [],
        "tables_checked": 0,
        "evidence_examples": [],
        "table_internal_checks": 0,
        "summary_internal_checks": 0,
    }


def _budget_candidate_is_component_total(candidate: dict[str, Any]) -> bool:
    return str(candidate.get("kind") or "") == "detail_table_total" and not candidate.get("included_in")


def _budget_candidate_is_grand_total(candidate: dict[str, Any]) -> bool:
    if str(candidate.get("kind") or "") == "investment_summary_total":
        return True
    line = str(candidate.get("line") or "")
    return _contains_fuzzy_keyword(
        line,
        (
            "投资估算总表",
            "投资估算汇总表",
            "项目投资估算",
            "总投资估算",
            "项目总投资",
            "项目投资总预算",
            "预算总额",
            "总预算",
        ),
        threshold=0.70,
    )


def _sum_candidate_amounts(candidates: list[dict[str, Any]]) -> float | None:
    amounts = [
        float(candidate["amount_wan"])
        for candidate in candidates
        if candidate.get("amount_wan") is not None
    ]
    if not amounts:
        return None
    return round(sum(amounts), 4)


def _amount_exceeds(left: float, right: float) -> bool:
    tolerance = 0.05  # 万元；仅用于概算摘要四舍五入，不随总投资放大
    return left - right > tolerance


def _iter_docx_table_rows_fast(doc: Any) -> Iterable[list[list[str]]]:
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    table_tag = namespace + "tbl"
    row_tag = namespace + "tr"
    cell_tag = namespace + "tc"
    text_tag = namespace + "t"

    for child in doc.element.body.iterchildren():
        if child.tag != table_tag:
            continue
        yield table_grid(child, repeat_vertical=False)


def _extract_docx_budget_table_checks(content: bytes | None, filename: str | None) -> dict[str, Any]:
    """解析 Word 预算表，校验明细合计与投资估算总表。

    parse_document 的纯文本会丢失表格列定位，容易漏掉“总金额/申报金额”列。这里直接读 DOCX
    表格，保留 DOCX表格X行Y 作为前端可展示的文章依据。
    """

    if not content or not (filename or "").lower().endswith(".docx"):
        return _empty_budget_table_checks()

    try:
        from docx import Document
    except Exception:
        return _empty_budget_table_checks()

    try:
        doc = Document(BytesIO(content))
    except Exception:
        return _empty_budget_table_checks()

    detail_total_candidates: list[dict[str, Any]] = []
    arithmetic_issues: list[RuleIssue] = []
    evidence_examples: list[str] = []
    tables_checked = 0
    table_internal_checks = 0
    summary_internal_checks = 0
    relation_records = []

    for table_index, rows in enumerate(_iter_docx_table_rows_fast(doc), start=1):
        if not rows:
            continue

        table_text = " ".join(" ".join(row) for row in rows)
        if not _looks_like_budget_table(table_text):
            continue

        location = f"DOCX表格{table_index}"
        if any(keyword in table_text for keyword in ("投资估算总表", "投资估算表", "投资估算汇总表", "项目投资估算", "总投资估算")) or any("总投资" in ' '.join(row[:3]) for row in rows) or (any(_budget_hierarchy_level(r) == 1 for r in rows) and any(_budget_hierarchy_level(r) == 3 for r in rows)):
            tables_checked += 1
            summary_result = _check_investment_summary_table(rows, location)
            detail_total_candidates.extend(summary_result["detail_total_candidates"])
            arithmetic_issues.extend(summary_result["arithmetic_issues"])
            evidence_examples.extend(summary_result["evidence_examples"])
            summary_internal_checks += summary_result["summary_internal_checks"]
            continue

        amount_info = _budget_amount_column(rows)
        if amount_info is None:
            continue
        header_index, amount_col, unit = amount_info
        multiplication = _check_budget_products(rows, header_index, amount_col, unit, location)
        arithmetic_issues.extend(multiplication)
        total_rows = [
            (row_index, row)
            for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2)
            if _budget_row_is_total(row)
        ]
        if not total_rows:
            continue

        tables_checked += 1
        total_row_index, total_row = total_rows[-1]
        total_amount = _budget_row_amount(total_row, amount_col, unit, allow_fallback=True)
        data_rows: list[list[str]] = []
        for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            if row_index == total_row_index or _budget_row_is_total(row):
                continue
            data_rows.append(row)
        data_amounts = [
            _budget_row_amount(row, amount_col, unit, allow_fallback=False)
            for row in data_rows
        ]
        data_amounts = [amount for amount in data_amounts if amount is not None]
        evidence_line = f"{location}行{total_row_index}：{_budget_row_preview(total_row)}"

        if total_amount is not None:
            from .budget_relations import entries
            relation_records.append({"source": location, "total": total_amount, "entries": entries(rows, header_index, amount_col, unit)})
            detail_total_candidates.append(
                {
                    "line": evidence_line,
                    "amount_wan": total_amount,
                    "source": location,
                    "kind": "detail_table_total",
                }
            )
            evidence_examples.append(evidence_line)

        if total_amount is not None and data_amounts and _budget_table_allows_simple_sum(rows, header_index, total_row_index):
            table_internal_checks += 1
            expected = round(sum(data_amounts), 4)
            if abs(expected - total_amount) > 0.0001:
                arithmetic_issues.append(
                    RuleIssue(
                        message=_budget_sum_message(data_amounts, total_amount, data_rows, total_row, amount_col, unit),
                        evidence=evidence_line,
                        section=location,
                    )
                )

    from .budget_relations import find_subsets
    relations = find_subsets(relation_records)
    included = {r['child']: r['parent'] for r in relations if r['complete']}
    for candidate in detail_total_candidates:
        if candidate['source'] in included:
            candidate['included_in'] = included[candidate['source']]
    return {
        "detail_total_candidates": detail_total_candidates,
        "arithmetic_issues": arithmetic_issues,
        "tables_checked": tables_checked,
        "table_relations": relations,
        "evidence_examples": _dedupe_preserve_order(evidence_examples),
        "table_internal_checks": table_internal_checks,
        "summary_internal_checks": summary_internal_checks,
    }


def _rounding_envelope(rows, total_row, amount_col, unit):
    """Bound independent rounding, without relaxing integer/unit mismatches."""
    cells = [r[amount_col].strip() for r in [*rows, total_row] if amount_col < len(r)]
    if len(cells) != len(rows) + 1 or len(rows) < 2:
        return 0.0
    if any(_parse_budget_number(c) is None for c in cells):
        return 0.0
    # Integer amounts are not assumed to be rounded estimates. A decimal
    # presentation is necessary, and a possible rounding difference remains
    # a review item rather than silently passing.
    if not all('.' in c for c in cells):
        return 0.0
    return sum(0.5 * 10 ** -len(c.split('.')[1]) for c in cells) / (10000 if unit == '元' else 1)


def _budget_sum_message(amounts, total, rows, total_row, amount_col, unit):
    expected = sum(amounts)
    if unit == '万元' and max(amounts) > max(1000, abs(total) * 10):
        normalized = sum(v / 10000 if v > max(1000, abs(total) * 10) else v for v in amounts)
        explanation = (f'将异常大额行暂按元换算后，合计为{_format_amount(normalized)}，与表列总计一致。'
                       if abs(normalized - total) < 0.0001 else '部分明细的金额量级与表头单位明显冲突。')
        return '单位待核实：预算明细表疑似元/万元混填（未自动换算）。' + explanation + '请确认并统一单位，不据此认定巨额预算差错。'
    prefix = '舍入待核实：显示精度允许累计舍入解释该差异；' if abs(expected - total) <= _rounding_envelope(rows, total_row, amount_col, unit) + 1e-9 else ''
    return prefix + f'预算明细表合计行内部加总不一致：明细金额合计为 {_format_amount(expected)}，合计行金额为 {_format_amount(total)}。'


def _check_investment_summary_table(rows: list[list[str]], location: str) -> dict[str, Any]:
    amount_info = _budget_amount_column(rows)
    amount_col = amount_info[1] if amount_info else (_numeric_column_with_most_values(rows[1:]) if "万元" in ' '.join(rows[0]) else None)
    if amount_col is None:
        return {"detail_total_candidates": [], "arithmetic_issues": [], "evidence_examples": [], "summary_internal_checks": 0}

    system_row = _find_budget_amount_row_any(rows, ("系统建设费", "工程建设费", "建设费"), amount_col, "万元")
    other_row = _find_budget_amount_row_any(rows, ("其他费用", "工程建设其他费", "其他费"), amount_col, "万元")
    total_row = _find_budget_amount_row_any(rows, ("总投资", "总计", "投资估算合计", "合计"), amount_col, "万元")
    candidates: list[dict[str, Any]] = []
    issues: list[RuleIssue] = []
    evidence: list[str] = []
    summary_internal_checks = 0

    if total_row:
        total_index, total_cells = total_row
        total_amount = _budget_row_amount(total_cells, amount_col, "万元", allow_fallback=True)
        line = f"{location}行{total_index}：{_budget_row_preview(total_cells)}"
        if total_amount is not None:
            candidates.append({"line": line, "amount_wan": total_amount, "source": location, "kind": "investment_summary_total"})
            evidence.append(line)
        roots = [_budget_row_amount(r, amount_col, "万元") for r in rows[:total_index - 1] if _budget_hierarchy_level(r) == 1]
        if len(roots) >= 2 and all(x is not None for x in roots) and total_amount is not None:
            summary_internal_checks += 1
            if abs(sum(roots) - total_amount) > 0.0001:
                issues.append(RuleIssue(message=f"投资估算总表一级分项合计 {_format_amount(sum(roots))} 与总计 {_format_amount(total_amount)} 不一致。", evidence=line, section=location))

        if system_row and other_row and total_amount is not None:
            system_amount = _budget_row_amount(system_row[1], amount_col, "万元", allow_fallback=True)
            other_amount = _budget_row_amount(other_row[1], amount_col, "万元", allow_fallback=True)
            if system_amount is not None and other_amount is not None:
                summary_internal_checks += 1
                expected = round(system_amount + other_amount, 4)
                if abs(expected - total_amount) > 0.0001:
                    issues.append(
                        RuleIssue(
                            message=(
                                "投资估算总表总计与一级费用合计不一致："
                                f"建设费+其他费用为 {_format_amount(expected)}，"
                                f"总投资为 {_format_amount(total_amount)}。"
                            ),
                            evidence=" | ".join(
                                [
                                    f"{location}行{system_row[0]}：{_budget_row_preview(system_row[1])}",
                                    f"{location}行{other_row[0]}：{_budget_row_preview(other_row[1])}",
                                    line,
                                ]
                            ),
                            section=location,
                        )
                    )

    hierarchy_result = _check_hierarchical_budget_rows(rows, amount_col, "万元", location)
    issues.extend(hierarchy_result["arithmetic_issues"])
    evidence.extend(hierarchy_result["evidence_examples"])
    summary_internal_checks += hierarchy_result["checks"]

    return {
        "detail_total_candidates": candidates,
        "arithmetic_issues": issues,
        "evidence_examples": evidence,
        "summary_internal_checks": summary_internal_checks,
    }


def _looks_like_budget_table(text: str) -> bool:
    extra_budget_keywords = ("\u603b\u6295\u8d44\u8d39\u7528", "\u5efa\u8bbe\u8d39\u7528\u5408\u8ba1")
    return _contains_any(
        text,
        extra_budget_keywords + (
            "投资估算",
            "预算",
            "总金额",
            "申报金额",
            "总价",
            "小计（万元）",
            "小计(万元)",
            "金额",
            "单价",
            "工作量",
            "系统建设费",
            "其他费用",
        ),
    )


def _budget_table_allows_simple_sum(rows: list[list[str]], header_index: int, total_row_index: int) -> bool:
    if total_row_index <= header_index + 1:
        return False

    body_rows = rows[header_index + 1 : total_row_index - 1]
    if not body_rows:
        return False

    return True


def _budget_amount_column(rows: list[list[str]]) -> tuple[int, int, str] | None:
    for row_index, row in enumerate(rows[:8]):
        for col_index, cell in enumerate(row):
            compact = re.sub(r"\s+", "", cell)
            if any(keyword in compact for keyword in ("总金额", "申报金额", "总价", "合价", "小计", "金额", "估算投资")) and not any(t in compact for t in ("单价", "说明", "备注")):
                unit = "元" if "元" in compact and "万元" not in compact else "万元"
                if "元" not in compact:
                    column_values = [
                        _parse_budget_number(candidate[col_index])
                        for candidate in rows[row_index + 1 :]
                        if col_index < len(candidate)
                    ]
                    column_values = [value for value in column_values if value is not None]
                    if column_values and max(column_values) >= 10000:
                        unit = "元"
                return row_index, col_index, unit
    return None


def _check_budget_products(rows, header_index, amount_col, unit, location):
    header = rows[header_index]
    price_col = next((i for i, c in enumerate(header) if "单价" in c), None)
    quantity_col = next((i for i, c in enumerate(header) if "数量" in c or "工作量" in c), None)
    if price_col is None or quantity_col is None:
        return []
    price_header = header[price_col]
    price_unit = "万元" if "万元" in price_header else "元" if "元" in price_header else unit
    errors = []
    for ri, row in enumerate(rows[header_index + 1:], header_index + 2):
        if _budget_row_is_total(row) or max(price_col, quantity_col, amount_col) >= len(row):
            continue
        q = _parse_budget_number(row[quantity_col])
        p = _parse_budget_number(row[price_col])
        a = _budget_row_amount(row, amount_col, unit)
        if q is None or p is None or a is None:
            continue
        expected = q * p / (10000 if price_unit == "元" else 1)
        # Compare at the stated monetary precision, not a 0.1% tolerance.
        if abs(expected - a) > 0.0001:
            errors.append(RuleIssue(
                message=f"预算行乘积不一致：数量/工作量 {q:g} × 单价 {p:g}（{price_unit}）为 {_format_amount(expected)}，填报 {_format_amount(a)}。",
                evidence=f"{location}行{ri}：{_budget_row_preview(row)}", section=location))
    return errors


def _numeric_column_with_most_values(rows: list[list[str]]) -> int | None:
    max_cols = max((len(row) for row in rows), default=0)
    best_col: int | None = None
    best_count = 0
    for col in range(max_cols):
        count = 0
        for row in rows:
            if col < len(row) and _parse_budget_number(row[col]) is not None:
                count += 1
        if count > best_count:
            best_col = col
            best_count = count
    return best_col if best_count else None


def _find_budget_amount_row_any(
    rows: list[list[str]],
    keywords: Iterable[str],
    amount_col: int,
    unit: str,
) -> tuple[int, list[str]] | None:
    for keyword in keywords:
        for row_index, row in enumerate(rows, start=1):
            row_text = " ".join(row)
            if keyword in row_text or _contains_fuzzy_keyword(row_text, (keyword,), threshold=0.75):
                if _budget_row_amount(row, amount_col, unit) is not None:
                    return row_index, row
    return None


def _check_hierarchical_budget_rows(
    rows: list[list[str]],
    amount_col: int,
    unit: str,
    location: str,
) -> dict[str, Any]:
    issues: list[RuleIssue] = []
    evidence: list[str] = []
    checks = 0
    parsed_rows: list[dict[str, Any]] = []

    for row_index, row in enumerate(rows, start=1):
        level = _budget_hierarchy_level(row)
        amount = _budget_row_amount(row, amount_col, unit)
        parsed_rows.append({"row_index": row_index, "row": row, "level": level, "amount": amount})

    for index, current in enumerate(parsed_rows):
        current_level = current["level"]
        current_amount = current["amount"]
        if current_level is None or current_amount is None:
            continue

        descendants: list[dict[str, Any]] = []
        for candidate in parsed_rows[index + 1 :]:
            candidate_level = candidate["level"]
            if candidate_level is None:
                continue
            if candidate_level <= current_level:
                break
            descendants.append(candidate)

        if not descendants:
            continue
        direct_level = min(item["level"] for item in descendants if item["level"] is not None)
        direct_amounts = [
            item["amount"]
            for item in descendants
            if item["level"] == direct_level and item["amount"] is not None
        ]
        if len(direct_amounts) < 2:
            continue

        checks += 1
        expected = round(sum(direct_amounts), 4)
        if not _amounts_close(expected, current_amount):
            evidence_line = f"{location}行{current['row_index']}：{_budget_row_preview(current['row'])}"
            direct_rows = [item['row'] for item in descendants if item['level'] == direct_level and item['amount'] is not None]
            rounding = abs(expected - current_amount) <= _rounding_envelope(direct_rows, current['row'], amount_col, unit) + 1e-9
            issues.append(
                RuleIssue(
                    message=(
                        ("舍入待核实：显示金额的差异可能来自累计四舍五入，请核对未舍入金额；" if rounding else "") + "投资估算表层级加总不一致："
                        f"下级分项合计为 {_format_amount(expected)}，"
                        f"本级金额为 {_format_amount(current_amount)}。"
                    ),
                    evidence=evidence_line,
                    section=location,
                    severity='warning' if rounding else 'risk',
                )
            )
            evidence.append(evidence_line)

    return {"arithmetic_issues": issues, "evidence_examples": evidence, "checks": checks}


def _budget_hierarchy_level(row: list[str]) -> int | None:
    first_cell = re.sub(r"\s+", "", row[0] if row else "")
    if not first_cell:
        return None

    chinese_number = r"一二三四五六七八九十"
    if re.match(rf"^[{chinese_number}]+(?:[、.．]|$)", first_cell):
        return 1
    if re.match(rf"^[（(][{chinese_number}]+[）)]", first_cell):
        return 2

    number_match = re.match(r"^(\d+(?:\.\d+)*)", first_cell)
    if number_match:
        token = number_match.group(1)
        return 3 + token.count(".")

    return None


def _budget_row_is_total(row: list[str]) -> bool:
    leading = " ".join(row[:3])
    return "合计" in leading or "总计" in leading or "总投资" in leading


def _budget_row_amount(
    row: list[str],
    amount_col: int,
    unit: str,
    allow_fallback: bool = False,
) -> float | None:
    amount = _parse_budget_number(row[amount_col] if amount_col < len(row) else "")
    if amount is None and allow_fallback:
        numeric_values = [_parse_budget_number(cell) for cell in row]
        numeric_values = [value for value in numeric_values if value is not None]
        amount = numeric_values[-1] if numeric_values else None
    if amount is None:
        return None
    return round(amount / 10000, 4) if unit == "元" else round(amount, 4)


def _parse_budget_number(value: str) -> float | None:
    text = unicodedata.normalize('NFKC', str(value or "")).replace(",", "").strip()
    if not text:
        return None
    match = re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)", text)
    if not match:
        return None
    return float(text)


def _budget_row_preview(row: list[str]) -> str:
    return " | ".join(cell for cell in row if cell)[:260]


def _distinct_budget_candidates(candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    seen: set[tuple[str, float]] = set()
    for candidate in candidates:
        line = str(candidate.get("line") or "")
        amount = round(float(candidate.get("amount_wan") or 0), 4)
        # Equal totals do not prove two tables describe the same component.
        key = (line, amount)
        if key in seen:
            continue
        seen.add(key)
        output.append(candidate)
    return output


def _extract_indicator_quantity_source(
    document: DocumentContent,
    content: bytes | None = None,
    filename: str | None = None,
) -> tuple[list[str], list[list[str]], str]:
    actual_filename = filename or document.filename
    if content and actual_filename.lower().endswith(".docx"):
        docx_lines, docx_table_rows = _extract_docx_2_6_indicator_source(content)
        if docx_lines or docx_table_rows:
            return docx_lines, docx_table_rows, "docx_2_6_section"

    raw_lines = _extract_text_2_6_section_lines(_meaningful_lines(document.raw_text))
    if raw_lines:
        return raw_lines, [], "raw_text_2_6_section"

    # Narrative evaluation chapters are evidence too, even without Shanghai's table format.
    # The shared parser stores Word heading styles in sections rather than raw_text.
    section_text = '\n'.join(s.title + '\n' + s.content for s in document.sections)
    all_lines = _meaningful_lines(section_text or document.raw_text)
    candidates = []
    for i, line in enumerate(all_lines):
        if len(line) <= 45 and re.search(r'项目评价指标分析|绩效目标|成效指标|社会效益分析|经济效益分析', line) and not re.search(r'\.{3}|…|\d{2,}$', line):
            candidates.append(all_lines[i:i + 65])
    if candidates:
        best = max(candidates, key=lambda ls: sum('指标' in x or '目标' in x for x in ls))
        return best, [], 'semantic_evaluation_text'

    return [], [], "2_6_section_not_found"


def _extract_docx_2_6_indicator_source(content: bytes) -> tuple[list[str], list[list[str]]]:
    try:
        from docx import Document
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except Exception:
        return [], []

    try:
        doc = Document(BytesIO(content))
    except Exception:
        return [], []

    candidates: list[tuple[list[str], list[list[str]]]] = []
    current_lines: list[str] | None = None
    current_table_rows: list[list[str]] = []
    current_heading_level: int | None = None

    def close_current() -> None:
        nonlocal current_lines, current_table_rows, current_heading_level
        if current_lines is not None:
            candidates.append((current_lines, current_table_rows))
        current_lines = None
        current_table_rows = []
        current_heading_level = None

    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, doc)
            text = paragraph.text.strip()
            style_match = re.search(r'heading\s*(\d+)', paragraph.style.name, re.I) if paragraph.style else None
            heading_level = int(style_match.group(1)) if style_match else None
            if not text:
                continue
            if _is_2_6_indicator_heading(text):
                close_current()
                current_lines = [text]
                current_table_rows = []
                current_heading_level = heading_level
                continue
            if current_lines is not None and (_is_after_2_6_heading(text)
                    or heading_level is not None and current_heading_level is not None and heading_level <= current_heading_level
                    or re.match(r'^第[一二三四五六七八九十\d]+章', text)):
                close_current()
                continue
            if current_lines is not None:
                current_lines.append(text)
        elif isinstance(child, CT_Tbl) and current_lines is not None:
            table = Table(child, doc)
            for cells in table_grid(table):
                if not cells:
                    continue
                current_table_rows.append(cells)
                current_lines.append(_indicator_cells_to_line(cells))

    close_current()
    # Some Word reports place the 2.6 heading in a text box or TOC, so the
    # body-order walk above never opens a section.  Indicator tables still have
    # stable header cells; use them as a bounded fallback instead of auto-pass.
    if not candidates or not any(
        _looks_like_indicator_table_data_cells(row)
        for _lines, rows in candidates
        for row in rows
    ):
        fallback_rows: list[list[str]] = []
        fallback_lines: list[str] = ["2.6 项目成效考核目标（规划指标）"]
        for table in doc.tables:
            table_rows = table_grid(table)
            if not any(_indicator_table_header_map(row) for row in table_rows):
                continue
            fallback_rows.extend(row for row in table_rows if row)
            fallback_lines.extend(_indicator_cells_to_line(row) for row in table_rows if row)
        if fallback_rows:
            candidates.append((fallback_lines, fallback_rows))
    return _select_best_indicator_candidate(candidates)


def _select_best_indicator_candidate(
    candidates: list[tuple[list[str], list[list[str]]]]
) -> tuple[list[str], list[list[str]]]:
    if not candidates:
        return [], []

    def score(candidate: tuple[list[str], list[list[str]]]) -> tuple[int, int, int]:
        lines, rows = candidate
        table_indicator_rows = sum(1 for row in rows if _looks_like_indicator_table_data_cells(row))
        indicator_line_count = sum(1 for line in lines if _looks_like_indicator_line(line))
        return table_indicator_rows, indicator_line_count, len(lines)

    return max(candidates, key=score)


def _extract_text_2_6_section_lines(lines: list[str]) -> list[str]:
    selected: list[str] = []
    in_section = False

    for line in lines:
        if _is_2_6_indicator_heading(line):
            selected = [line]
            in_section = True
            continue
        if in_section and _is_after_2_6_heading(line):
            break
        if in_section:
            selected.append(line)

    return selected


def _is_2_6_indicator_heading(text: str) -> bool:
    if re.search(r'\t\s*\d+\s*$|[.．…]{3,}\s*\d+\s*$', text):
        return False  # TOC entries do not open a body evidence section.
    compact = re.sub(r"\s+", "", text)
    return len(compact) <= 80 and ("项目成效考核目标" in compact or compact.startswith("规划指标") or bool(re.search(r'绩效目标|成效指标|项目评价指标分析', compact)))


def _is_after_2_6_heading(text: str) -> bool:
    compact = re.sub(r"\s+", "", text)
    return bool(
        re.match(r"^2[\.．](?:7|8|9|[1-9]\d)(?:[\.．、:：]|$)", compact)
        or re.match(r"^[3-9][\.．]\d*", compact)
    )


def _clean_indicator_cell(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\r", "\n").replace("\n", " ")).strip()


def _indicator_cells_to_line(cells: list[str]) -> str:
    return " | ".join(cell for cell in cells if cell)


def _count_indicator_table_rows(
    table_rows: list[list[str]],
) -> tuple[int, int, int, list[str]]:
    common_benefit_count = 0
    business_output_or_benefit_count = 0
    achievement_count = 0
    counted_lines: list[str] = []
    header_map: dict[str, int] = {}
    # Word 表格常用合并单元格，后续行的“一级/二级指标”单元格可能为空。
    # 保留上一行的层级，避免合并单元格导致指标被漏计。
    current_scope: str | None = None
    current_type: str | None = None

    seen_indicators = set()
    inherited_subject = ''
    for cells in table_rows:
        normalized_cells = [_clean_indicator_cell(cell) for cell in cells]
        if not normalized_cells or _is_indicator_table_title_row(normalized_cells):
            continue
        current_header_map = _indicator_table_header_map(normalized_cells)
        if current_header_map:
            header_map = current_header_map
            current_scope = None
            current_type = None
            inherited_subject = ''
            continue
        if not _looks_like_indicator_table_data_cells(normalized_cells) and not (
            header_map and _looks_like_indicator_merged_data_row(normalized_cells)
        ):
            continue

        line = _indicator_cells_to_line(normalized_cells)
        scope = _detect_indicator_table_scope(normalized_cells, header_map)
        indicator_type = _detect_indicator_table_type(normalized_cells, header_map)
        if scope:
            if scope != current_scope:
                current_type = None
                inherited_subject = ''
            current_scope = scope
        else:
            scope = current_scope
        if _get_indicator_table_cell(normalized_cells, header_map, "type", 2):
            current_type = indicator_type
        else:
            indicator_type = current_type
        if not scope or not indicator_type:
            continue

        value_cell = _get_indicator_table_cell(normalized_cells, header_map, "value", 4)
        if not _indicator_has_target_value(value_cell, line):
            continue

        name_index = header_map.get('name', header_map.get('level3', max(0, len(normalized_cells) - 2)))
        name = normalized_cells[name_index] if name_index < len(normalized_cells) else line
        identity = (scope, indicator_type, re.sub(r'[\s、，,：:（）()]', '', name).lower())
        if identity in seen_indicators:
            continue
        seen_indicators.add(identity)
        if 'name' in header_map and 'level3' in header_map:
            subject = _get_indicator_table_cell(normalized_cells, header_map, 'level3', 3)
            if subject:
                inherited_subject = subject
        contextual_line = inherited_subject + ' | ' + line if inherited_subject else line

        counted = False
        if scope == "common" and indicator_type == "benefit":
            common_benefit_count += 1
            counted = True
        if scope == "business" and indicator_type in {"output", "benefit"}:
            business_output_or_benefit_count += 1
            counted = True
        if _is_intelligent_outcome_indicator(contextual_line) and indicator_type in {"output", "benefit"}:
            achievement_count += 1
            counted = True
        if counted:
            counted_lines.append(contextual_line)

    return (
        common_benefit_count,
        business_output_or_benefit_count,
        achievement_count,
        _dedupe_preserve_order(counted_lines),
    )


def _indicator_table_header_map(cells: list[str]) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for index, cell in enumerate(cells):
        compact = re.sub(r"\s+", "", cell)
        if "一级指标" in compact:
            header_map["scope"] = index
        elif "二级指标" in compact:
            header_map["type"] = index
        elif "三级指标" in compact:
            header_map["level3"] = index
        elif "四级指标" in compact or "指标名称" in compact:
            header_map["name"] = index
        elif "指标值" in compact or "目标值" in compact:
            header_map["value"] = index
    if {"scope", "type"} <= set(header_map):
        return header_map
    return {}


def _is_indicator_table_title_row(cells: list[str]) -> bool:
    unique_cells = {re.sub(r"\s+", "", cell) for cell in cells if cell}
    return len(unique_cells) == 1 and any(
        keyword in next(iter(unique_cells), "")
        for keyword in ("规划指标参数", "项目成效考核目标", "绩效目标")
    )


def _looks_like_indicator_table_data_cells(cells: list[str]) -> bool:
    compact_line = re.sub(r"\s+", "", _indicator_cells_to_line(cells))
    if not compact_line or _indicator_table_header_map(cells):
        return False
    if "序号" in compact_line and ("一级指标" in compact_line or "二级指标" in compact_line):
        return False
    return bool(_detect_indicator_scope(compact_line) and _detect_indicator_type(compact_line))


def _looks_like_indicator_merged_data_row(cells: list[str]) -> bool:
    """识别因合并层级单元格而缺少 scope/type 文本的指标数据行。"""

    compact_line = re.sub(r"\s+", "", _indicator_cells_to_line(cells))
    if not compact_line or ("指标名称" in compact_line and "目标值" in compact_line):
        return False
    return len(cells) >= 2 and bool(re.search(r"\d|%|率|量|数|提升|降低|满意|覆盖|完成", compact_line))


def _detect_indicator_table_scope(cells: list[str], header_map: dict[str, int]) -> str | None:
    scope_cell = _get_indicator_table_cell(cells, header_map, "scope", 1)
    return _detect_indicator_scope(scope_cell) or _detect_indicator_scope(_indicator_cells_to_line(cells))


def _detect_indicator_table_type(cells: list[str], header_map: dict[str, int]) -> str | None:
    type_cell = _get_indicator_table_cell(cells, header_map, "type", 2)
    return _detect_indicator_type(type_cell) or _detect_indicator_type(_indicator_cells_to_line(cells))


def _get_indicator_table_cell(
    cells: list[str],
    header_map: dict[str, int],
    key: str,
    fallback_index: int,
) -> str:
    index = header_map.get(key, fallback_index)
    if 0 <= index < len(cells):
        return cells[index]
    return ""


def _collect_indicator_lines(lines: list[str]) -> list[str]:
    selected: list[str] = []
    in_indicator_section = False

    for line in lines:
        if _contains_any(line, ("2.6", "绩效目标", "项目成效", "成效考核", "规划指标")):
            in_indicator_section = True

        has_indicator_keyword = _contains_any(line, ("通用指标", "业务指标", "产出指标", "效益指标", "成效指标", "指标"))
        if (in_indicator_section or has_indicator_keyword) and has_indicator_keyword:
            if _looks_like_indicator_line(line):
                selected.append(line)

        if in_indicator_section and _contains_any(line, ("2.7", "投资概况", "建设内容", "预算编制")):
            in_indicator_section = False

    return _dedupe_preserve_order(selected)


def _looks_like_indicator_line(line: str) -> bool:
    compact = re.sub(r"\s+", "", line)
    if len(compact) < 5:
        return False

    label_only = compact
    for token in ("通用指标", "业务指标", "产出指标", "效益指标", "成效指标", "指标名称", "指标类型", "一级指标", "二级指标", "三级指标", "目标值", "单位"):
        label_only = label_only.replace(token, "")
    label_only = re.sub(r"[:：|,，;；\-/\\（）()【】\[\]\t]", "", label_only)

    return bool(label_only) or bool(re.search(r"\d|%|≥|≤|>|<", compact))


def _count_indicator_rows(
    lines: list[str],
    scope_keywords: tuple[str, ...],
    type_keywords: tuple[str, ...],
) -> int:
    count = 0
    target_scope = "common" if "通用" in "".join(scope_keywords) else "business"
    target_types = {
        _detect_indicator_type(keyword)
        for keyword in type_keywords
        if _detect_indicator_type(keyword)
    }
    current_scope: str | None = None
    current_type: str | None = None

    for line in dict.fromkeys(lines):
        detected_scope = _detect_indicator_scope(line)
        detected_type = _detect_indicator_type(line)
        if detected_scope:
            current_scope = detected_scope
        if detected_type:
            current_type = detected_type

        scoped = current_scope == target_scope
        typed = current_type in target_types

        if scoped and typed and _looks_like_indicator_data_row(line) and _indicator_has_target_value("", line):
            count += 1

    return count


def _detect_indicator_scope(line: str) -> str | None:
    if any(x in line for x in INDICATOR_SCOPE_ALIASES['common']):
        return "common"
    if any(x in line for x in INDICATOR_SCOPE_ALIASES['business']):
        return "business"
    return None


def _detect_indicator_type(line: str) -> str | None:
    if "产出指标" in line or "产出" in line:
        return "output"
    if "效益指标" in line or "效益" in line:
        return "benefit"
    return None


def _looks_like_indicator_data_row(line: str) -> bool:
    compact = re.sub(r"\s+", "", line)
    header_tokens = ("指标名称", "指标类型", "一级指标", "二级指标", "三级指标", "目标值", "单位")
    if all(token in compact for token in ("指标", "名称")) and not re.search(r"\d|%|率|量|数|提升|降低|满意|覆盖|完成", compact):
        return False
    if compact in header_tokens:
        return False
    return len(compact) >= 8


def _is_intelligent_outcome_indicator(line: str) -> bool:
    context = _contains_intelligent_keyword(line) or _has_positive_indicator_agent_context(line)
    if not context:
        return False
    # Construction scale and metadata inventory are not application outcomes.
    if re.search(r'形成语料库条目数|(?:开发|购置|建设|建成|部署)数量|目录覆盖率|编目更新|知识库专题数量', line):
        return False
    return bool(re.search(r'准确率|召回率|精确率|识别率|正确率|响应时间|推理速度|交互.*时间|耗时|节省|节约|效率|满意度|采纳率|成功率|完成率|误报率|漏报率|使用率|调用量|服务人次', line))


def _has_positive_indicator_agent_context(line: str) -> bool:
    return bool(list(positive_mentions(line, ('智能体', '模型推理', '模型回答'))))


def _count_achievement_indicator_rows(lines: list[str]) -> int:
    count = 0
    in_achievement_section = False

    for line in dict.fromkeys(lines):
        if "成效指标" in line:
            in_achievement_section = True
        if _is_intelligent_outcome_indicator(line) and _looks_like_indicator_data_row(line) and _indicator_has_target_value("", line):
            count += 1

    return count


def _indicator_has_target_value(value_cell: str, line: str) -> bool:
    """Count only indicators with an explicit target or measurable criterion."""

    value = str(value_cell or "").strip()
    if value in {"-", "—", "/", "待定", "待补充", "无", "未填写"} or re.search(r'待(?:定|补充|明确|确认)|尚未|另行确定', value):
        return False
    if value and not all(token in value for token in ("指标", "名称", "类型")):
        return True
    return bool(
        re.search(
            r"(?:\d+(?:\.\d+)?\s*(?:%|％|人|项|次|天|月|年|万元|元|个|家|户|件|套|台)|"
            r"(?:不少于|不低于|不超过|达到|提升|降低)\d+)",
            line,
        )
    )


def _join_evidence(lines: list[str], limit: int = 5) -> str | None:
    if not lines:
        return None
    return " | ".join(lines[:limit])


def _dedupe_preserve_order(lines: list[str]) -> list[str]:
    seen: set[str] = set()
    output: list[str] = []
    for line in lines:
        key = re.sub(r"\s+", "", line)
        if key not in seen:
            seen.add(key)
            output.append(line)
    return output




def _retain_data_evidence_locations(result, content, filename):
    if not content or not str(filename or '').lower().endswith('.docx'):
        return result
    from app.modules.resource.quantity_parser import DocumentIndex, build_evidence_location
    rows = result.issues if isinstance(result, RuleResult) else result.get('issues', [])
    if not rows:
        return result
    try:
        index = DocumentIndex(content, filename)
    except Exception:
        return result
    suggestions = result.suggestions if isinstance(result, RuleResult) else result.get('suggestions', [])
    advice = next((s for s in suggestions if s), '')
    locations = []
    for issue in rows:
        data = asdict(issue) if isinstance(issue, RuleIssue) else issue
        evidence = data.get('evidence') or data.get('section') or ''
        locations.append(build_evidence_location(index, [evidence], data.get('section') or '', advice=advice))
    if isinstance(result, RuleResult):
        return replace(result, evidence_locations=locations)
    return {**result, 'evidence_locations': locations}


def _evaluate_budget_consistency(document, content=None, filename=None):
    result = _evaluate_budget_consistency_without_locations(document, content=content, filename=filename)
    return _retain_data_evidence_locations(result, content, filename)


def _evaluate_indicator_quantity(document, content=None, filename=None, project_name=None):
    result = _evaluate_indicator_quantity_without_locations(document, content=content, filename=filename, project_name=project_name)
    if result.severity == 'warning':
        result = replace(result, issues=[replace(issue, severity='warning') for issue in result.issues])
    return _retain_data_evidence_locations(result, content, filename)


def _build_data_reporting_results(content, filename, project_name='', selected_rule_ids=None, parsed_document=None):
    results = _build_data_reporting_results_without_locations(content, filename, project_name, selected_rule_ids, parsed_document)
    return [_retain_data_evidence_locations(result, content, filename) if str(result.get('rule_id')) == 'DATA_REASON_024' else result for result in results]
