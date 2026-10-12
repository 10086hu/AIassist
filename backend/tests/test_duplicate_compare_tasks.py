from io import BytesIO

from docx import Document
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.api import evaluate
from app.db.models import Base
from app.modules.duplicate import service
from app.modules.duplicate.compact_extractor import extract_compact_modules
from app.modules.duplicate.document_parser import parse_document
from app.modules.duplicate.llm_judge import LLMJudgement


def client(monkeypatch, handler):
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    app = FastAPI()
    app.include_router(evaluate.router)

    def db():
        with Session(engine) as session:
            yield session

    app.dependency_overrides[evaluate.get_db] = db
    monkeypatch.setattr(evaluate, "SessionLocal", lambda: Session(engine))
    monkeypatch.setattr(evaluate, "_evaluate_duplicate_compare_sync", handler)
    return TestClient(app)


def test_compare_task_reports_progress_and_returns_complete_result(monkeypatch):
    observed = []

    def compare(*args, progress_reporter=None):
        assert args[1] == b"current"
        assert args[3] == [(b"history", "history.docx")]
        for fraction, message in [(.35, "已整理6个模块"), (.1, "解析往期报告"), (.8, "已完成2/3批复核")]:
            progress_reporter(fraction, message)
            observed.append(next(reversed(evaluate.TASKS.values()))["progress"])
        return {"status": "completed", "summary": {"history_imported_count": 3}}

    api = client(monkeypatch, compare)
    response = api.post("/duplicate/compare/tasks", files=[("current_file", ("current.docx", b"current")),
                                                        ("history_files", ("history.docx", b"history"))])
    assert response.status_code == 200
    task = api.get("/tasks/" + response.json()["task_id"]).json()
    assert task["status"] == "completed"
    assert task["progress"] == 100
    assert task["result"]["summary"]["history_imported_count"] == 3
    assert observed == sorted(observed)


def test_failed_compare_task_is_not_reported_as_complete(monkeypatch):
    def fail(*args, **kwargs):
        raise HTTPException(status_code=400, detail="批量结果不完整")

    api = client(monkeypatch, fail)
    response = api.post("/duplicate/compare/tasks", files=[("current_file", ("current.docx", b"current")),
                                                        ("history_files", ("history.docx", b"history"))])
    task = api.get("/tasks/" + response.json()["task_id"]).json()
    assert task["status"] == "failed"
    assert "批量结果不完整" in task["error"]
    assert "result" not in task


def test_custom_style_numbered_long_report_restores_modules_and_preserves_late_functions():
    doc = Document()
    doc.add_paragraph("第一章 项目概述")
    doc.add_paragraph("背景说明" * 10000)
    doc.add_paragraph("第六章 项目建设内容")
    doc.add_paragraph("6.1.应用系统建设")
    doc.add_paragraph("6.1.1.数据平台中枢")
    for i in range(200):
        doc.add_paragraph(f"6.1.1.{i + 1}.数据中台-任务管理-任务功能{i}")
        doc.add_paragraph(f"为业务人员提供任务功能{i}，本期新建。")
    output = BytesIO()
    doc.save(output)
    parsed = parse_document(output.getvalue(), "custom.docx")
    modules = extract_compact_modules(parsed)
    assert len(modules) == 1
    assert "任务功能199" in modules[0].description
    assert "任务功能199" in modules[0].source_location["review_context"]


def test_internal_and_cross_candidates_share_one_review_queue_and_keep_types(monkeypatch):
    from app.modules.duplicate.excel_parser import ParsedFunctionPoint
    from app.schemas import DuplicatePairOut
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    points = [ParsedFunctionPoint(i, f"系统{i}", "统一登录") for i in (1, 2)]
    monkeypatch.setattr(service, "_parse_points_from_file", lambda *args, **kwargs: points)
    calls = []

    def review(candidates, project_id, comparison_type, context_builder, use_llm):
        calls.append(comparison_type)
        assert comparison_type == "mixed"
        assert {p.right.source for p in candidates} == {"current", "history"}
        return [DuplicatePairOut(item_id=c.left.point_id, related_item_id=c.right.point_id,
                                 item_name=c.left.name, related_item_name=c.right.name,
                                 similarity=c.similarity, result_label="重复", severity="risk", reason="独立建设重叠",
                                 comparison_type="cross_report" if c.right.source == "history" else "internal")
                for c in candidates]

    monkeypatch.setattr(service, "_judge_report_pairs", review)
    with Session(engine) as db:
        result = service.run_duplicate_compare_check(db, b"current", "current.csv", [(b"history", "history.csv")],
                                                       None, "项目", None)
    assert calls == ["mixed"]
    assert result.internal_pairs and result.cross_pairs
