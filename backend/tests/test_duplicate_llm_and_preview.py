from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document

from app.api.evaluate import _render_docx_preview
from app.modules.duplicate import llm_extractor, llm_judge, service
from app.modules.duplicate.document_parser import DocumentContent, DocumentSection, parse_document
from app.modules.duplicate.excel_parser import ParsedFunctionPoint


def test_long_word_preview_is_paged_and_locates_later_content():
    document = Document()
    for index in range(190):
        document.add_paragraph(f"第 {index} 段建设内容")
    document.add_paragraph("目标功能点：统一身份认证与权限管理")
    output = BytesIO()
    document.save(output)

    first = _render_docx_preview(output.getvalue(), "报告.docx")
    located = _render_docx_preview(output.getvalue(), "报告.docx", quote="统一身份认证与权限管理")

    assert "第 1 / 3 页" in first
    assert "第 3 / 3 页" in located
    assert "目标功能点：统一身份认证与权限管理" not in first
    assert 'class="source-block source-target"' in located
    assert "目标功能点：统一身份认证与权限管理" in located


def test_word_preview_prefers_occurrence_location_over_repeated_quote():
    document = Document()
    document.add_paragraph("重复功能原文")
    document.add_paragraph("重复功能原文")
    output = BytesIO()
    document.save(output)

    located = _render_docx_preview(
        output.getvalue(), "报告.docx", quote="重复功能原文", paragraph=2, line=2,
    )

    assert '<p class="source-block source-target" data-paragraph="2" data-line="2"' in located
    assert '<p class="source-block" data-paragraph="1" data-line="1"' in located


def test_duplicate_document_uses_llm_even_when_local_sections_exist(monkeypatch):
    point = ParsedFunctionPoint(1, "模型功能点", "描述")
    monkeypatch.setattr(service, "_extract_function_sections", lambda _: [point] * 3)
    monkeypatch.setattr(service, "extract_function_points", lambda *args, **kwargs: [point])
    document = DocumentContent("docx", "报告内容", [DocumentSection("功能", "报告内容")], "报告.docx")

    assert service._parse_document_function_points(document, "项目", use_llm=True) == [point]


def test_duplicate_llm_failure_does_not_return_rule_result(monkeypatch):
    monkeypatch.setattr(llm_judge, "settings", SimpleNamespace(deepseek_api_key="configured"))
    monkeypatch.setattr(llm_judge, "_call_deepseek_api", lambda _: (_ for _ in ()).throw(RuntimeError("模型连接失败")))

    with pytest.raises(ValueError, match="重复建设大模型判定失败"):
        llm_judge.judge_pair_with_llm("A", "描述", "B", "描述", 0.9)


def test_v4_flash_extraction_uses_low_reasoning_and_user_message_only(monkeypatch):
    captured = {}

    class Response:
        status_code = 200

        def raise_for_status(self):
            return None

        def json(self):
            return {"choices": [{"finish_reason": "stop", "message": {"content": '{"function_points": []}'}}]}

    monkeypatch.setattr(llm_extractor, "settings", SimpleNamespace(
        deepseek_api_url="https://example.invalid/v1",
        deepseek_api_key="key",
        deepseek_model="DeepSeek-V4-Flash",
        duplicate_llm_max_tokens=6000,
        duplicate_llm_timeout_seconds=30,
    ))
    monkeypatch.setattr(llm_extractor, "post_direct", lambda url, **kwargs: captured.update(kwargs) or Response())

    assert llm_extractor._call_deepseek_api("extract") == '{"function_points": []}'
    assert captured["json"]["reasoning_effort"] == "low"
    assert [message["role"] for message in captured["json"]["messages"]] == ["user"]


def test_extraction_chunks_stay_within_v4_flash_limit():
    document = DocumentContent(
        "docx",
        "甲" * 9000,
        [DocumentSection("功能章节", "甲" * 9000)],
        "报告.docx",
    )
    chunks = llm_extractor._build_extraction_chunks(document)
    assert len(chunks) >= 3
    assert max(len(text) for _, text in chunks) <= llm_extractor.MAX_SECTION_CHARS


def test_repeated_function_points_are_bound_to_distinct_paragraphs():
    document = Document()
    document.add_paragraph("统一身份认证用于业务系统登录")
    document.add_paragraph("中间说明")
    document.add_paragraph("统一身份认证用于业务系统登录")
    output = BytesIO()
    document.save(output)
    parsed = parse_document(output.getvalue(), "报告.docx")
    points = [
        ParsedFunctionPoint(
            index, "统一身份认证", "用于业务系统登录", source_location={"quote": "统一身份认证用于业务系统登录"},
        )
        for index in (1, 2)
    ]

    located = service._annotate_document_points(points, parsed)

    assert [point.source_location["paragraph"] for point in located] == [1, 3]
    assert [point.source_location["line_start"] for point in located] == [1, 3]


def test_duplicate_highlights_keep_paragraph_number():
    highlights = service._duplicate_pair_highlights(
        {"quote": "功能一", "paragraph": 12, "line_start": 10},
        {"quote": "功能二", "paragraph": 28, "line_start": 21},
    )
    assert [(item["role"], item["paragraph"]) for item in highlights] == [
        ("current", 12), ("related", 28),
    ]


def test_points_from_same_paragraph_are_not_treated_as_duplicate_occurrences():
    assert service._same_source_occurrence(
        {"file_name": "报告.docx", "paragraph": 12},
        {"file_name": "报告.docx", "paragraph": 12},
    )
    assert not service._same_source_occurrence(
        {"file_name": "报告.docx", "paragraph": 12},
        {"file_name": "报告.docx", "paragraph": 28},
    )
