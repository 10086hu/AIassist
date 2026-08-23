from __future__ import annotations

import sys
import time
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session


BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from app.db.models import Base, PriceBenchmark  # noqa: E402
from app.modules.price.parser import ParsedPriceItem, PriceProjectContext, parse_price_document  # noqa: E402
from app.modules.price.service import (  # noqa: E402
    PRICE_REF_001,
    _evaluate_products,
    _evaluate_software,
    import_price_benchmarks,
)
from app.modules.price.standards import evaluate_class2_fees  # noqa: E402


def _item(name: str, amount_wan: float) -> ParsedPriceItem:
    return ParsedPriceItem(item_name=name, quantity=1, total_price=amount_wan * 10_000, evidence=f"{name} {amount_wan}万元")


class Class2FeeStandardTests(unittest.TestCase):
    def test_realistic_municipal_fee_calculation(self) -> None:
        items = [
            _item("咨询费", 75),
            _item("监理费", 28),
            _item("软测费", 45.51),
            _item("安全测评费", 60),
            _item("密测费", 60),
        ]
        context = PriceProjectContext(
            project_type="construction",
            total_investment=5_010.4038 * 10_000,
            direct_construction_cost=4_741.8938 * 10_000,
            software_development_cost=2_957.75 * 10_000,
            hardware_software_purchase_cost=1_447.1438 * 10_000,
            application_development_ratio=0.59,
            security_level=2,
        )
        evaluations, metadata = evaluate_class2_fees(items, context, "市级项目")
        self.assertEqual(metadata["standard_id"], "SHANGHAI_MUNICIPAL_CLASS2_2023")
        self.assertTrue(evaluations)
        self.assertTrue(all(item.status == "pass" for item in evaluations))
        amounts = {item.fee_key: item.standard_amount for item in evaluations}
        self.assertAlmostEqual(amounts["consulting"] or 0, 926_284.07, places=2)
        self.assertAlmostEqual(amounts["supervision"] or 0, 386_673.81, places=2)
        self.assertAlmostEqual(amounts["software_testing"] or 0, 658_550.0, places=2)

    def test_infrastructure_adjustment_and_caps(self) -> None:
        items = [_item("咨询费", 7), _item("软件测试费", 100), _item("安全测评费", 150)]
        context = PriceProjectContext(
            project_type="construction",
            total_investment=40_000 * 10_000,
            direct_construction_cost=40_000 * 10_000,
            software_development_cost=10_000 * 10_000,
            application_development_ratio=0.2,
            security_level=3,
        )
        evaluations, _ = evaluate_class2_fees(items, context, "市级项目")
        amounts = {item.fee_key: item.standard_amount for item in evaluations}
        self.assertAlmostEqual(amounts["consulting"] or 0, 341.5 * 0.8 * 10_000, places=2)
        self.assertEqual(amounts["software_testing"], 100 * 10_000)
        self.assertEqual(amounts["security_assessment"], 150 * 10_000)

    def test_uncovered_supervision_range_requires_manual_review(self) -> None:
        context = PriceProjectContext(project_type="construction", direct_construction_cost=55_000 * 10_000)
        evaluations, _ = evaluate_class2_fees([_item("监理费", 600)], context, "市级项目")
        self.assertEqual(evaluations[0].status, "manual")

    def test_missing_fee_table_is_not_reported_as_pass(self) -> None:
        evaluations, _ = evaluate_class2_fees([], PriceProjectContext(), "市级项目")
        self.assertEqual(evaluations[0].status, "manual")

    def test_fee_names_are_not_overridden_by_remark_text(self) -> None:
        items = [
            ParsedPriceItem(
                item_name="应用软件开发",
                category="投资估算汇总",
                total_price=600 * 10_000,
                evidence="应用软件开发 | 600 | 软件测试费计费基数",
            ),
            ParsedPriceItem(item_name="软件测试费", total_price=100 * 10_000, evidence="软件测试费 | 100"),
            ParsedPriceItem(item_name="安全测评费", total_price=150 * 10_000, evidence="安全测评费 | 150"),
            ParsedPriceItem(
                item_name="等级保护测评费",
                total_price=150 * 10_000,
                evidence="等级保护测评费 | 150 | 与安全测评费不得重复申报",
            ),
        ]
        context = PriceProjectContext(
            project_type="construction",
            total_investment=2_000 * 10_000,
            direct_construction_cost=1_000 * 10_000,
            software_development_cost=600 * 10_000,
            security_level=2,
        )
        evaluations, _ = evaluate_class2_fees(items, context, "市级项目")
        by_key = {evaluation.fee_key: evaluation for evaluation in evaluations}
        self.assertEqual(by_key["software_testing"].source_item.item_name, "软件测试费")
        self.assertEqual(by_key["grade_assessment"].source_item.item_name, "等级保护测评费")
        self.assertIn("security_grade_mutual_exclusion", by_key)


class PriceParserTests(unittest.TestCase):
    def test_summary_and_price_tables_are_extracted(self) -> None:
        document = Document()
        document.add_paragraph("本项目为新建项目，按等保二级建设。")
        summary = document.add_table(rows=1, cols=3)
        summary.cell(0, 0).text = "投资估算总表（单位：万元）"
        for values in (
            ("一、系统建设费", "1000", ""),
            ("1.1应用软件开发", "200", ""),
            ("1.2硬件购置", "300", ""),
            ("1.3产品软件", "100", ""),
            ("二、其他费用", "50", ""),
            ("1", "咨询费", "20"),
            ("总计", "1050", ""),
        ):
            cells = summary.add_row().cells
            for cell, value in zip(cells, values):
                cell.text = value
        irrelevant = document.add_table(rows=1, cols=3)
        for cell, value in zip(irrelevant.rows[0].cells, ("序号", "章节名称", "说明")):
            cell.text = value
        for index in range(500):
            cells = irrelevant.add_row().cells
            cells[0].text = str(index)
            cells[1].text = f"章节{index}"
            cells[2].text = "普通内容"
        product = document.add_table(rows=1, cols=7)
        for cell, value in zip(product.rows[0].cells, ("序号", "产品大类", "产品小类", "产品配置", "单价（元）", "数量", "总金额（元）")):
            cell.text = value
        row = product.add_row().cells
        for cell, value in zip(row, ("1", "硬件产品", "应用服务器", "32核", "10000", "2", "20000")):
            cell.text = value
        output = BytesIO()
        document.save(output)

        started = time.perf_counter()
        parsed = parse_price_document(output.getvalue(), "sample.docx")
        self.assertLess(time.perf_counter() - started, 3)
        self.assertEqual(parsed.context.total_investment, 10_500_000)
        self.assertEqual(parsed.context.direct_construction_cost, 10_000_000)
        self.assertEqual(parsed.context.software_development_cost, 2_000_000)
        self.assertTrue(any(item.item_name == "应用服务器" for item in parsed.items))
        self.assertFalse(any(item.item_name.startswith("章节") for item in parsed.items))


class PriceReferenceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.engine = create_engine("sqlite+pysqlite:///:memory:", future=True)
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)

    def tearDown(self) -> None:
        self.db.close()
        self.engine.dispose()

    def test_software_total_uses_effort_person_month(self) -> None:
        item = ParsedPriceItem(
            item_name="普通软件功能点",
            effort_person_month=2,
            unit_price=25_000,
            total_price=60_000,
            evidence="普通软件功能点 2人月 25000元 60000元",
        )
        result = _evaluate_software([item])
        self.assertTrue(any(issue["rule_id"] == PRICE_REF_001["rule_id"] and "50000" in issue["reason"] for issue in result["issues"]))

    def test_software_check_ignores_totals_and_investment_summary(self) -> None:
        items = [
            ParsedPriceItem(
                item_name="5",
                effort_person_month=150_000,
                evidence="合计 | 5 | 150000",
            ),
            ParsedPriceItem(
                item_name="1.1应用软件开发",
                category="投资估算汇总",
                total_price=2_957.75 * 10_000,
                evidence="1.1应用软件开发 | 2957.75",
            ),
        ]
        result = _evaluate_software(items)
        self.assertEqual(result["issues"], [])

    def test_missing_unit_prices_are_grouped_by_table(self) -> None:
        items = [
            ParsedPriceItem(item_name=f"功能{i}", effort_person_month=2, source_section="表14", source_row=i, evidence=f"功能{i} | 人月 | 2")
            for i in range(1, 11)
        ]
        result = _evaluate_software(items)
        self.assertEqual(len(result["issues"]), 1)
        self.assertEqual(result["issues"][0]["affected_item_count"], 10)
        self.assertEqual(result["issues"][0]["result_label"], "表格缺少单价")

    def test_new_technology_adjustment_is_optional(self) -> None:
        items = [
            ParsedPriceItem(item_name="隐私计算接口", effort_person_month=2, unit_price=25_000, total_price=50_000, evidence="隐私计算接口 | 2 | 25000 | 50000"),
            ParsedPriceItem(item_name="智能问答知识库", effort_person_month=2, unit_price=30_000, total_price=60_000, evidence="智能问答知识库 | 2 | 30000 | 60000"),
        ]
        self.assertEqual(_evaluate_software(items)["issues"], [])

    def test_unexplained_adjusted_prices_are_grouped_by_table(self) -> None:
        items = [
            ParsedPriceItem(item_name=f"普通功能{i}", effort_person_month=2, unit_price=30_000, total_price=60_000, source_section="表8", source_row=i, evidence=f"普通功能{i} | 2 | 30000 | 60000")
            for i in range(1, 6)
        ]
        result = _evaluate_software(items)
        self.assertEqual(len(result["issues"]), 1)
        self.assertEqual(result["issues"][0]["affected_item_count"], 5)
        self.assertEqual(result["issues"][0]["result_label"], "调整依据待确认")

    def test_adjusted_price_uses_clear_table_level_new_technology_context(self) -> None:
        items = [
            ParsedPriceItem(item_name="智能问答知识库", effort_person_month=2, unit_price=30_000, total_price=60_000, source_section="表9", source_row=1, evidence="智能问答知识库 | 2 | 30000 | 60000"),
            ParsedPriceItem(item_name="模型训练服务", effort_person_month=2, unit_price=30_000, total_price=60_000, source_section="表9", source_row=2, evidence="模型训练服务 | 2 | 30000 | 60000"),
            ParsedPriceItem(item_name="核心字段定义库", effort_person_month=2, unit_price=30_000, total_price=60_000, source_section="表9", source_row=3, evidence="核心字段定义库 | 2 | 30000 | 60000"),
            ParsedPriceItem(item_name="规则维护", effort_person_month=2, unit_price=30_000, total_price=60_000, source_section="表9", source_row=4, evidence="规则维护 | 2 | 30000 | 60000"),
        ]
        self.assertEqual(_evaluate_software(items)["issues"], [])

    def test_price_above_new_technology_ceiling_is_reported(self) -> None:
        item = ParsedPriceItem(item_name="大模型功能", effort_person_month=2, unit_price=32_000, total_price=64_000, evidence="大模型功能 | 2 | 32000 | 64000")
        result = _evaluate_software([item])
        self.assertEqual(len(result["issues"]), 1)
        self.assertIn("超过", result["issues"][0]["reason"])

    def test_empty_benchmark_library_returns_status_notice(self) -> None:
        item = ParsedPriceItem(item_name="应用服务器", spec="32核", quantity=2, unit_price=15_000)
        result = _evaluate_products(self.db, [item])
        self.assertEqual(result["benchmark_count"], 0)
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["library_status"], "empty")
        self.assertEqual(result["notices"][0]["code"], "PRICE_LIBRARY_EMPTY")

    def test_benchmark_import_updates_existing_record(self) -> None:
        first = [{"item_name": "应用服务器", "brand": "厂商A", "model": "S1", "unit_price": 10_000, "source": "询价"}]
        second = [{"item_name": "应用服务器", "brand": "厂商A", "model": "S1", "unit_price": 12_000, "source": "询价"}]
        self.assertEqual(import_price_benchmarks(self.db, first)["imported"], 1)
        result = import_price_benchmarks(self.db, second)
        self.assertEqual(result["updated"], 1)
        self.assertEqual(self.db.query(PriceBenchmark).count(), 1)
        self.assertEqual(self.db.query(PriceBenchmark).one().unit_price, 12_000)


if __name__ == "__main__":
    unittest.main()
