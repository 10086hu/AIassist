from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.modules.resource.quantity_parser import DocumentIndex, build_evidence_location
from app.modules.data_rules.document_rules import _retain_data_evidence_locations
from app.api.evaluate import _normalize_document_rule_result, _attach_finding_locations


def content(document):
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def result(evidence, rule='18', section='', module='data_reasonableness'):
    finding = dict(rule_id=rule, review_opinion='原判断不变', risk_level='需人工确认',
                   source_section=section, evidence_examples=[evidence], revision_advice='核对原文')
    return {'module_code': module, 'findings': [finding], 'summary': {'total_findings': 1}}


def locate(document, evidence, **kwargs):
    index = DocumentIndex(content(document), '报告.docx')
    return build_evidence_location(index, [evidence], kwargs.get('section', ''))


def test_four_level_path_and_paragraph_quote():
    d = Document()
    for level, title in enumerate(['建设方案', '资源申请', '计算资源', '数据库配置'], 1):
        d.add_heading(title, level)
    d.add_paragraph('本项目数据库服务器数量为四台。')
    f = locate(d, '本项目数据库服务器数量为四台。')
    loc = f['source_locations'][0]
    assert loc['heading_path'] == ['建设方案', '资源申请', '计算资源', '数据库配置']
    assert loc['paragraph'] == 5
    assert '数据库配置' in f['location_hint']
    assert 'line_start' not in loc and 'page' not in loc
    assert f['source_highlights'][0]['quote'] == '本项目数据库服务器数量为四台。'


def test_numbered_five_level_titles_without_styles():
    d = Document()
    for title in ['第六章 建设方案', '6.2 资源申请', '6.2.1 计算资源', '6.2.1.3 数据库配置', '6.2.1.3.1 授权安排']:
        d.add_paragraph(title)
    d.add_paragraph('数据库授权覆盖所有计算节点。')
    assert len(locate(d, '数据库授权覆盖所有计算节点。')['source_locations'][0]['heading_path']) == 5


def test_numbered_sibling_does_not_inherit_intervening_sublist():
    d = Document()
    for title in ['第六章 建设方案', '6.1 系统功能', '6.1.4 应用系统', '6.1.4.34 业务处理',
                  '十一、其他指标', '3.质量目标', '6.1.4.35 数据核验']:
        d.add_paragraph(title)
    d.add_paragraph('数据核验系统服务数量为两套。')
    path = locate(d, '数据核验系统服务数量为两套。')['source_locations'][0]['heading_path']
    assert path == ['第六章 建设方案', '6.1 系统功能', '6.1.4 应用系统', '6.1.4.35 数据核验']


def test_acceptance_requirement_uses_actual_heading_as_scope():
    d = Document()
    d.add_paragraph('4.2.2 验收标准、具体指标和内容')
    f = locate(d, '验收要求', rule='24')
    assert f['source_location']['match_status'] == 'scope'
    assert f['source_highlights'][0]['quote'] == '4.2.2 验收标准、具体指标和内容'


def test_toc_does_not_steal_body_heading_or_quote():
    d = Document()
    d.add_paragraph('6.3.2 数据治理内容........12')
    toc = d.add_paragraph('本项目的迁移单位为人月。')
    toc.style = d.styles.add_style('TOC 1', 1)
    d.add_paragraph('6.3.2 数据治理内容')
    d.add_paragraph('本项目的迁移单位为人月。')
    f = locate(d, '本项目的迁移单位为人月。', rule='24')
    assert len(f['source_locations']) == 1
    assert f['source_locations'][0]['paragraph'] == 4


def test_duplicate_text_is_disambiguated_by_section():
    d = Document()
    for title in ['一期建设', '二期建设']:
        d.add_heading(title, 1)
        d.add_paragraph('项目预算总金额为100万元。')
    f = locate(d, '项目预算总金额为100万元。', section='二期建设')
    assert len(f['source_locations']) == 1
    assert f['source_locations'][0]['paragraph'] == 4


def test_duplicate_text_without_scope_is_explicitly_ambiguous():
    d = Document()
    for title in ['一期建设', '二期建设']:
        d.add_heading(title, 1)
        d.add_paragraph('项目预算总金额为100万元。')
    f = locate(d, '项目预算总金额为100万元。')
    assert len(f['source_locations']) == 2
    assert f['source_location']['match_status'] == 'ambiguous'
    assert '候选位置' in f['location_hint']


def test_missing_sentence_is_not_invented_as_a_highlight():
    d = Document()
    d.add_paragraph('本报告不包含所要求的额外说明。')
    f = locate(d, '需补充数据归集和质量自评估报告', rule='24')
    assert f['source_highlights'] == []
    assert f['source_location']['match_status'] == 'unresolved'


def test_missing_field_points_to_checked_table_scope():
    d = Document()
    d.add_heading('数据治理服务', 2)
    d.add_paragraph('表6.8 数据服务事项清单')
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = '服务事项'
    t.cell(0, 1).text = '测算单位'
    f = locate(d, '表6.8 数据服务事项清单', rule='24')
    assert f['revision_target']['target_type'] == 'section_or_document'
    assert f['source_locations'][0]['table_index'] == 1
    assert f['source_highlights'][0]['quote'] == '服务事项 | 测算单位'


def test_budget_discrepancy_retains_prose_and_table():
    d = Document()
    d.add_heading('投资概算', 1)
    d.add_paragraph('应用软件开发费用为100万元。')
    t = d.add_table(rows=2, cols=3)
    for c, value in zip(t.rows[1].cells, ['一', '软件开发费用', '104']):
        c.text = value
    f = locate(d, '应用软件开发费用为100万元。 | DOCX表格1行2：一 | 软件开发费用 | 104')
    assert len(f['source_locations']) == 2
    assert {l['precision'] for l in f['source_locations']} == {'paragraph', 'table_row'}


def test_same_row_number_in_other_table_is_not_selected():
    d = Document()
    for title, value in [('投资概算', '软件开发费用'), ('其他内容', '设备采购费用')]:
        d.add_heading(title, 1)
        table = d.add_table(rows=2, cols=2)
        table.cell(1, 0).text = value
        table.cell(1, 1).text = '104'
    f = locate(d, 'DOCX表格1行2：软件开发费用 | 104')
    assert len(f['source_locations']) == 1
    assert f['source_locations'][0]['table_index'] == 1


def test_vertical_and_horizontal_merge_keep_real_cell_coordinates():
    d = Document()
    t = d.add_table(rows=3, cols=3)
    t.cell(0, 0).merge(t.cell(0, 1)).text = '服务说明'
    t.cell(1, 0).merge(t.cell(2, 0)).text = '历史数据迁移'
    t.cell(2, 1).text = '人月'
    t.cell(2, 2).text = '20'
    f = locate(d, 'DOCX表格1行3：历史数据迁移 | 人月 | 20')
    loc = f['source_locations'][0]
    assert loc['row_index'] == 3
    assert loc['cells'][0]['origin_row'] == 2
    assert loc['cells'][1]['column_start'] == 2
    index = DocumentIndex(content(d), '报告.docx')
    first = next(b for b in index.blocks if b.get('row_index') == 1)
    assert first['cells'][0]['column_end'] == 2


def test_nested_table_gets_distinct_path():
    d = Document()
    outer = d.add_table(rows=1, cols=1)
    nested = outer.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = '嵌套服务测算数量20人月'
    f = locate(d, '嵌套服务测算数量20人月')
    assert len(f['source_locations']) == 1
    assert f['source_locations'][0]['table_path'] == '1/1/1'


def test_expanded_merged_total_maps_back_to_physical_cells():
    d = Document()
    table = d.add_table(rows=1, cols=6)
    table.cell(0, 0).text = '6'
    table.cell(0, 1).text = '总计'
    table.cell(0, 4).merge(table.cell(0, 5)).text = '55.746'
    f = locate(d, 'DOCX表格1行1：6 | 总计 | 55.746 | 55.746')
    assert len(f['source_locations']) == 1
    assert f['source_highlights'][0]['quote'] == '6 | 总计 |  |  | 55.746'
    assert f['source_locations'][0]['cells'][-1]['column_end'] == 6


def test_rule_evidence_can_omit_inherited_vertical_cells():
    d = Document()
    table = d.add_table(rows=2, cols=4)
    table.cell(0, 0).merge(table.cell(1, 0)).text = '应用软件开发'
    for c, value in zip(table.rows[1].cells[1:], ['权限控制', '0.5', '2']):
        c.text = value
    f = locate(d, 'DOCX表格1行2：权限控制 | 0.5 | 2')
    assert f['source_locations'][0]['row_index'] == 2
    assert f['source_locations'][0]['cells'][0]['origin_row'] == 1


def test_long_spec_does_not_hide_quantity_cell_from_highlight():
    d = Document()
    table = d.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ['应用服务器', '设备技术说明' * 120, '31']):
        cell.text = text
    f = locate(d, 'DOCX表格1行1：应用服务器 | '+'设备技术说明' * 120+' | 31')
    assert any(h['quote'] == '31' and h['column_start'] == 3 and h['precision'] == 'cell'
               for h in f['source_highlights'])


def test_heading_in_composite_evidence_limits_repeated_prose():
    d = Document()
    repeated = '通过本项目的建设，能够有效减少企业和市民的办事成本，缩短审批周期。'
    d.add_heading('项目概述', 1)
    d.add_paragraph(repeated)
    d.add_paragraph('克拉玛依市数据资源和政务服务中心')
    d.add_heading('经济效益分析', 1)
    d.add_paragraph('本项目提高克拉玛依市数据资源和政务服务中心办事效率。')
    d.add_paragraph(repeated)
    f = locate(d, '经济效益分析 | 本项目提高克拉玛依市数据资源和政务服务中心办事效率。 | '+repeated,
               rule='21', section='项目评价/成效指标')
    assert [loc['paragraph'] for loc in f['source_locations']] == [5, 6]


@pytest.mark.parametrize('module,rule', [('function_correspondence', '18'), ('data_reasonableness', 'DATA_REASON_003'), ('resource', 'MAINT_OPS_ALL'), ('price', '24')])
def test_other_modules_and_rules_unchanged(module, rule):
    response = result('正文证据', module=module, rule=rule)
    _attach_finding_locations(response, b'invalid', '报告.docx')
    assert response['findings'][0]['source_location']['precision'] not in {'document_blocks', 'scope'}
    assert response['findings'][0]['review_opinion'] == '原判断不变'


def test_non_docx_keeps_existing_locations():
    response = {'issues': [{'message': '原判断', 'evidence': '正文证据'}]}
    original = deepcopy(response)
    _retain_data_evidence_locations(response, b'pdf', '报告.pdf')
    assert response == original


def test_only_location_fields_change_and_nested_rule_results_get_same_locations():
    d = Document()
    d.add_paragraph('本项目的软件开发费为104万元。')
    response = result('本项目的软件开发费为104万元。')
    response['findings'][0].update(build_evidence_location(DocumentIndex(content(d), '报告.docx'), ['本项目的软件开发费为104万元。'], advice='核对原文'))
    original = deepcopy(response)
    response['rule_results'] = [{'rule_id': '18', 'findings': [deepcopy(response['findings'][0])]}]
    _attach_finding_locations(response, content(d), '报告.docx')
    assert response['summary'] == original['summary']
    for key, value in original['findings'][0].items():
        assert response['findings'][0][key] == value
    assert response['findings'][0] == response['rule_results'][0]['findings'][0]


def test_paragraph_outline_overrides_heading_style():
    d = Document()
    p = d.add_heading('作为正文引用的标题', 4)
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), '9')
    p._p.get_or_add_pPr().append(outline)
    index = DocumentIndex(content(d), '报告.docx')
    assert index.blocks[0]['heading_path'] == []


def test_quoted_regulation_chapters_do_not_replace_report_outline():
    d = Document()
    d.add_heading('第一章 单位概况', 1)
    d.add_heading('第二章 项目概述', 1)
    d.add_paragraph('第二章资产评估')
    d.add_heading('2.6 项目成效指标', 2)
    d.add_paragraph('智能搜索响应时间不超过五秒。')
    f = locate(d, '智能搜索响应时间不超过五秒。', rule='21')
    assert f['source_locations'][0]['heading_path'] == ['第二章 项目概述', '2.6 项目成效指标']


def test_outline_parent_is_preserved_when_printed_numbering_is_inconsistent():
    d = Document()
    d.add_heading('第六章 建设内容', 1)
    d.add_heading('6.3 数据治理内容', 2)
    d.add_heading('6.4.2 数据服务', 3)
    d.add_heading('6.4.2.1 数据服务清单', 4)
    d.add_paragraph('本项目申报数据治理服务。')
    f = locate(d, '本项目申报数据治理服务。', rule='24')
    assert f['source_locations'][0]['heading_path'][1] == '6.3 数据治理内容'


def test_full_raw_evidence_survives_display_excerpt_truncation():
    d = Document()
    table = d.add_table(rows=3, cols=2)
    for row, name in zip(table.rows, ['历史数据迁移', '数据融合', '数据购买']):
        row.cells[0].text, row.cells[1].text = name, '人月'
    evidence = '；'.join(f'DOCX表格1行{i}：{name} | 人月' for i, name in enumerate(['历史数据迁移', '数据融合', '数据购买'], 1))
    raw = {'rule_excel_row': 24, 'passed': False, 'issues': [{'message': '原判断不变', 'evidence': evidence}]}
    raw = _retain_data_evidence_locations(raw, content(d), '报告.docx')
    response = _normalize_document_rule_result({'results': [raw]}, 'data_reasonableness', False)
    _attach_finding_locations(response, content(d), '报告.docx')
    assert [l['row_index'] for l in response['findings'][0]['source_locations']] == [1, 2, 3]


def test_resource_rule_retains_locations_before_api_formatting():
    from app.modules.resource.document_check import check_resource_document
    d = Document()
    d.add_heading('6.7 PaaS服务', 2)
    d.add_paragraph('本项目不涉及PaaS服务。')
    findings = check_resource_document(content(d), '报告.docx')
    finding = next(f for f in findings if f.rule_code == 'R15_SECURITY_PAAS_CRYPTO_QUANTITY')
    assert finding.evidence_location['source_locations'][0]['paragraph'] == 2
    assert finding.evidence_location['source_highlights'][0]['quote'] == '本项目不涉及PaaS服务。'
    response = result('本项目不涉及PaaS服务。', rule=finding.rule_code, module='resource')
    response['findings'][0].update(finding.evidence_location)
    _attach_finding_locations(response, content(d), '报告.docx')
    assert response['findings'][0]['revision_target']['advice'] == '核对原文'


def test_budget_rule_retains_locations_before_api_formatting():
    from app.modules.data_rules.service import run_data_rules_check_from_document
    d = Document()
    d.add_heading('7.2 投资估算', 2)
    table = d.add_table(rows=3, cols=4)
    values = [['项目名称', '数量', '单价（万元）', '合价（万元）'],
              ['软件开发', '2', '3', '9'], ['合计', '', '', '9']]
    for row, texts in zip(table.rows, values):
        for cell, text in zip(row.cells, texts):
            cell.text = text
    raw = run_data_rules_check_from_document(content(d), '报告.docx', selected_rule_ids=['18'])
    rule = raw['results'][0]
    assert rule['issues']
    assert len(rule['evidence_locations']) == len(rule['issues'])
    assert any(loc.get('table_index') == 1 and loc.get('row_index') == 2
               for evidence in rule['evidence_locations'] for loc in evidence['source_locations'])
