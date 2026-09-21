from io import BytesIO
import json
from uuid import uuid4

from docx import Document
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api import completeness
from app.api.projects import _project_tree_to_dict
from app.db.models import Base, Project, CheckResult, SourceArtifact
from app.db.session import get_db
from app.modules.completeness.service import check_completeness


SECTIONS = [
    ("立项依据与背景", "本项目根据政务服务改革要求建设，解决群众跨部门办事的材料重复提交问题。"),
    ("现状与需求分析", "现有业务系统分散，群众需重复填写材料。需要统一申报、材料复用和跨部门审批。"),
    ("技术方案", "本次开发统一申报、数据交换及业务审批三个模块，为群众提供在线办理服务。"),
    ("投资估算", "本项目总投资120万元，其中软件开发费100万元，实施服务费20万元。"),
]


def document(sections=SECTIONS, extra=(), table=False):
    doc = Document()
    for title, body in sections:
        doc.add_heading(title, level=1)
        if body is not None:
            doc.add_paragraph(body)
    for text in extra:
        doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=1, cols=3)
        for cell, text in zip(t.rows[0].cells, ("项目", "金额（万元）", "说明")):
            cell.text = text
        for row in (("应用开发", "100", "含接口开发"), ("实施服务", "20", "部署培训")):
            for cell, text in zip(t.add_row().cells, row):
                cell.text = text
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def result(raw=None, **kwargs):
    return check_completeness(raw or document(), "报告.docx", **kwargs)


def group(r, key):
    return next(i for i in r["details"] + r["items"] if i["key"] == key)


def test_narrative_equivalence_no_mandatory_tables():
    assert result()["status"] == "pass"


@pytest.mark.parametrize("missing", [2, 3])
def test_major_content_absent(missing):
    r = result(document([s for n, s in enumerate(SECTIONS) if n != missing]))
    assert r["status"] == "missing"
    assert r["missing_count"] == 1


def test_title_only_is_not_a_section():
    sections = [s if n != 2 else (s[0], None) for n, s in enumerate(SECTIONS)]
    assert group(result(document(sections)), "construction")["status"] == "missing"


@pytest.mark.parametrize("placeholder", ["待补充", "略", "请在此填写建设方案", "详见附件", "TBD"])
def test_template_prompts_do_not_prove_presence(placeholder):
    sections = [s if n != 2 else (s[0], placeholder) for n, s in enumerate(SECTIONS)]
    assert group(result(document(sections)), "construction")["status"] == ("review" if placeholder == "详见附件" else "missing")


def test_toc_only():
    doc = Document()
    for title, _ in SECTIONS:
        doc.add_paragraph(title + "........12")
    out = BytesIO(); doc.save(out)
    r = result(out.getvalue())
    assert r["status"] == "missing"
    assert len(r["items"]) == 4


def test_arithmetic_error_does_not_fail_precheck():
    sections = SECTIONS[:-1] + [("投资估算", "本项目总投资999万元，其中开发费100万元，实施服务费20万元。")]
    assert result(document(sections))["status"] == "pass"


def test_detail_table_can_replace_prose():
    assert result(document(SECTIONS[:-1] + [("投资估算", "预算明细如下。")], table=True))["status"] == "pass"


def test_total_without_major_details():
    r = result(document(SECTIONS[:-1] + [("投资估算", "本项目总投资120万元。")]))
    assert r["status"] == "pass"


def test_budget_business_numbers_are_not_costs():
    sections = SECTIONS[:-1] + [("预算管理", "本系统预算编制每月300次，每年存储120MB，预计覆盖2025年至2027年。")]
    r = result(document(sections))
    assert group(r, "budget")["status"] == "missing"


def test_agency_name_not_needs():
    r = result(document([SECTIONS[0], ("业务需求单位概况", "中华人民共和国上海海关。"), *SECTIONS[2:]]))
    assert r["status"] == "pass"
    assert not any(i["key"] == "needs" for i in r["details"])


def test_policy_implementation_plan_not_project_construction():
    r = result(document([SECTIONS[0], SECTIONS[1], ("编制依据", "《国务院关于政务服务整合实施方案的通知》（2025年1号）"), SECTIONS[3]]))
    assert group(r, "construction")["status"] == "missing"


def test_attachment_reference_is_not_automatically_missing():
    r = result(document(extra=["相关政策规定了业务审批流程，详见附件三。"])); assert r["status"] == "pass"
    assert r["info_count"] == 1
    assert r["missing_count"] == 0


def test_explicit_required_missing_attachment():
    r = result(document(extra=["本报告必须提交完整的批复材料，详见附件三。"])); assert r["status"] == "missing"


def test_repeated_reference_cannot_resolve_itself():
    r = result(document(extra=["相关政策规定了审批流程，详见附件三。", "相关审批要求和条件详见附件三。"])); assert r["review_count"] == 0 and r["info_count"] == 1


def test_supplied_numbered_attachment_resolves():
    raw = document(extra=["本报告必须提交完整的批复材料，详见附件三。"])
    assert result(raw, attachments=[("附件三.docx", document())])["status"] == "pass"


def test_attachment_thirty_not_three():
    raw = document(extra=["本报告必须提交完整的批复材料，详见附件三。"])
    assert result(raw, attachments=[("附件三十.docx", document())])["status"] == "missing"


def test_material_inside_attachment_counts():
    raw = document(SECTIONS[:-1])
    assert result(raw, attachments=[("投资明细.docx", document([SECTIONS[3]]))])["status"] == "pass"


def test_unreadable_attachment_is_review_not_missing():
    r = result(attachments=[("说明.pdf", b"not a pdf")]); assert r["status"] == "pass" and r["info_count"] == 1


@pytest.mark.parametrize("name,raw", [("坏文件.docx", b"broken"), ("旧报告.doc", b"legacy"), ("空.txt", b"")])
def test_unreadable_or_empty_files(name, raw):
    r = check_completeness(raw, name)
    assert r["status"] in {"unreadable", "missing"}
    assert r["status"] != "pass"


def test_no_agency_or_filename_whitelist():
    raw = document(SECTIONS[:-1])
    a = check_completeness(raw, "上海海关.docx")
    b = check_completeness(raw, "外地某部门.docx")
    assert a["status"] == b["status"] == "missing"


def test_chinese_amounts_and_merged_sections():
    r = result(document(SECTIONS[:-1] + [("投资估算", "本项目总投资一百二十万元，其中软件开发费一百万元，实施服务费二十万元。")]))
    assert r["status"] == "pass"


def test_txt_heading_context():
    raw = "\n".join(title + "\n" + body for title, body in SECTIONS).encode()
    assert check_completeness(raw, "报告.txt")["status"] == "pass"


def test_empty_function_table_is_not_construction():
    doc = Document(BytesIO(document([SECTIONS[0], SECTIONS[1], ("功能设计", None), SECTIONS[3]])))
    # A header inside a genuine function-design section has no functional content.
    doc.add_heading("功能设计", level=1)
    t = doc.add_table(rows=2, cols=3)
    for cell, label in zip(t.rows[0].cells, ("序号", "模块名称", "功能描述")): cell.text = label
    out = BytesIO(); doc.save(out)
    assert group(result(out.getvalue()), "construction")["status"] == "missing"


def test_generic_principles_are_not_construction():
    raw = document([SECTIONS[0], SECTIONS[1], ("总体方案", "本项目遵循先进适用、安全可靠、统筹规划的原则。"), SECTIONS[3]])
    assert group(result(raw), "construction")["status"] == "missing"


def test_currency_unit_header_is_not_an_amount():
    from app.modules.completeness.service import MONEY
    assert not MONEY.search("单价（万元） | 小计（万元）")
    assert MONEY.search("一百万元")


def test_policy_cost_threshold_is_not_project_budget_detail():
    doc = Document(BytesIO(document(SECTIONS[:-1] + [("投资估算", "本项目总投资120万元。")])) )
    doc.add_heading("建设依据", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0,0).text = "估算模块"
    table.cell(0,1).text = "根据资产评估政策，账面原值低于500万元的资产可以不进行评估，100万元以下另有要求。"
    out = BytesIO(); doc.save(out)
    assert result(out.getvalue())["status"] == "pass"


@pytest.fixture
def api():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine)
    with sessions() as db:
        db.add_all([Project(id="p", name="测试项目"), Project(id="q", name="另一个项目")]); db.commit()
    app = FastAPI(); app.include_router(completeness.router, prefix="/c")
    def override():
        with sessions() as db:
            yield db
    app.dependency_overrides[get_db] = override
    with TestClient(app) as client:
        yield client, sessions
    engine.dispose()


def post(client, request_id=None, raw=None):
    return client.post("/c/reports", data={"project_id": "p", "request_id": request_id or str(uuid4())}, files={"file": ("报告.docx", raw or document())})


def test_history_snapshot_export_and_project_isolation(api):
    client, sessions = api
    r = post(client); assert r.status_code == 200, r.text
    report = r.json(); ident = report["id"]
    assert report["status"] == "pass"
    assert client.get("/c/reports?project_id=p").json()["total"] == 1
    assert client.get("/c/reports?project_id=q").json()["total"] == 0
    assert client.get(f"/c/reports/{ident}?project_id=q").status_code == 404
    assert client.get(f"/c/reports/{ident}/sources/0?project_id=p").content[:2] == b"PK"
    assert client.get(f"/c/reports/{ident}/sources/0?project_id=q").status_code == 404
    exported = client.get(f"/c/reports/{ident}/export?project_id=p")
    assert exported.status_code == 200 and "完整性检查报告" in exported.text
    # Reading history never re-runs the checker or rewrites the result.
    assert client.get(f"/c/reports/{ident}?project_id=p").json() == report
    with sessions() as db:
        assert db.query(CheckResult).count() == 1
        assert db.query(SourceArtifact).count() == 1
        assert _project_tree_to_dict(db.get(Project, "p"))["check_runs"] == []


def test_retry_is_one_report_new_run_creates_second(api):
    client, _ = api; token = str(uuid4()); raw = document()
    a = post(client, token, raw); b = post(client, token, raw)
    assert a.json()["id"] == b.json()["id"]
    assert post(client).json()["id"] != a.json()["id"]
    assert client.get("/c/reports?project_id=p&limit=1").json()["total"] == 2
    assert len(client.get("/c/reports?project_id=p&offset=1").json()["items"]) == 1


def test_failed_parse_is_saved_as_report(api):
    client, _ = api
    r = post(client, raw=b"broken"); assert r.json()["status"] == "unreadable"
    assert client.get("/c/reports?project_id=p").json()["total"] == 1


def test_export_escapes_untrusted_content():
    r = result(); r["filename"] = '<script>alert("x")</script>'
    assert "<script>" not in completeness.render_html(r)


def test_scan_cannot_prove_absence():
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>", b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Resources << >> >>"]
    raw = b"%PDF-1.4\n"; offsets = [0]
    for n, obj in enumerate(objects, 1):
        offsets.append(len(raw)); raw += str(n).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    start = len(raw); raw += b"xref\n0 4\n0000000000 65535 f \n"
    for pos in offsets[1:]: raw += f"{pos:010d} 00000 n \n".encode()
    raw += f"trailer\n<< /Size 4 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF".encode()
    r = check_completeness(raw, "扫描.pdf")
    assert r["status"] == "review"
    assert r["missing_count"] == 0


@pytest.mark.parametrize('omitted', [0, 1])
def test_background_and_needs_not_independent_gates(omitted):
    r = result(document([s for n, s in enumerate(SECTIONS) if n != omitted]))
    assert r['status'] == 'pass'
    assert r['missing_count'] == 0


def test_specialist_topics_not_required_by_precheck():
    r = result(document(SECTIONS[2:]))
    assert r['status'] == 'pass'
    assert set(i['key'] for i in r['details']) == {'construction', 'budget'}


def test_budget_narrative_without_amount_is_material():
    r = result(document(SECTIONS[:-1] + [('投资估算', '本项目经费纳入已有财政预算统筹安排，不新增投资。')]))
    assert r['status'] == 'pass'


def test_one_budget_row_is_sufficient_for_precheck():
    doc = Document(BytesIO(document(SECTIONS[:-1] + [('投资估算', '预算如下。')])))
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = '项目', '金额（万元）'
    t.cell(1, 0).text, t.cell(1, 1).text = '系统建设', '120'
    out = BytesIO(); doc.save(out)
    assert result(out.getvalue())['status'] == 'pass'


def test_whole_budget_template_empty_is_detected_once():
    doc = Document(BytesIO(document(SECTIONS[:-1] + [('投资估算', '预算明细如下。')])))
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = '项目', '金额（万元）'
    out = BytesIO(); doc.save(out)
    r = result(out.getvalue())
    assert r['status'] == 'missing' and r['missing_count'] == 1


def test_core_numbered_attachment_missing_despite_summary():
    r = result(document(extra=['本项目主要预算明细详见附件三。']))
    assert r['status'] == 'missing' and r['missing_count'] == 1


def test_core_named_attachment_supplied():
    raw = document(SECTIONS[:-1] + [('投资估算', '本项目预算详见附件《投资估算表》。')])
    missing = result(raw)
    assert missing['status'] == 'missing' and missing['missing_count'] == 1
    supplied = result(raw, attachments=[('投资估算表.docx', document([SECTIONS[3]]))])
    assert supplied['status'] == 'pass'


def test_policy_reference_does_not_gate_whole_report():
    r = result(document(extra=['相关政策依据：《监管工作指引》规定了联合监管流程，内部文件未公开，详见附件三。']))
    assert r['status'] == 'pass' and r['info_count'] == 1 and r['review_count'] == 0


def test_unreadable_core_attachment_review_not_missing():
    raw = document(SECTIONS[:-1] + [('投资估算', '本项目预算详见附件三。')])
    r = result(raw, attachments=[('附件三.pdf', b'broken')])
    assert r['status'] == 'review' and r['missing_count'] == 0 and r['review_count'] == 1


def test_unrelated_bad_attachment_does_not_excuse_missing_budget():
    r = result(document(SECTIONS[:-1]), attachments=[('宣传照片.pdf', b'broken')])
    assert group(r, 'budget')['status'] == 'missing'
    assert r['info_count'] == 1


def picture_document(image_scope):
    from PIL import Image
    picture = BytesIO(); Image.new('RGB', (10, 10), 'white').save(picture, format='PNG'); picture.seek(0)
    doc = Document()
    if image_scope == 'cover':
        doc.add_picture(picture)
    for title, body in SECTIONS[:-1]:
        doc.add_heading(title, 1); doc.add_paragraph(body)
    if image_scope == 'budget':
        doc.add_heading('投资估算', 1); doc.add_picture(picture)
    out = BytesIO(); doc.save(out)
    return out.getvalue()


def test_cover_logo_does_not_excuse_missing_budget():
    r = result(picture_document('cover'))
    assert group(r, 'budget')['status'] == 'missing'


def test_budget_picture_is_local_review():
    r = result(picture_document('budget'))
    assert group(r, 'budget')['status'] == 'review'
    assert group(r, 'construction')['status'] == 'pass'
    assert r['missing_count'] == 0


def test_old_history_keeps_original_result_and_version(api):
    client, sessions = api
    old = result(); old['version'] = 'completeness-1.0'; old['status'] = 'review'; old['label'] = '待核实'; old.pop('info_count')
    original = json.dumps(old, ensure_ascii=False)
    with sessions() as db:
        row = CheckResult(project_id='p', module='completeness', check_subtype='completeness_report', result_label='待核实', reason='旧报告', reference_data=original)
        db.add(row); db.commit(); old_id = row.id
    new = post(client).json()
    assert new['version'] == 'completeness-2.1' and new['status'] == 'pass'
    history = client.get('/c/reports?project_id=p').json()['items']
    old_entry = next(i for i in history if i['id'] == old_id)
    assert old_entry['version'] == 'completeness-1.0' and old_entry['status'] == 'review'
    assert old_entry['info_count'] == 0
    detail = client.get(f'/c/reports/{old_id}?project_id=p').json()
    assert detail['version'] == 'completeness-1.0' and detail['status'] == 'review'
    with sessions() as db:
        assert db.get(CheckResult, old_id).reference_data == original


def test_internal_chapter_reference_accepts_equivalent_appendix():
    doc = Document(BytesIO(document(extra=['明细详见附件章节13.3.1及13.3.2。'])))
    doc.add_heading('附件1 投资估算表', 1)
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = '项目', '金额（万元）'
    t.cell(1, 0).text, t.cell(1, 1).text = '软件开发', '120'
    out = BytesIO(); doc.save(out)
    r = result(out.getvalue())
    assert r['status'] == 'pass'
    assert '不校验引用章节编号' in group(r, 'attachment_附件')['reason']


def test_internal_chapter_reference_without_appendix_stays_review():
    r = result(document(extra=['明细详见附件章节13.3.1。']))
    assert r['status'] == 'review'


def test_unrelated_appendix_cannot_resolve_budget_reference():
    doc = Document(BytesIO(document(extra=['明细详见附件章节13.3.1。'])))
    doc.add_heading('附件1 单位情况', 1)
    doc.add_paragraph('本单位成立于2020年，现有工作人员120人。')
    out = BytesIO(); doc.save(out)
    assert result(out.getvalue())['status'] == 'review'


def test_merged_caption_does_not_fill_empty_budget_table():
    doc = Document(BytesIO(document(SECTIONS[:-1] + [('投资估算', '本项目总投资120万元。')])))
    t = doc.add_table(rows=3, cols=2)
    t.cell(0, 0).merge(t.cell(0, 1)).text = '项目投资估算表'
    t.cell(1, 0).text, t.cell(1, 1).text = '项目', '金额（万元）'
    out = BytesIO(); doc.save(out)
    r = result(out.getvalue())
    assert group(r, 'tables')['status'] == 'missing' and r['missing_count'] == 1


def checklist(title, headers, rows=(), prose=None, sections=SECTIONS, picture=False):
    doc = Document(BytesIO(document(sections)))
    doc.add_heading(title, 1)
    if prose:
        doc.add_paragraph(prose)
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    if picture:
        from PIL import Image
        data = BytesIO(); Image.new('RGB', (10, 10), 'white').save(data, format='PNG'); data.seek(0)
        table.add_row().cells[0].paragraphs[0].add_run().add_picture(data)
    out = BytesIO(); doc.save(out)
    return out.getvalue()


@pytest.mark.parametrize('title,headers,row,key', [
    ('建设功能清单', ('模块名称', '功能描述'), ('审批模块', '支持在线申请及审批'), 'functions'),
    ('设备采购清单', ('设备名称', '数量', '规格型号'), ('服务器', '2', ''), 'equipment'),
    ('服务内容清单', ('服务名称', '服务内容'), ('运行维护', '每月巡检并处理故障'), 'services'),
])
def test_major_checklist_data_can_supply_construction(title, headers, row, key):
    r = result(checklist(title, headers, [row], sections=[SECTIONS[3]]))
    assert r['status'] == 'pass'
    assert group(r, 'table_' + key)['status'] == 'pass'


@pytest.mark.parametrize('title,headers,key', [
    ('建设功能清单', ('模块名称', '功能描述'), 'functions'),
    ('设备采购清单', ('设备名称', '数量', '规格型号'), 'equipment'),
    ('服务内容清单', ('服务名称', '服务内容'), 'services'),
])
def test_whole_empty_checklist_without_equivalent_is_one_gap(title, headers, key):
    r = result(checklist(title, headers, sections=[SECTIONS[3]]))
    assert group(r, 'table_' + key)['status'] == 'missing'
    assert r['missing_count'] == 1


def test_function_prose_can_replace_empty_function_form():
    r = result(checklist('建设功能清单', ('模块名称', '功能描述')))
    assert r['status'] == 'pass'
    assert '等效表达' in group(r, 'table_functions')['reason']


def test_unrelated_function_prose_cannot_fill_empty_equipment_form():
    r = result(checklist('设备采购清单', ('设备名称', '数量')))
    assert r['status'] == 'missing' and r['missing_count'] == 1
    assert group(r, 'construction')['status'] == 'pass'


@pytest.mark.parametrize('prose', [
    '本项目不涉及设备采购，继续使用现有设施。',
    '本次采购两台服务器与一台交换机，配置于市级机房承担应用运行。',
])
def test_equipment_prose_or_nonapplicability_is_accepted(prose):
    assert result(checklist('设备采购清单', ('设备名称', '数量'), prose=prose))['status'] == 'pass'


def test_serial_numbers_and_placeholders_do_not_fill_a_form():
    r = result(checklist('设备采购清单', ('序号', '设备名称', '数量'), [('1', '待填写', '—')]))
    assert group(r, 'table_equipment')['status'] == 'missing'


def test_checklist_picture_needs_review_not_missing():
    r = result(checklist('设备采购清单', ('设备名称', '数量'), picture=True))
    assert group(r, 'table_equipment')['status'] == 'review'
    assert r['missing_count'] == 0 and r['review_count'] == 1


@pytest.mark.parametrize('title,headers', [
    ('绩效目标表', ('指标名称', '目标值')),
    ('实施进度表', ('任务', '时间')),
    ('资源申请表', ('资源名称', '申请数量')),
])
def test_optional_forms_are_not_new_mandatory_gates(title, headers):
    assert result(checklist(title, headers))['status'] == 'pass'


def test_whole_equipment_appendix_dependency_is_detected():
    r = result(document(extra=['本项目设备采购清单全部详见附表三。']))
    assert r['status'] == 'missing' and r['missing_count'] == 1


def test_equipment_appendix_accepts_same_topic_prose():
    r = result(document(extra=['本次采购两台服务器和一台交换机，统一部署于现有机房。', '本项目设备采购清单详见附表三。']))
    assert r['status'] == 'pass'
    assert '等效表达' in group(r, 'attachment_附表三')['reason']


def test_explicit_mandatory_attachment_is_not_waived_by_prose():
    r = result(document(extra=['本次采购两台服务器和一台交换机，统一部署于现有机房。', '本报告必须提交设备采购清单，详见附表三。']))
    assert r['status'] == 'missing'


def test_resource_table_only_required_on_explicit_dependency():
    r = result(document(extra=['本项目资源申请内容全部详见附表三。']))
    assert r['status'] == 'missing'
    ordinary = result(document(extra=['资源申请的填写格式可参见附表三。']))
    assert ordinary['status'] == 'pass' and ordinary['info_count'] == 1


def test_policy_checklist_reference_is_not_current_submission():
    r = result(document(extra=['相关政策《设备采购清单填报通知》中的格式详见附表三。']))
    assert r['status'] == 'pass' and r['info_count'] == 1


def test_supplied_checklist_is_equivalent_material():
    raw = document(extra=['本项目设备采购清单详见附表三。'])
    supplied = checklist('设备采购清单', ('设备名称', '数量'), [('服务器', '2')], sections=[])
    r = result(raw, attachments=[('采购说明.docx', supplied)])
    assert r['status'] == 'pass'


def test_interface_data_fields_are_not_equipment_headers():
    r = result(checklist('接口设计', ('对接方', '获取数据项'), [('外部系统', '产品名称、产品规格、申请数量、进口数量等')]))
    assert not any(i['key'] == 'table_equipment' for i in r['details'])


def test_service_fees_do_not_replace_missing_service_content():
    r = result(checklist('服务内容清单', ('服务名称', '服务内容'), sections=[SECTIONS[3]]))
    assert group(r, 'table_services')['status'] == 'missing'
