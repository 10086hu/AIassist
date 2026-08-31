from __future__ import annotations

import sys
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document


BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from app.modules.data_rules.service import run_data_rules_check_from_document  # noqa: E402


def _run_rule_18(document: Document) -> dict:
    stream = BytesIO()
    document.save(stream)
    return run_data_rules_check_from_document(
        content=stream.getvalue(),
        filename="budget_case.docx",
        project_name="预算一致性测试项目",
        selected_rule_ids=["18"],
    )["results"][0]


def _add_budget_table(document: Document, title: str, amounts: list[float], total: float) -> None:
    table = document.add_table(rows=len(amounts) + 3, cols=3)
    table.rows[0].cells[0].text = title
    table.rows[1].cells[0].text = "序号"
    table.rows[1].cells[1].text = "项目"
    table.rows[1].cells[2].text = "总金额（万元）"
    for index, amount in enumerate(amounts, start=1):
        row = table.rows[index + 1]
        row.cells[0].text = str(index)
        row.cells[1].text = f"子项{index}"
        row.cells[2].text = str(amount)
    total_row = table.rows[len(amounts) + 2]
    total_row.cells[0].text = "合计"
    total_row.cells[2].text = str(total)


class DataBudgetConsistencyRuleTests(unittest.TestCase):
    def test_only_project_total_is_insufficient_not_calculation_error(self) -> None:
        document = Document()
        document.add_paragraph("2.7 投资概况 本项目总投资100万元。")

        rule = _run_rule_18(document)

        self.assertFalse(rule["passed"])
        self.assertEqual(rule["status"], "资料不足")
        self.assertIn("仅识别到项目总投资金额", rule["issues"][0]["message"])

    def test_detail_table_internal_sum_can_pass_without_forcing_all_details(self) -> None:
        document = Document()
        _add_budget_table(document, "应用软件开发投资估算明细表", [40, 60], 100)

        rule = _run_rule_18(document)

        self.assertTrue(rule["passed"], rule)
        self.assertEqual(rule["metrics"]["docx_table_internal_checks"], 1)

    def test_detail_table_internal_sum_error_fails_rule_18(self) -> None:
        document = Document()
        _add_budget_table(document, "应用软件开发投资估算明细表", [40, 60], 120)

        rule = _run_rule_18(document)

        self.assertFalse(rule["passed"])
        self.assertEqual(rule["status"], "failed")
        self.assertIn("预算明细表合计行内部加总不一致", rule["issues"][0]["message"])

    def test_component_detail_totals_sum_to_project_total(self) -> None:
        document = Document()
        document.add_paragraph("2.7 投资概况 本项目总投资100万元。")
        _add_budget_table(document, "应用软件开发投资估算明细表", [15, 25], 40)
        _add_budget_table(document, "硬件设备购置投资估算明细表", [60], 60)

        rule = _run_rule_18(document)

        self.assertTrue(rule["passed"], rule)
        self.assertEqual(rule["metrics"]["component_detail_total_sum_wan"], 100)
        self.assertGreaterEqual(rule["metrics"]["calculation_checks_performed"], 3)

    def test_hierarchical_investment_summary_table_passes(self) -> None:
        document = Document()
        document.add_paragraph("2.7 投资概况 本项目总投资1932.562万元。")
        table = document.add_table(rows=8, cols=3)
        rows = [
            ["松江区投资估算表", "", ""],
            ["序号", "名称", "总价（万元）"],
            ["一", "工程建设费", "1747.671"],
            ["（一）", "软件建设", "435.000"],
            ["（二）", "信息化配套系统", "1189.771"],
            ["（三）", "密码应用建设费", "122.900"],
            ["二", "工程建设其他费", "184.891"],
            ["三", "总投资估算（一+二）", "1932.562"],
        ]
        for row_index, values in enumerate(rows):
            for col_index, value in enumerate(values):
                table.rows[row_index].cells[col_index].text = value

        rule = _run_rule_18(document)

        self.assertTrue(rule["passed"], rule)
        self.assertEqual(rule["metrics"]["distinct_project_totals_wan"], [1932.562])
        self.assertEqual(rule["metrics"]["detail_total_candidates"][0]["amount_wan"], 1932.562)
        self.assertGreaterEqual(rule["metrics"]["docx_summary_internal_checks"], 2)


if __name__ == "__main__":
    unittest.main()
