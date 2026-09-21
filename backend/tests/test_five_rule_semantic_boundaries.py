"""Counterexamples for deployment coverage and evidence-based review levels."""
from io import BytesIO
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pytest
from docx import Document
from app.modules.resource.document_check import check_resource_document, _coverage_statement
from app.modules.resource.quantity_parser import ParsedResourceItem
from app.modules.resource.quantity_rules import check_security_service_quantities
from app.modules.data_rules.document_rules import _count_indicator_table_rows, _check_hierarchical_budget_rows
from app.modules.data_rules.service import run_data_rules_check_from_document
from app.modules.data_rules.data_governance_service import (
    DATA_GOVERNANCE_SERVICE_SPECS, _check_governance_acceptance_requirements,
    validate_data_governance_service_design,
)


def doc_bytes(d):
    b = BytesIO(); d.save(b); return b.getvalue()


@pytest.mark.parametrize('sentence,field,n', [
    ('所有服务器配同一套操作系统。', '操作系统', 8),
    ('本项目服务器统一安装银河麒麟操作系统。', '操作系统', 8),
    ('所有服务器统一使用银河麒麟。', '操作系统', 8),
    ('所有服务器及操作系统由政务云平台统一提供。', '操作系统', 8),
    ('6台数据库服务器共享一套数据库。', '数据库', 6),
    ('本项目多个数据库服务器配一套数据库。', '数据库', 6),
    ('多个数据库服务器配一套数据库。', '数据库', 0),
    ('全部数据库节点共用一套达梦数据库集群。', '数据库', 4),
    ('六台数据库服务器共享一套数据库。', '数据库', 6),
    ('8个数据库节点共享一套数据库。', '数据库', 8),
    ('数据库服务器（共六台）共享一套数据库。', '数据库', 6),
])
def test_explicit_coverage(sentence, field, n):
    assert _coverage_statement(sentence, field, n)


@pytest.mark.parametrize('sentence', [
    '所有服务器仅支持安装操作系统。',
    '建议所有服务器统一安装操作系统。',
    '例如所有服务器配同一套操作系统。',
    '上期所有服务器配同一套操作系统。',
    '所有服务器尚未配置操作系统。',
    '所有服务器未全部配置操作系统。',
    '其中2台服务器共享一套操作系统。',
    '部分服务器统一安装操作系统。',
    '2台服务器统一安装操作系统。',
    '两台服务器统一安装操作系统。',
    '所有服务器无需配置操作系统。',
    '操作系统支持所有服务器统一部署。',
    '每台服务器均需配备系统盘，用于操作系统、运行环境及存储关键日志和配置文件，所有服务器的系统盘均采用RAID5方案进行冗余配置。',
])
def test_no_blanket_pass_for_non_coverage(sentence):
    assert not _coverage_statement(sentence, '操作系统', 8)


def test_real_doc_many_to_one_with_quantities_and_location():
    d = Document(); d.add_paragraph('本项目为市级上云项目。')
    d.add_heading('6.2.1.3 部署配置', 4)
    d.add_paragraph('所有服务器配同一套操作系统。')
    d.add_paragraph('本项目多个数据库服务器共用一套数据库。')
    items = [ParsedResourceItem(1, '资源清单', '计算资源', '应用服务器', 8, '台'),
             ParsedResourceItem(2, '资源清单', '计算资源', '数据库服务器', 6, '台'),
             ParsedResourceItem(3, '软件清单', 'PaaS服务清单', '操作系统', 1, '套'),
             ParsedResourceItem(4, '软件清单', 'PaaS服务清单', '数据库软件', 1, '套')]
    results = check_resource_document(doc_bytes(d), '市级报告.docx', items)
    for result in results[1:]:
        assert result.result_label == '通过', result
        assert result.evidence_location
        assert '部署覆盖' in result.reason


def test_explicit_unconfigured_row_blocks_shared_claim():
    d = Document(); d.add_paragraph('所有服务器统一安装操作系统。')
    d.add_paragraph('新增的服务器未配置操作系统。')
    findings = check_resource_document(doc_bytes(d), 'report.docx', [])
    assert findings[1].result_label != '通过'


def test_no_paas_does_not_hide_security_crypto_mismatch():
    d = Document(); d.add_paragraph('本项目为市级项目。')
    d.add_paragraph('本项目不涉及PaaS服务。')
    items = [ParsedResourceItem(1, '安全表', '安全服务需求表', '数字证书服务', 2, '个'),
             ParsedResourceItem(1, '密码表', '密码服务资源内容清单', '数字证书服务', 3, '个')]
    f = check_resource_document(doc_bytes(d), 'report.docx', items)[0]
    assert f.result_label == '不一致'


def test_two_actual_linked_lists_can_pass_without_third():
    items = [ParsedResourceItem(1, '安全表', '安全服务需求表', '数字证书服务', 2, '个'),
             ParsedResourceItem(1, '密码表', '密码服务资源内容清单', '数字证书服务', 2, '个')]
    assert check_security_service_quantities(items)[0].result_label == '通过'


def test_different_project_stages_are_not_direct_mismatch():
    items = [ParsedResourceItem(1, '安全表', '安全服务需求表', '数字证书服务', 2, '个', spec='新增配置'),
             ParsedResourceItem(1, '密码表', '密码服务资源内容清单', '数字证书服务', 3, '个', spec='存量配置')]
    assert check_security_service_quantities(items)[0].result_label == '资料不足'


def test_rounding_is_review_but_large_difference_remains_error():
    rows = [['一', '软件', '1.0'], ['1', '模块甲', '0.4'], ['2', '模块乙', '0.3'], ['3', '模块丙', '0.4']]
    result = _check_hierarchical_budget_rows(rows, 2, '万元', '表1')
    assert result['arithmetic_issues'][0].severity == 'warning'
    rows[0][2] = '2.0'
    assert _check_hierarchical_budget_rows(rows, 2, '万元', '表1')['arithmetic_issues'][0].severity == 'risk'


def test_duplicate_indicators_do_not_inflate_counts_and_aliases_work():
    rows = [['序号', '一级指标', '二级指标', '指标名称', '目标值'],
            ['1', '共性指标', '社会效益', '公众满意度', '90%'],
            ['2', '共性指标', '社会效益', '公众满意度', '90%'],
            ['3', '行业指标', '产出指标', '服务事项数', '4项']]
    common, business, ai, _ = _count_indicator_table_rows(rows)
    assert (common, business, ai) == (1, 1, 0)


def test_ai_subject_in_merged_parent_is_preserved():
    rows = [['序号', '一级指标', '二级指标', '三级指标', '指标名称', '目标值'],
            ['1', '业务指标', '产出指标', '智能问答应用', '回答准确率', '90%'],
            ['2', '', '', '', '平均响应时间', '3秒']]
    assert _count_indicator_table_rows(rows)[2] == 2


def specs(*names):
    return [s for s in DATA_GOVERNANCE_SERVICE_SPECS if s.name in names]


def test_semantic_acceptance_and_non_service_logs():
    text = '本项目数据融合服务完成政务目录链登记、归集至公共数据平台，并自行开展数据质量评估。'
    assert not _check_governance_acceptance_requirements(text, specs('数据融合'))
    assert not _check_governance_acceptance_requirements('数据标签服务验收提交标注操作记录。', specs('数据标签'))
    assert any(e.code == 'DGS-011' for e in _check_governance_acceptance_requirements('安全运维系统提供执行日志。', specs('数据标签')))
    assert any(e.code == 'DGS-011' for e in _check_governance_acceptance_requirements('数据标签服务不提供执行日志。', specs('数据标签')))
    missing = _check_governance_acceptance_requirements('6.4 数据上链内容', specs('数据融合'))
    assert any('数据上链' in e.message for e in missing)
    explained = _check_governance_acceptance_requirements('6.4 数据上链内容\n项目建设后将根据市级要求配合开展上链工作。', specs('数据融合'))
    assert not any('交付要求：数据上链' in e.message for e in explained)


def test_budget_unit_review_only_does_not_become_hard_arithmetic_failure():
    d = Document()
    t = d.add_table(rows=4, cols=3)
    for cells, values in zip(t.rows, [['项目', '数量', '总价（万元）'], ['软件', '1', '4.0'], ['设备', '1', '20000'], ['总计', '', '6.0']]):
        for c,v in zip(cells.cells, values): c.text = v
    r = run_data_rules_check_from_document(doc_bytes(d), '预算.docx', '项目', ['18'])['results'][0]
    assert r['status'] != 'failed', r
    assert r['severity'] == 'warning'
    assert any('单位待核实' in i['message'] for i in r['issues'])


def test_known_non_municipal_scope_skips_resource_rules():
    d = Document(); d.add_paragraph('本项目属于区级项目。')
    f = check_resource_document(doc_bytes(d), '项目.docx', [])
    assert all(x.result_label == '不适用' for x in f)


def test_reference_example_does_not_exclude_cloud_scope():
    d = Document(); d.add_paragraph('本项目为市级上云项目。')
    d.add_paragraph('例如本项目全部本地部署时另行测算。')
    f = check_resource_document(doc_bytes(d), '项目.docx', [])
    assert f[1].result_label != '不适用'


def test_2025_report_is_reference_not_retroactive_failure():
    d = Document(); d.add_paragraph('预算单位：上海市市级预算单位')
    d.add_paragraph('本项目申报数据治理服务，包括数据融合。')
    r = run_data_rules_check_from_document(doc_bytes(d), '2025年建设报告.docx', '项目', ['24'])['results'][0]
    assert r['status'] == '待复核'
    assert r['metrics']['policy_temporal_mode'] == 'reference_check'


@pytest.mark.parametrize('description,flag', [
    ('对政策文件按法律法规类别、时效性和行业领域进行特征标注', False),
    ('为算法模型进行训练数据标注', True),
])
def test_feature_label_is_not_automatically_ai_training(description, flag):
    d = Document(); d.add_paragraph('预算单位：上海市市级预算单位')
    d.add_heading('6.3 数据治理服务', 2)
    t=d.add_table(rows=2, cols=4)
    for row, values in zip(t.rows, [['服务事项','测算单位','工作量','服务内容'],['数据标签','个标签','500',description]]):
        for c,v in zip(row.cells,values):c.text=v
    r=run_data_rules_check_from_document(doc_bytes(d),'报告.docx','项目',['24'])['results'][0]
    assert any('训练数据标注' in i['message'] for i in r['issues']) == flag
