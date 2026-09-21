from __future__ import annotations

import sys
import unittest
from io import BytesIO
from pathlib import Path


BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from app.modules.data_rules.service import (  # noqa: E402
    DATA_REASONABLENESS_RULES,
    run_data_reporting_check_from_document,
)


def _run_rule_24(text: str) -> dict:
    return run_data_reporting_check_from_document(
        content=("预算单位：上海市市级预算单位\n" + text).encode("utf-8"),
        filename="sample.txt",
        project_name="数据治理测试项目",
        selected_rule_ids=["24"],
    )


def _run_rule_24_bytes(content: bytes, filename: str) -> dict:
    return run_data_reporting_check_from_document(
        content=content,
        filename=filename,
        project_name="数据治理测试项目",
        selected_rule_ids=["24"],
    )


class DataGovernanceServiceRuleTests(unittest.TestCase):
    def test_rule_24_is_registered_without_replacing_existing_rules(self) -> None:
        rules = {str(rule["rule_id"]): rule for rule in DATA_REASONABLENESS_RULES}
        self.assertIn("DATA_REASON_024", rules)
        self.assertIn("DATA_REASON_004", rules)
        self.assertEqual(rules["DATA_REASON_024"]["excel_row"], 24)
        self.assertEqual(rules["DATA_REASON_024"]["rule_name"], "数据治理服务内容设计合规性审查规则")

    def test_valid_6_3_governance_annex_passes(self) -> None:
        text = """
6.3 数据治理内容附表
序号 服务类型 服务事项 测算单位 标准 适用项目阶段 申报条件
1 数据采集 数据购买 项 每类第三方数据集为1项 建设/运维 完成数据上链
2 数据采集 历史数据迁移 千万条 数据库记录总数按千万条计 建设 满足系统整合重构需求
3 数据加工 数据标签 个标签 新增数据标签个数 建设/运维 满足分类分级要求
4 数据加工 数据融合 万条 需融合的数据条数 建设/运维 生成高质量数据集
5 数据退役 历史数据归档及销毁 张 周期内归档销毁的数据表数 运维 满足退役要求
6 数据安全管理 数据安全风险评估 项 按预算单位评估工作计为1项 运维 开展安全风险评估
6.4 数据上链内容
本项目按政务目录链要求完成新数据资源归集和上链。\n数据购买、历史数据迁移和数据融合服务交付数据质量自评估报告。\n数据标签和历史数据归档及销毁服务交付执行日志。\n数据安全风险评估由具备信息安全服务资质的第三方交付评估报告和整改报告。
"""
        result = _run_rule_24(text)
        self.assertEqual(len(result["results"]), 1)
        rule = result["results"][0]
        self.assertEqual(rule["rule_id"], "DATA_REASON_024")
        self.assertTrue(rule["passed"], rule)
        self.assertEqual(rule["issues"], [])

    def test_negative_list_and_wrong_unit_are_reported(self) -> None:
        text = """
6.3 数据治理内容附表
序号 服务类型 服务事项 测算单位 标准 适用项目阶段 申报条件
1 数据加工 数据融合 项 需融合的数据条数 建设/运维 生成高质量数据集
2 数据采集 数据抽取服务 项 按接口接入数量测算 建设 纳入数据治理服务申报
6.4 数据上链内容
本项目说明数据资源归集计划。
"""
        result = _run_rule_24(text)
        rule = result["results"][0]
        self.assertFalse(rule["passed"])
        messages = "\n".join(issue["message"] for issue in rule["issues"])
        self.assertIn("数据抽取服务", messages)
        self.assertIn("负面清单", messages)
        self.assertIn("数据融合", messages)
        self.assertIn("测算单位", messages)

    def test_missing_6_3_annex_is_reported_when_governance_service_is_in_scope(self) -> None:
        text = """
5.1 建设内容
本项目申报数据治理服务，包括数据融合和历史数据迁移。
6.4 数据上链内容
本项目将按政务目录链要求完成数据上链。
"""
        result = _run_rule_24(text)
        rule = result["results"][0]
        self.assertFalse(rule["passed"])
        self.assertEqual(rule["issues"][0]["section"], "6.3")
        self.assertIn("未识别到可核验的服务附表", rule["issues"][0]["message"])

    def test_6_3_table_candidate_is_detected_without_6_3_2_heading(self) -> None:
        text = """
6.3 数据库建设和数据治理内容
序号 数据治理服务事项 测算单位 工作量 治理对象及范围 服务内容 治理成效及交付物 关联模块
1 数据采集-历史数据迁移-特殊区域 人月 2 特殊区域库 实现历史数据迁移 《数据接入方案》 6.1.1
2 数据加工-数据融合-报关单宽表 人月 1 报关单库 形成报关单宽表 形成高质量数据集 6.1.1
6.4 数据上链内容
本项目将按政务目录链要求完成数据上链。
"""
        result = _run_rule_24(text)
        rule = result["results"][0]
        messages = "\n".join(issue["message"] for issue in rule["issues"])
        self.assertFalse(rule["passed"])
        self.assertNotIn("申报核验信息", messages)
        self.assertNotIn("服务类型、标准、适用项目阶段、申报条件", messages)
        self.assertIn("历史数据迁移", messages)
        self.assertIn("千万条 / 百万个", messages)
        self.assertIn("数据融合", messages)
        self.assertIn("万条", messages)
        self.assertNotIn("未识别到《数据治理服务配置指引》允许的4类6项", messages)

    def test_docx_6_3_tables_are_used_for_governance_annex_detection(self) -> None:
        try:
            from docx import Document
        except Exception as exc:  # pragma: no cover - dependency guard for minimal envs
            self.skipTest(f"python-docx is unavailable: {exc}")

        document = Document()
        document.add_paragraph("预算单位：上海市市级预算单位")
        document.add_paragraph("6.3 数据库建设和数据治理内容")
        document.add_paragraph("6.3.1 数据内容分析")
        other_table = document.add_table(rows=2, cols=3)
        other_table.rows[0].cells[0].text = "场景"
        other_table.rows[0].cells[1].text = "主要数据项"
        other_table.rows[0].cells[2].text = "是否需要数据治理"
        other_table.rows[1].cells[0].text = "外贸气象站"
        other_table.rows[1].cells[1].text = "报关单申报数据"
        other_table.rows[1].cells[2].text = "是"
        document.add_paragraph("6.3.2 数据治理内容")
        governance_table = document.add_table(rows=3, cols=8)
        headers = [
            "序号",
            "数据治理服务事项",
            "测算单位",
            "工作量",
            "治理对象及范围",
            "服务内容",
            "治理成效及交付物",
            "关联模块",
        ]
        for index, header in enumerate(headers):
            governance_table.rows[0].cells[index].text = header
        rows = [
            ["1", "数据采集-历史数据迁移-特殊区域", "人月", "2", "特殊区域库", "实现历史数据迁移", "《数据接入方案》", "6.1.1"],
            ["2", "数据加工-数据融合-报关单宽表", "人月", "1", "报关单库", "形成报关单宽表", "形成高质量数据集", "6.1.1"],
        ]
        for row_index, row_values in enumerate(rows, start=1):
            for col_index, value in enumerate(row_values):
                governance_table.rows[row_index].cells[col_index].text = value
        document.add_paragraph("6.4 数据上链内容")

        stream = BytesIO()
        document.save(stream)
        result = _run_rule_24_bytes(stream.getvalue(), "sample.docx")
        rule = result["results"][0]
        messages = "\n".join(issue["message"] for issue in rule["issues"])
        self.assertFalse(rule["passed"])
        self.assertNotIn("申报核验信息", messages)
        self.assertNotIn("服务类型、标准、适用项目阶段、申报条件", messages)
        self.assertIn("历史数据迁移", messages)
        self.assertIn("千万条 / 百万个", messages)
        self.assertIn("数据融合", messages)
        self.assertIn("万条", messages)
        self.assertNotIn("未识别到《数据治理服务配置指引》允许的4类6项", messages)


if __name__ == "__main__":
    unittest.main()
