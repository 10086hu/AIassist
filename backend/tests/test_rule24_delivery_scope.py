"""Rule 24: conditional service deliverables and evidence scope, not project acceptance."""
from io import BytesIO

import pytest
from docx import Document

from app.modules.data_rules.data_governance_service import (
    DATA_GOVERNANCE_SERVICE_SPECS, GovernanceAnnexTable,
    _check_governance_acceptance_requirements, _produces_new_governance_data,
)
from app.modules.data_rules.service import run_data_rules_check_from_document


def specs(*names):
    return [s for s in DATA_GOVERNANCE_SERVICE_SPECS if s.name in names]


def checks(text, names=('数据融合',), tables=None):
    return _check_governance_acceptance_requirements(text, specs(*names), tables)


def table(*rows):
    return GovernanceAnnexTable(
        rows=(('服务事项', '测算单位', '服务内容', '治理成效及交付物'), *rows),
        heading_text='6.3.2 数据治理内容', heading_number='6.3.2')


@pytest.mark.parametrize('text', [
    '数据融合', '6.4 数据上链内容',
    '本项目提供历史数据迁移服务。',
    '历史数据迁移仅原样搬迁已有数据，不产生新数据。',
    '本项目数据融合不形成新数据集，复用现有成果。',
    '指引规定数据融合形成新数据时应归集。',
    '数据融合服务例如可以形成新数据集。',
    '后续项目的数据融合将形成新数据集。',
    '本项目数据融合是否形成新数据集尚待确认。',
    '本项目数据融合可能形成新数据集。',
    '本项目新增数据融合服务，复用已有成果。',
    '数据融合服务形成数据质量评价报告。',
    '数据融合服务不新增400万条数据。',
])
def test_service_name_or_uncertain_new_data_does_not_trigger_missing_acceptance(text):
    assert not any(e.code == 'DGS-010' for e in checks(text, ('数据融合', '历史数据迁移')))


@pytest.mark.parametrize('text', [
    '本项目数据融合形成报关单宽表。',
    '本期历史数据迁移将生成新数据集。',
    '数据购买服务引入后将形成新的高质量数据集。',
    '数据融合不新增业务数据，但形成关联宽表。',
    '数据融合服务产生新的业务数据。',
    '数据融合服务新增400万条业务数据。',
    '数据购买服务引入第三方数据集。',
])
def test_explicit_planned_new_data_checks_delivery_agreements(text):
    result = checks(text, ('数据融合', '历史数据迁移', '数据购买'))
    assert any(e.code == 'DGS-010' for e in result)
    assert all(e.severity == 'warning' for e in result)
    assert all(e.section != '验收要求' for e in result)
    assert '不等同于缺少竣工验收材料' in result[0].message


def test_service_table_deliverables_and_linked_prose_can_jointly_satisfy_requirement():
    annex = table(('数据融合', '万条', '形成新的高质量数据集', '提交数据质量自评估报告'))
    prose = '本项目数据治理服务统一完成政务目录链登记和资源归集。'
    assert not checks(prose, tables=[annex])


def test_multi_paragraph_service_delivery_and_equivalent_phrasing_are_accepted():
    text = ('7.2 数据融合服务\n形成业务宽表。\n'
            '交付时完成政务目录链编目，归集至公共数据平台。\n'
            '项目单位自行开展数据质量评估。')
    assert not checks(text)


def test_generic_acceptance_heading_never_satisfies_service_delivery():
    result = checks('数据融合形成新的数据集。\n8.4 项目验收要求\n完成验收后交付。')
    assert any(e.code == 'DGS-010' for e in result)
    assert all('验收要求' not in e.section for e in result)


@pytest.mark.parametrize('heading', ['6.4 数据上链内容', '6.4.数据上链内容', '六、数据上链内容', '（六）数据上链内容', '数据上链内容'])
def test_numbered_heading_alone_is_not_an_on_chain_commitment(heading):
    result = checks('数据融合形成新的业务数据集。\n' + heading + '\n本项目不涉及对接政务目录链的数据上链，后续计划接入国家航贸链。')
    assert '政务目录链上链' in result[0].message


def test_other_service_cannot_supply_this_services_missing_deliverables():
    text = ('数据融合形成业务宽表，交付宽表清单。\n'
            '数据购买服务完成政务目录链登记、资源归集和数据质量自评估。')
    assert any(e.code == 'DGS-010' for e in checks(text))


def test_quoted_policy_and_unrelated_security_logs_are_not_delivery_evidence():
    text = ('数据融合形成新数据集。\n'
            '指引规定：数据融合完成数据上链、资源归集和质量自评估。\n'
            '安全运维提供执行日志。')
    assert any(e.code == 'DGS-010' for e in checks(text))
    assert any(e.code == 'DGS-011' for e in checks(text, ('数据标签',)))


@pytest.mark.parametrize('chain_text', [
    '本项目不涉及对接政务目录链的数据上链，计划接入国家航贸链。',
    '本项目数据治理服务在国家航贸链完成数据上链。',
    '本项目数据治理服务不开展数据上链。',
])
def test_other_chain_or_negative_statement_is_not_government_directory_chain(chain_text):
    result = checks('数据融合形成业务宽表。\n' + chain_text)
    assert '政务目录链上链' in result[0].message


def test_explicit_new_output_is_not_hidden_by_other_rows_no_new_data():
    annex = table(
        ('数据融合-复用', '万条', '不形成新数据集', '复用原有成果'),
        ('数据融合-新增宽表', '万条', '形成业务宽表', '宽表清单'),
    )
    result = checks('', tables=[annex])
    assert result and '表内行3' in result[0].section


def test_each_log_service_is_checked_without_borrowing_other_services_log():
    result = checks('数据标签服务交付执行日志。', ('数据标签', '历史数据归档及销毁'))
    assert len(result) == 1 and result[0].code == 'DGS-011'
    assert result[0].message.startswith('历史数据归档及销毁')


def test_own_table_logging_agreement_is_enough_at_feasibility_stage():
    annex = table(('数据标签', '个标签', '开展分类分级标签工作', '交付标注操作记录'))
    assert not checks('', ('数据标签',), [annex])


def docx_rule(heading='6.3.2 数据治理内容'):
    doc = Document()
    doc.add_paragraph('预算单位：上海市市级预算单位')
    doc.add_heading('4.2.2 验收标准、具体指标和内容', 3)
    doc.add_paragraph('项目通过整体测评后验收。')
    doc.add_heading(heading, 3)
    grid = doc.add_table(rows=2, cols=5)
    for cell, text in zip(grid.rows[0].cells, ['序号', '数据治理服务事项', '测算单位', '服务内容', '治理成效及交付物']):
        cell.text = text
    for cell, text in zip(grid.rows[1].cells, ['1', '数据融合', '万条', '形成新的高质量业务数据集', '《成果清单》']):
        cell.text = text
    doc.add_heading('8.4 项目验收阶段', 2)
    doc.add_paragraph('项目建成后组织竣工验收。')
    stream = BytesIO(); doc.save(stream)
    result = run_data_rules_check_from_document(stream.getvalue(), '本项目2026年可研.docx', '市级项目', ['24'])
    return result['results'][0]


@pytest.mark.parametrize('heading', ['6.3.2 数据治理内容', '7.2 数据治理服务方案'])
def test_missing_deliverable_locates_service_table_as_scope_not_generic_acceptance(heading):
    result = docx_rule(heading)
    i = next(i for i, issue in enumerate(result['issues']) if issue.get('code') == 'DGS-010')
    issue = result['issues'][i]
    evidence = result['evidence_locations'][i]
    assert '表内行2' in issue['section']
    assert evidence['source_locations']
    for loc in evidence['source_locations']:
        assert loc['table_index'] == 1 and loc['row_index'] == 2
        assert loc['role'] == 'scope' and loc['match_status'] == 'scope'
        assert '验收' not in loc['quote'] and '验收' not in loc['section']
    assert evidence['source_location']['precision'] == 'scope'
    assert '相关服务及交付物核查范围' in evidence['location_hint']
    assert '可研阶段' in evidence['revision_target']['advice']
