from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import html

from app.api.evaluate import (
    _precise_rule_module, _render_docx_preview, _render_rule_evidence_preview,
    evaluate_result_source_preview,
)
from app.modules.resource.quantity_parser import DocumentIndex


def raw(document):
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def page(document):
    return html.fromstring(_render_rule_evidence_preview(raw(document), "报告.docx"))


def test_table_paragraphs_do_not_shift_body_coordinates():
    doc = Document()
    doc.add_paragraph("第一段")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "服务器"
    table.cell(0, 1).text = "2"
    doc.add_paragraph("第二段")
    root = page(doc)
    assert root.xpath('//*[@data-e-paragraph="2"]')[0].get("data-e-text") == "第二段"
    assert root.xpath('//*[@data-e-kind="row"]')[0].get("data-e-text") == "服务器 | 2"
    assert len(root.xpath('//*[@data-e-paragraph]')) == 2


def test_four_level_title_and_repeated_text_keep_distinct_addresses():
    doc = Document()
    for section in ["一期", "二期"]:
        doc.add_heading(section, 1)
        for title, level in [("资源申请", 2), ("计算资源", 3), ("数据库配置", 4)]:
            doc.add_heading(title, level)
        doc.add_paragraph("数据库统一覆盖所有节点。")
    root = page(doc)
    nodes = root.xpath('//*[@data-e-text="数据库统一覆盖所有节点。"]')
    assert [node.get("data-e-paragraph") for node in nodes] == ["5", "10"]
    assert nodes[1].get("data-e-section") == "二期 → 资源申请 → 计算资源 → 数据库配置"
    assert root.xpath('//h4[@data-e-text="数据库配置"]')


def test_same_row_in_two_tables_and_nested_table_are_distinct():
    doc = Document()
    first = doc.add_table(rows=1, cols=1)
    first.cell(0, 0).text = "同名项目"
    nested = first.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套明细"
    second = doc.add_table(rows=1, cols=1)
    second.cell(0, 0).text = "同名项目"
    root = page(doc)
    assert len(root.xpath('//*[@data-e-kind="row" and @data-e-path=""]')) == 2
    assert root.xpath('//*[@data-e-kind="row" and @data-e-table="1" and @data-e-path="1/1/1"]')[0].get("data-e-text") == "嵌套明细"


def test_merged_cells_use_physical_origin_and_spans():
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "纵向合并"
    table.cell(0, 0).merge(table.cell(2, 0))
    table.cell(0, 1).text = "横向合并"
    table.cell(0, 1).merge(table.cell(0, 2))
    root = page(doc)
    vertical = root.xpath('//*[@data-e-kind="cell" and @data-e-column="1"]')
    assert len(vertical) == 1
    assert vertical[0].get("rowspan") == "3"
    horizontal = root.xpath('//*[@data-e-kind="cell" and @data-e-column="2" and @data-e-row="1"]')[0]
    assert horizontal.get("colspan") == "2"


def test_toc_is_visible_but_has_no_false_evidence_anchor():
    doc = Document()
    paragraph = doc.add_paragraph("目录中的同名文本")
    paragraph.style = doc.styles.add_style("TOC 1", 1)
    doc.add_paragraph("实际正文")
    root = page(doc)
    assert "目录中的同名文本" in root.text_content()
    assert not root.xpath('//*[@data-e-text="目录中的同名文本"]')
    assert root.xpath('//*[@data-e-paragraph="2"]')[0].get("data-e-text") == "实际正文"


def test_content_control_paragraph_uses_existing_index_numbering():
    doc = Document()
    paragraph = doc.add_paragraph("内容控件正文")
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    paragraph._p.addprevious(control)
    content.append(paragraph._p)
    control.append(content)
    doc.add_paragraph("普通正文")
    root = page(doc)
    assert root.xpath('//*[@data-e-paragraph="1"]')[0].get("data-e-text") == "内容控件正文"
    assert root.xpath('//*[@data-e-paragraph="2"]')[0].get("data-e-text") == "普通正文"


def test_document_text_is_escaped_and_index_not_mutated():
    doc = Document()
    doc.add_paragraph('<script>alert("x")</script>')
    content = raw(doc)
    before = DocumentIndex(content, "报告.docx").blocks
    result = _render_rule_evidence_preview(content, "报告.docx")
    root = html.fromstring(result)
    assert root.xpath('//*[@data-e-paragraph="1"]')[0].text_content() == '<script>alert("x")</script>'
    assert before == DocumentIndex(content, "报告.docx").blocks
    assert "rule-evidence-v1" in result


@pytest.mark.parametrize("rule,module", [
    ("R15_SECURITY_PAAS_CRYPTO_QUANTITY", "resource"),
    ("R16_SERVER_OS_QUANTITY", "resource"),
    ("R16_DB_SERVER_DATABASE_QUANTITY", "resource"),
    ("DATA_REASON_001", "data_reasonableness"),
    ("DATA_REASON_002", "data_reasonableness"),
    ("DATA_REASON_024", "data_reasonableness"),
    ("FUNC_CORR_004", None), ("BASIS_007", None),
])
def test_only_target_rules_opt_in(rule, module):
    assert _precise_rule_module(rule) == module


class FakeDb:
    def __init__(self, content, module="resource", filename="报告.docx"):
        self.artifact = SimpleNamespace(content=content, filename=filename)
        self.result = SimpleNamespace(module=module)

    def query(self, *args): return self
    def filter(self, *args): return self
    def first(self): return self.artifact
    def get(self, *args): return self.result


def test_default_preview_response_remains_legacy_renderer():
    doc = Document()
    doc.add_paragraph("旧预览仍保持原样")
    content = raw(doc)
    response = evaluate_result_source_preview("result", "source", db=FakeDb(content))
    assert response.body.decode() == _render_docx_preview(content, "报告.docx")
    assert b"locateRuleEvidence" not in response.body


def test_precise_endpoint_is_opt_in_and_checks_module():
    doc = Document()
    doc.add_paragraph("原始报告")
    content = raw(doc)
    response = evaluate_result_source_preview("result", "source", db=FakeDb(content), precision_rule="R16_SERVER_OS_QUANTITY")
    assert b"locateRuleEvidence" in response.body
    with pytest.raises(Exception) as caught:
        evaluate_result_source_preview("result", "source", db=FakeDb(content, "basis"), precision_rule="R16_SERVER_OS_QUANTITY")
    assert caught.value.status_code == 400


def test_non_docx_is_not_silently_treated_as_precise_word():
    with pytest.raises(Exception) as caught:
        evaluate_result_source_preview("result", "source", db=FakeDb(b"test", filename="表.xlsx"), precision_rule="R16_SERVER_OS_QUANTITY")
    assert caught.value.status_code == 400


def test_precise_source_snapshot_belongs_to_requested_record():
    doc = Document()
    doc.add_paragraph("历史记录对应原文")
    missing = FakeDb(raw(doc))
    missing.artifact = None
    with pytest.raises(Exception) as caught:
        evaluate_result_source_preview("wrong-record", "source", db=missing, precision_rule="R16_SERVER_OS_QUANTITY")
    assert caught.value.status_code == 404


def test_unknown_rule_cannot_switch_renderer():
    doc = Document()
    doc.add_paragraph("普通规则继续使用旧预览")
    with pytest.raises(Exception) as caught:
        evaluate_result_source_preview("result", "source", db=FakeDb(raw(doc)), precision_rule="FUNC_CORR_004")
    assert caught.value.status_code == 400


def test_broken_docx_reports_unavailable_instead_of_highlighting_other_text():
    with pytest.raises(Exception) as caught:
        evaluate_result_source_preview("result", "source", db=FakeDb(b"not a docx"), precision_rule="R16_SERVER_OS_QUANTITY")
    assert caught.value.status_code == 422
