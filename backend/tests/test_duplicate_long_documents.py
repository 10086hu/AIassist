import json
from dataclasses import replace
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document

from app.modules.duplicate import document_cache, llm_extractor, llm_judge, service
from app.modules.duplicate.candidate_search import CandidateIndex
from app.modules.duplicate.compact_extractor import extract_compact_modules, focused_document
from app.modules.duplicate.document_parser import parse_document
from app.modules.duplicate.excel_parser import ParsedFunctionPoint


def serialize(document):
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def add_function_table(document, scope="道路挖掘一件事", module="申报模块"):
    table = document.add_table(rows=1, cols=5)
    for cell, value in zip(table.rows[0].cells, ["序号", "系统名称", "模块名称", "子模块名称", "功能描述"]):
        cell.text = value
    for i, (core, detail) in enumerate([("综合表单", "定制道路施工申请表"), ("流程定制", "新建道路施工审批流程")], 1):
        for cell, value in zip(table.add_row().cells, [str(i), scope, module, core, detail]):
            cell.text = value
    table.cell(1, 1).merge(table.cell(2, 1)).text = scope
    table.cell(1, 2).merge(table.cell(2, 2)).text = module
    return table


def test_long_report_merges_subfunctions_and_design_budget_repetitions_without_model(monkeypatch):
    document = Document()
    document.add_heading("项目概况", 1)
    for i in range(200):
        document.add_paragraph("背景与建设必要性" * 80)
    document.add_heading("项目建设内容", 1)
    document.add_heading("软件开发方案", 2)
    document.add_paragraph("本期复用现有表单引擎，仅定制道路施工审批流程。")
    add_function_table(document)
    add_function_table(document, "犬类治理一件事")
    document.add_heading("投资概算", 1)
    add_function_table(document)
    parsed = parse_document(serialize(document), "长报告.docx")
    monkeypatch.setattr(llm_extractor, "_call_deepseek_api", lambda _: pytest.fail("结构化模块无需模型抽取"))
    points = llm_extractor.extract_function_points(parsed)
    assert len(points) == 2
    assert {p.source_location["business_scope"] for p in points} == {"道路挖掘一件事", "犬类治理一件事"}
    assert all("综合表单" in p.description and "流程定制" in p.description for p in points)
    assert "复用" in points[0].source_location["report_context"]
    table_sections = [s for s in parsed.sections if "综合表单" in s.content]
    assert table_sections[0].title == "软件开发方案"
    assert table_sections[-1].title == "投资概算"


def test_modules_after_120_are_kept_and_unlisted_heading_modules_are_supplemented():
    document = Document()
    document.add_heading("项目建设内容", 1)
    for i in range(125):
        add_function_table(document, f"业务系统{i}", "查询模块")
    document.add_heading("其他软件开发", 2)
    document.add_heading("异常监控模块", 3)
    document.add_paragraph("为监管人员新建异常告警和事件跟踪功能。")
    document.add_heading("告警通知", 4)
    document.add_paragraph("提供告警通知功能。")
    points = llm_extractor.extract_function_points(parse_document(serialize(document), "报告.docx"))
    assert len(points) == 126
    assert any("业务系统124" in p.name for p in points)
    assert any("告警通知" in p.name for p in points)
    assert len(service._limit_function_points(points)) == 126


def test_fallback_focus_preserves_parent_business_scope_and_omits_background():
    document = Document()
    document.add_heading("项目概况", 1)
    document.add_paragraph("背景说明" * 10000)
    document.add_heading("建设内容", 1)
    document.add_heading("道路挖掘一件事", 2)
    document.add_heading("处理能力", 3)
    document.add_paragraph("新建综合表单和道路施工审批功能。")
    focused = focused_document(parse_document(serialize(document), "报告.docx"))
    assert "背景说明" not in focused.raw_text
    chunks = llm_extractor._build_extraction_chunks(focused)
    assert "道路挖掘一件事" in chunks[0][1]


def test_identical_heading_module_name_under_other_system_is_not_discarded():
    document = Document()
    document.add_heading("建设内容", 1)
    document.add_heading("系统甲", 2)
    add_function_table(document, "系统甲", "权限管理模块")
    document.add_heading("系统乙", 2)
    document.add_heading("权限管理模块", 3)
    document.add_paragraph("为相同业务人员另建角色权限管理功能。")
    points = extract_compact_modules(parse_document(serialize(document), "报告.docx"))
    assert len(points) == 2
    assert any("系统乙" in p.name for p in points)


def test_similar_system_names_do_not_merge_distinct_construction_units():
    document = Document()
    document.add_heading("建设内容", 1)
    document.add_heading("城市监管平台甲", 2)
    add_function_table(document, "城市监管平台甲", "权限管理模块")
    document.add_heading("城市监管平台乙", 2)
    document.add_heading("权限管理模块", 3)
    document.add_paragraph("另行独立开发权限管理模块，服务同一批监管人员。")
    points = extract_compact_modules(parse_document(serialize(document), "报告.docx"))
    assert len(points) == 2


def test_existing_system_table_is_not_a_new_construction_item():
    document = Document()
    document.add_heading("现状分析", 1)
    add_function_table(document, "既有系统", "权限管理模块")
    document.add_heading("建设内容", 1)
    add_function_table(document, "待建系统", "查询模块")
    points = extract_compact_modules(parse_document(serialize(document), "报告.docx"))
    assert len(points) == 1
    assert "待建系统" in points[0].name


def test_identical_names_and_strong_matches_survive_top_k_limit():
    points = [ParsedFunctionPoint(i, "统一身份认证", "对同一批工作人员提供单点登录与角色权限") for i in range(9)]
    index = CandidateIndex(points)
    assert len(index.search(points[0], 2, 0.5, query_index=0)) == 8


def test_partial_subfunction_overlap_survives_different_large_module_descriptions():
    left = ParsedFunctionPoint(1, "档案监管平台", "档案统计分析" * 30,
                               source_location={"core_functions": ["操作日志审计", "档案统计分析"]})
    right = ParsedFunctionPoint(2, "资产运维系统", "资产设备巡检" * 30,
                                source_location={"core_functions": ["操作日志审计", "资产设备巡检"]})
    hits = CandidateIndex([right]).search(left, 5, .5)
    assert hits and hits[0][0] == 0


def test_cached_document_avoids_reextraction_but_content_and_model_changes_invalidate(tmp_path, monkeypatch):
    config = SimpleNamespace(duplicate_cache_enabled=True, duplicate_cache_dir=str(tmp_path),
                             deepseek_model="model-a", deepseek_api_url="url")
    monkeypatch.setattr(document_cache, "settings", config)
    calls = []

    def load():
        calls.append(1)
        return [ParsedFunctionPoint(1, "模块", "核心功能", source_location={"paragraph": 5})]

    def cached(content=b"a"):
        return document_cache.cached_document_points(content, "报告.docx", "项目", True, load)

    assert cached() == cached()
    assert len(calls) == 1
    cached(b"changed")
    config.deepseek_model = "model-b"
    cached()
    assert len(calls) == 3


def test_failed_extraction_is_never_cached(tmp_path, monkeypatch):
    monkeypatch.setattr(document_cache, "settings", SimpleNamespace(
        duplicate_cache_enabled=True, duplicate_cache_dir=str(tmp_path), deepseek_model="model", deepseek_api_url="url"))
    with pytest.raises(ValueError):
        document_cache.cached_document_points(b"a", "报告.docx", "项目", True,
                                               lambda: (_ for _ in ()).throw(ValueError("抽取失败")))
    assert not list(tmp_path.glob("*.json"))


def test_same_history_is_reused_across_projects_and_names_with_correct_source_locations(tmp_path, monkeypatch):
    monkeypatch.setattr(document_cache, "settings", SimpleNamespace(
        duplicate_cache_enabled=True, duplicate_cache_dir=str(tmp_path), deepseek_model="model", deepseek_api_url="url"))
    first = ParsedFunctionPoint(1, "模块", "功能", source_location={"file_name": "原名.docx", "paragraph": 4,
                               "source_spans": [{"file_name": "原名.docx", "paragraph": 5, "quote": "功能原文"}]})
    document_cache.cached_document_points(b"same", "原名.docx", "项目甲 往期1", True, lambda: [first])
    def must_not_extract():
        pytest.fail("同一历史文档换项目或文件名无需重抽取")
    result = document_cache.cached_document_points(b"same", "副本.docx", "项目乙 一期", True, must_not_extract)
    assert result[0].source_location["file_name"] == "副本.docx"
    assert result[0].source_location["source_spans"][0]["file_name"] == "副本.docx"


def items(count=2):
    return [dict(left_name="系统甲", left_description="新建统一登录", right_name="系统乙",
                 right_description="另建统一登录", similarity=.8, context="两个独立系统",
                 left_source="系统甲新建统一登录", right_source="系统乙另建统一登录", use_llm=True)
            for _ in range(count)]


def result(identity, duplicate=False):
    return dict(id=identity, label="重复" if duplicate else "无关",
                relation="construction_duplicate" if duplicate else "different_scope", reason="判断依据",
                left_unit="系统甲", right_unit="系统乙", overlap_scope="统一登录",
                left_evidence="系统甲新建统一登录", right_evidence="系统乙另建统一登录")


def test_batch_results_are_matched_by_id_and_positive_evidence_is_checked(monkeypatch):
    monkeypatch.setattr(llm_judge, "settings", SimpleNamespace(deepseek_api_key="test", deepseek_model="test"))
    responses = [result(1), result(0, True)]
    monkeypatch.setattr(llm_judge, "_call_deepseek_api", lambda _: json.dumps({"results": responses}))
    judged = llm_judge.judge_pairs_with_llm(items())
    assert [j.label for j in judged] == ["重复", "无关"]
    responses[1]["left_evidence"] = "虚构证据"
    assert llm_judge.judge_pairs_with_llm(items())[0].label == "无关"


def test_missing_batch_results_are_retried_and_cannot_become_pass(monkeypatch):
    monkeypatch.setattr(llm_judge, "settings", SimpleNamespace(deepseek_api_key="test", deepseek_model="test"))
    monkeypatch.setattr(llm_judge, "_call_deepseek_api", lambda _: '{"results":[]}')
    monkeypatch.setattr(llm_judge, "judge_pair_with_llm", lambda **kwargs: (_ for _ in ()).throw(ValueError("缺少判定")))
    with pytest.raises(ValueError, match="缺少判定"):
        llm_judge.judge_pairs_with_llm(items())


def test_evidence_location_points_to_matching_later_subfunction():
    location = {"paragraph": 5, "quote": "用户认证", "source_spans": [
        {"paragraph": 5, "quote": "用户认证"}, {"paragraph": 10, "quote": "新建统一角色管理"}]}
    assert service._evidence_location(location, "新建统一角色管理")["paragraph"] == 10


def test_service_batches_candidates_and_keeps_business_context(monkeypatch):
    points = [service.ReportPoint(str(i), f"系统{i}", "统一登录", None, i, "报告", "本期", "current",
                                 {"quote": f"系统{i}新建统一登录", "report_context": "本期独立新建"}, "统一登录")
              for i in range(8)]
    candidates = [service.ReportCandidatePair(points[0], p, .8) for p in points[1:]]
    batches = []

    def judge(batch):
        batches.append(batch)
        assert all("本期独立新建" in item["context"] for item in batch)
        return [llm_judge.LLMJudgement("无关", "pass", "业务范围不同") for _ in batch]

    monkeypatch.setattr(service, "judge_pairs_with_llm", judge)
    monkeypatch.setattr(service, "judge_pair_with_llm", lambda **item: llm_judge.LLMJudgement("无关", "pass", "业务范围不同"))
    assert service._judge_report_pairs(candidates, "p", "internal", lambda *args: "内部比较") == []
    assert len(batches) == 1
    assert len(batches[0]) == 6
