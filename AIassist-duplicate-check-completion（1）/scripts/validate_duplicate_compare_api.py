from __future__ import annotations

import json
import os
import sys
from pathlib import Path
from typing import Any

import requests
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_DIR = ROOT / "samples" / "duplicate_report_test"
CURRENT_FILE = SAMPLE_DIR / "current_feasibility_report.docx"
HISTORY_FILE = SAMPLE_DIR / "history_phase1_feasibility_report.docx"
API_URL = os.getenv(
    "DUPLICATE_COMPARE_API_URL",
    "http://127.0.0.1:8000/api/evaluate/duplicate/compare",
)


def u(text: str) -> str:
    return text.encode("ascii").decode("unicode_escape")


CURRENT_PARAGRAPHS = [
    u(r"\u57ce\u5e02\u7efc\u5408\u670d\u52a1\u5e73\u53f0\uff08\u4e8c\u671f\uff09\u53ef\u884c\u6027\u7814\u7a76\u62a5\u544a"),
    u(r"\u672c\u9879\u76ee\u62df\u5efa\u8bbe\u57ce\u5e02\u7efc\u5408\u670d\u52a1\u5e73\u53f0\uff08\u4e8c\u671f\uff09\uff0c\u670d\u52a1\u5bf9\u8c61\u5305\u62ec\u5e02\u7ea7\u4e1a\u52a1\u90e8\u95e8\u3001\u533a\u7ea7\u7ecf\u529e\u4eba\u5458\u548c\u793e\u4f1a\u516c\u4f17\u3002"),
    u(r"\u4e00\u671f\u9879\u76ee\u5df2\u5b8c\u6210\u57fa\u7840\u95e8\u6237\u548c\u90e8\u5206\u4e8b\u9879\u529e\u7406\u80fd\u529b\uff0c\u4f46\u8de8\u90e8\u95e8\u6570\u636e\u534f\u540c\u3001\u7edf\u4e00\u8eab\u4efd\u8ba4\u8bc1\u6269\u5c55\u548c\u8fd0\u884c\u6001\u52bf\u76d1\u6d4b\u80fd\u529b\u4ecd\u9700\u5b8c\u5584\u3002"),
    u(r"\u4e1a\u52a1\u90e8\u95e8\u9700\u8981\u7edf\u4e00\u8eab\u4efd\u8ba4\u8bc1\u80fd\u529b\uff0c\u5b9e\u73b0\u7edf\u4e00\u767b\u5f55\u3001\u5355\u70b9\u767b\u5f55\u3001\u6743\u9650\u6821\u9a8c\u548c\u7528\u6237\u4f1a\u8bdd\u7ba1\u7406\u3002"),
    u(r"\u4e1a\u52a1\u90e8\u95e8\u9700\u8981\u6570\u636e\u5171\u4eab\u4ea4\u6362\u80fd\u529b\uff0c\u5b9e\u73b0\u8de8\u90e8\u95e8\u6570\u636e\u4ea4\u6362\u3001\u63a5\u53e3\u53d1\u5e03\u3001\u4ea4\u6362\u76d1\u63a7\u548c\u6570\u636e\u8ba2\u9605\u7ba1\u7406\u3002"),
    u(r"\u7ba1\u7406\u4eba\u5458\u9700\u8981\u7efc\u5408\u6001\u52bf\u5206\u6790\u80fd\u529b\uff0c\u5c55\u793a\u4e1a\u52a1\u529e\u7406\u91cf\u3001\u6570\u636e\u4ea4\u6362\u91cf\u3001\u7cfb\u7edf\u8fd0\u884c\u72b6\u6001\u548c\u5f02\u5e38\u544a\u8b66\u8d8b\u52bf\u3002"),
    u(r"\u5efa\u8bbe\u7edf\u4e00\u8eab\u4efd\u8ba4\u8bc1\u6a21\u5757\uff0c\u63d0\u4f9b\u8d26\u53f7\u7edf\u4e00\u7ba1\u7406\u3001\u5355\u70b9\u767b\u5f55\u3001\u6743\u9650\u6821\u9a8c\u3001\u4f1a\u8bdd\u7ba1\u7406\u548c\u767b\u5f55\u5ba1\u8ba1\u80fd\u529b\u3002"),
    u(r"\u5efa\u8bbe\u6570\u636e\u5171\u4eab\u4ea4\u6362\u6a21\u5757\uff0c\u63d0\u4f9b\u8de8\u90e8\u95e8\u6570\u636e\u4ea4\u6362\u3001\u63a5\u53e3\u53d1\u5e03\u3001\u63a5\u53e3\u8ba2\u9605\u3001\u4ea4\u6362\u76d1\u63a7\u548c\u5f02\u5e38\u544a\u8b66\u80fd\u529b\u3002"),
    u(r"\u5efa\u8bbe\u7efc\u5408\u6001\u52bf\u5206\u6790\u6a21\u5757\uff0c\u63d0\u4f9b\u8fd0\u884c\u6307\u6807\u770b\u677f\u3001\u8d8b\u52bf\u5206\u6790\u3001\u544a\u8b66\u7edf\u8ba1\u548c\u7ba1\u7406\u9a7e\u9a76\u8231\u80fd\u529b\u3002"),
    u(r"\u7cfb\u7edf\u91c7\u7528\u524d\u540e\u7aef\u5206\u79bb\u67b6\u6784\uff0c\u7edf\u4e00\u63a5\u5165\u8ba4\u8bc1\u670d\u52a1\u3001\u6570\u636e\u4ea4\u6362\u670d\u52a1\u548c\u76d1\u63a7\u5206\u6790\u670d\u52a1\u3002"),
    u(r"\u9879\u76ee\u6295\u8d44\u4e3b\u8981\u5305\u62ec\u8f6f\u4ef6\u5f00\u53d1\u3001\u7cfb\u7edf\u96c6\u6210\u3001\u6d4b\u8bd5\u90e8\u7f72\u548c\u8fd0\u7ef4\u57f9\u8bad\u8d39\u7528\u3002"),
]

HISTORY_PARAGRAPHS = [
    u(r"\u57ce\u5e02\u7efc\u5408\u670d\u52a1\u5e73\u53f0\uff08\u4e00\u671f\uff09\u53ef\u884c\u6027\u7814\u7a76\u62a5\u544a"),
    u(r"\u4e00\u671f\u9879\u76ee\u5efa\u8bbe\u57ce\u5e02\u7efc\u5408\u670d\u52a1\u5e73\u53f0\u57fa\u7840\u80fd\u529b\uff0c\u9762\u5411\u5e02\u7ea7\u4e1a\u52a1\u90e8\u95e8\u548c\u533a\u7ea7\u7ecf\u529e\u4eba\u5458\uff0c\u652f\u6491\u7edf\u4e00\u5165\u53e3\u3001\u8eab\u4efd\u8ba4\u8bc1\u3001\u6570\u636e\u4ea4\u6362\u548c\u57fa\u7840\u62a5\u8868\u7edf\u8ba1\u3002"),
    u(r"\u4e3a\u89e3\u51b3\u591a\u7cfb\u7edf\u91cd\u590d\u767b\u5f55\u3001\u6570\u636e\u63a5\u53e3\u5206\u6563\u548c\u4e1a\u52a1\u8fd0\u884c\u72b6\u6001\u4e0d\u53ef\u89c6\u7b49\u95ee\u9898\uff0c\u4e00\u671f\u9879\u76ee\u5efa\u8bbe\u5e73\u53f0\u57fa\u7840\u652f\u6491\u80fd\u529b\u3002"),
    u(r"\u7cfb\u7edf\u9700\u8981\u7edf\u4e00\u8eab\u4efd\u8ba4\u8bc1\u80fd\u529b\uff0c\u5b9e\u73b0\u7edf\u4e00\u767b\u5f55\u3001\u5355\u70b9\u767b\u5f55\u3001\u6743\u9650\u6821\u9a8c\u548c\u7528\u6237\u4f1a\u8bdd\u7ba1\u7406\u3002"),
    u(r"\u7cfb\u7edf\u9700\u8981\u6570\u636e\u4ea4\u6362\u5e73\u53f0\u80fd\u529b\uff0c\u5b9e\u73b0\u8de8\u90e8\u95e8\u6570\u636e\u4ea4\u6362\u3001\u63a5\u53e3\u53d1\u5e03\u3001\u4ea4\u6362\u76d1\u63a7\u548c\u6570\u636e\u8ba2\u9605\u7ba1\u7406\u3002"),
    u(r"\u7cfb\u7edf\u9700\u8981\u57fa\u7840\u62a5\u8868\u7edf\u8ba1\u80fd\u529b\uff0c\u5b9e\u73b0\u4e1a\u52a1\u529e\u7406\u6570\u91cf\u7edf\u8ba1\u3001\u6570\u636e\u4ea4\u6362\u6570\u91cf\u7edf\u8ba1\u548c\u5e38\u7528\u62a5\u8868\u5bfc\u51fa\u3002"),
    u(r"\u5efa\u8bbe\u7edf\u4e00\u8eab\u4efd\u8ba4\u8bc1\u6a21\u5757\uff0c\u63d0\u4f9b\u8d26\u53f7\u7edf\u4e00\u7ba1\u7406\u3001\u5355\u70b9\u767b\u5f55\u3001\u6743\u9650\u6821\u9a8c\u3001\u4f1a\u8bdd\u7ba1\u7406\u548c\u767b\u5f55\u5ba1\u8ba1\u80fd\u529b\u3002"),
    u(r"\u5efa\u8bbe\u6570\u636e\u4ea4\u6362\u5e73\u53f0\u6a21\u5757\uff0c\u63d0\u4f9b\u8de8\u90e8\u95e8\u6570\u636e\u4ea4\u6362\u3001\u63a5\u53e3\u53d1\u5e03\u3001\u63a5\u53e3\u8ba2\u9605\u3001\u4ea4\u6362\u76d1\u63a7\u548c\u5f02\u5e38\u544a\u8b66\u80fd\u529b\u3002"),
    u(r"\u5efa\u8bbe\u57fa\u7840\u62a5\u8868\u7edf\u8ba1\u6a21\u5757\uff0c\u63d0\u4f9b\u4e1a\u52a1\u6570\u91cf\u7edf\u8ba1\u3001\u6570\u636e\u4ea4\u6362\u7edf\u8ba1\u548c\u62a5\u8868\u5bfc\u51fa\u80fd\u529b\u3002"),
    u(r"\u7cfb\u7edf\u91c7\u7528\u7edf\u4e00\u95e8\u6237\u548c\u516c\u5171\u652f\u6491\u670d\u52a1\u67b6\u6784\uff0c\u63a5\u5165\u8ba4\u8bc1\u3001\u4ea4\u6362\u548c\u62a5\u8868\u80fd\u529b\u3002"),
    u(r"\u9879\u76ee\u6295\u8d44\u4e3b\u8981\u5305\u62ec\u5e73\u53f0\u57fa\u7840\u80fd\u529b\u5f00\u53d1\u3001\u90e8\u7f72\u548c\u6d4b\u8bd5\u8d39\u7528\u3002"),
]


def write_docx(path: Path, paragraphs: list[str]) -> None:
    document = Document()
    for index, paragraph in enumerate(paragraphs):
        if index == 0:
            document.add_heading(paragraph, level=1)
        else:
            document.add_paragraph(paragraph)
    document.save(path)


def generate_reports() -> None:
    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)
    write_docx(CURRENT_FILE, CURRENT_PARAGRAPHS)
    write_docx(HISTORY_FILE, HISTORY_PARAGRAPHS)


def call_api() -> tuple[int, Any]:
    with CURRENT_FILE.open("rb") as current, HISTORY_FILE.open("rb") as history:
        response = requests.post(
            API_URL,
            data={
                "project_name": u(r"\u57ce\u5e02\u7efc\u5408\u670d\u52a1\u5e73\u53f0\u91cd\u590d\u5efa\u8bbe\u9a8c\u8bc1"),
                "department": u(r"\u9a8c\u8bc1\u90e8\u95e8"),
                "current_stage": u(r"\u4e8c\u671f"),
                "history_stages": json.dumps([u(r"\u4e00\u671f")], ensure_ascii=False),
            },
            files=[
                (
                    "current_file",
                    (
                        CURRENT_FILE.name,
                        current,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                ),
                (
                    "history_files",
                    (
                        HISTORY_FILE.name,
                        history,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                ),
            ],
            timeout=600,
        )

    try:
        payload: Any = response.json()
    except ValueError:
        payload = response.text
    return response.status_code, payload


def build_success_summary(result: dict[str, Any]) -> dict[str, Any]:
    summary = result.get("summary") or {}
    return {
        "status": result.get("status"),
        "imported_count": result.get("imported_count"),
        "history_imported_count": result.get("history_imported_count"),
        "total_findings": summary.get("total_findings"),
        "cross_report_duplicate_count": summary.get("cross_report_duplicate_count"),
        "llm_model": summary.get("llm_model"),
        "current_file": str(CURRENT_FILE),
        "history_file": str(HISTORY_FILE),
        "findings": [
            {
                "title": item.get("display_title"),
                "risk_level": item.get("risk_level"),
                "review_opinion": item.get("review_opinion"),
                "evidence_summary": item.get("evidence_summary"),
                "model_name": item.get("model_name"),
            }
            for item in result.get("findings", [])[:10]
        ],
    }


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    generate_reports()
    status_code, payload = call_api()
    print(f"HTTP_STATUS {status_code}")

    if status_code >= 400:
        print(json.dumps(payload, ensure_ascii=False, indent=2) if isinstance(payload, dict) else payload)
        raise SystemExit(1)

    print(json.dumps(build_success_summary(payload), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
