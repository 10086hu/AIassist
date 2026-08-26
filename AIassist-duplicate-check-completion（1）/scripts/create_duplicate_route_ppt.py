from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


TEMPLATE = Path(r"C:/Users/ASUS/Documents/xwechat_files/wxid_8j51kiy3if8o22_6c1e/msg/file/2026-08/技术路线模板.pptx")
OUTPUT = Path(r"D:/AIassist/重复建设检查技术路线.pptx")

NAVY = RGBColor(0x00, 0x5A, 0xA0)
NAVY_DARK = RGBColor(0x08, 0x36, 0x5C)
INK = RGBColor(0x1F, 0x2D, 0x3D)
MUTED = RGBColor(0x5B, 0x6B, 0x7A)
LIGHT = RGBColor(0xF4, 0xF8, 0xFC)
PALE_BLUE = RGBColor(0xEA, 0xF2, 0xF9)
PALE_GREEN = RGBColor(0xEA, 0xF8, 0xF1)
PALE_ORANGE = RGBColor(0xFF, 0xF3, 0xE8)
ORANGE = RGBColor(0xC4, 0x51, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD9, 0xE4, 0xEE)


def clear_slide(slide):
    tree = slide.shapes._spTree
    for child in list(tree):
        if child.tag.endswith("}nvGrpSpPr") or child.tag.endswith("}grpSpPr"):
            continue
        tree.remove(child)


def rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=False):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.7)
    return shape


def text(slide, value, x, y, w, h, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT,
         valign=MSO_ANCHOR.TOP, margin=0.05, font="微软雅黑"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def line(slide, x1, y1, x2, y2, color=NAVY, width=1.2):
    shape = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def pill(slide, label, x, y, w, fill=PALE_BLUE, color=NAVY):
    rect(slide, x, y, w, 0.30, fill, fill, radius=True)
    text(slide, label, x, y + 0.015, w, 0.26, 9.5, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0.01)


def bullet_list(slide, items, x, y, w, line_h=0.34, size=11, color=INK, bullet_color=NAVY):
    for i, item in enumerate(items):
        cy = y + i * line_h
        rect(slide, x, cy + 0.09, 0.08, 0.08, bullet_color, bullet_color)
        text(slide, item, x + 0.16, cy, w - 0.16, line_h, size, color)


def step_card(slide, number, title, body, x, y, w=7.85, fill=WHITE, number_fill=NAVY):
    rect(slide, x, y, w, 0.52, fill, LINE, radius=True)
    rect(slide, x + 0.14, y + 0.12, 0.28, 0.28, number_fill, number_fill, radius=True)
    text(slide, str(number), x + 0.14, y + 0.125, 0.28, 0.25, 8.5, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0.0)
    text(slide, title, x + 0.56, y + 0.07, 1.75, 0.23, 11.2, NAVY_DARK, True)
    text(slide, body, x + 2.25, y + 0.07, w - 2.40, 0.36, 10.0, MUTED, False, valign=MSO_ANCHOR.MIDDLE)


def header(slide, title_value, subtitle):
    rect(slide, -0.45, 0.27, 0.88, 0.57, NAVY, NAVY)
    text(slide, title_value, 0.53, 0.15, 8.6, 0.55, 28, NAVY_DARK, True, valign=MSO_ANCHOR.MIDDLE)
    text(slide, subtitle, 11.30, 0.27, 4.0, 0.28, 9.5, MUTED, False, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
    line(slide, 0, 1.08, 16, 1.08, NAVY, 1.4)
    rect(slide, 0.55, 1.42, 14.90, 0.58, NAVY, NAVY)
    text(slide, "规则内容与技术实现", 0.78, 1.49, 4.5, 0.35, 16, WHITE, True, valign=MSO_ANCHOR.MIDDLE)


def panel_title(slide, label, x, y, w):
    text(slide, label, x, y, w, 0.30, 12.5, NAVY_DARK, True)
    line(slide, x, y + 0.36, x + w, y + 0.36, LINE, 0.8)


def footer_band(slide, left, mid, right):
    rect(slide, 0.55, 6.62, 14.90, 1.02, LIGHT, LIGHT)
    text(slide, left[0], 0.78, 6.78, 1.3, 0.25, 10.5, NAVY_DARK, True)
    text(slide, left[1], 2.05, 6.75, 2.65, 0.45, 10.2, INK, False)
    text(slide, mid[0], 5.55, 6.78, 1.3, 0.25, 10.5, NAVY_DARK, True)
    text(slide, mid[1], 6.85, 6.75, 3.15, 0.45, 10.2, INK, False)
    text(slide, right[0], 10.55, 6.78, 1.35, 0.25, 10.5, NAVY_DARK, True)
    text(slide, right[1], 11.85, 6.75, 3.25, 0.45, 10.2, INK, False)


def build_slide_1(slide):
    header(slide, "重复建设检查技术路线", "可研报告智能审查助手 · Module 2")
    rect(slide, 0.55, 2.15, 4.95, 4.25, WHITE, LINE)
    rect(slide, 5.72, 2.15, 9.73, 4.25, WHITE, LINE)
    panel_title(slide, "规则名称", 0.78, 2.34, 4.4)
    text(slide, "2.1 重复建设判断规则", 0.78, 2.72, 4.30, 0.35, 17, INK, True)
    panel_title(slide, "规则内容", 0.78, 3.22, 4.4)
    text(slide, "抽取当前报告及关联报告的功能点、建设内容和系统能力，识别本报告内部重复，以及跨阶段重复申报或边界不清。", 0.78, 3.62, 4.25, 0.82, 12, INK, True)
    panel_title(slide, "判定重点", 0.78, 4.63, 4.4)
    bullet_list(slide, ["同一能力的重复表述或重复申报", "相近功能的范围、层级与阶段边界", "输出证据、原因、风险和复核建议"], 0.82, 5.03, 4.2, 0.31, 10.5)
    panel_title(slide, "核心技术", 0.78, 5.98, 4.4)
    pill(slide, "文档解析", 1.60, 6.08, 1.00)
    pill(slide, "混合相似度", 2.70, 6.08, 1.28, PALE_GREEN, RGBColor(0x2B, 0x7A, 0x4B))
    pill(slide, "LLM复核", 4.08, 6.08, 0.92, PALE_ORANGE, ORANGE)
    panel_title(slide, "规则（系统）实现方式", 6.02, 2.34, 8.9)
    steps = [
        ("输入与关联", "当前报告 + 可选往期报告；记录项目、阶段和报告关系"),
        ("解析与抽取", "DOCX/PDF/XLSX/CSV 解析，形成结构化功能点"),
        ("候选筛选", "向量、关键词、同义词混合相似度，阈值 0.50"),
        ("语义复核", "DeepSeek 判断：重复 / 高度相似 / 无关"),
        ("双范围比对", "当前报告内部两两比对；当前 vs 往期跨报告比对"),
        ("结果闭环", "入库 CheckResult，API 归一化，WPF 分组展示"),
    ]
    for idx, (title_value, body) in enumerate(steps, 1):
        step_card(slide, idx, title_value, body, 6.02, 2.83 + (idx - 1) * 0.56, 8.9,
                  fill=WHITE if idx % 2 else LIGHT, number_fill=ORANGE if idx == 4 else NAVY)
    footer_band(slide, ("输入", "报告 / 功能点清单"), ("处理", "解析 → 抽取 → 筛选 → 复核"), ("输出", "疑似重复项 + 审查建议"))


def build_slide_2(slide):
    header(slide, "输入材料与功能点抽取", "从非结构化报告到可比对对象")
    rect(slide, 0.55, 2.15, 4.95, 4.25, WHITE, LINE)
    rect(slide, 5.72, 2.15, 9.73, 4.25, WHITE, LINE)
    panel_title(slide, "输入范围", 0.78, 2.34, 4.4)
    bullet_list(slide, ["当前可研报告：DOCX / PDF / XLSX / CSV", "往期、一期、二期、续建、补充或历史版本", "项目阶段、报告名称和关联关系作为上下文"], 0.82, 2.80, 4.2, 0.40, 10.8)
    panel_title(slide, "功能点字段", 0.78, 4.15, 4.4)
    text(slide, "名称 · 描述 · 分类 · 来源章节 · 报告 · 阶段 · 行号", 0.78, 4.58, 4.2, 0.34, 11.5, INK, True)
    panel_title(slide, "抽取策略", 0.78, 5.18, 4.4)
    bullet_list(slide, ["章节标题可识别时优先按章节形成点", "长文档按章节分块调用模型", "名称去重，保留来源和上下文"], 0.82, 5.60, 4.2, 0.30, 10.2)
    panel_title(slide, "处理链路", 6.02, 2.34, 8.9)
    flow = [
        ("1", "文件接收", "校验扩展名\n保存当前/往期关系"),
        ("2", "文档解析", "提取正文、章节\n表格优先结构化"),
        ("3", "功能点提取", "章节规则或 DeepSeek\n输出固定 JSON"),
        ("4", "后处理", "字段校验、名称去重\n生成 point_id"),
        ("5", "进入比对", "统一 ReportPoint\n携带报告/阶段"),
    ]
    for i, (n, t, b) in enumerate(flow):
        x = 6.05 + i * 1.87
        rect(slide, x, 3.10, 1.62, 1.60, PALE_BLUE if i != 2 else PALE_ORANGE, LINE, radius=True)
        rect(slide, x + 0.58, 3.28, 0.42, 0.42, NAVY if i != 2 else ORANGE, NAVY if i != 2 else ORANGE, radius=True)
        text(slide, n, x + 0.58, 3.31, 0.42, 0.30, 10, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0.0)
        text(slide, t, x + 0.10, 3.86, 1.42, 0.27, 11.2, NAVY_DARK, True, PP_ALIGN.CENTER)
        text(slide, b, x + 0.10, 4.20, 1.42, 0.40, 9.2, MUTED, False, PP_ALIGN.CENTER)
        if i < 4:
            text(slide, "→", x + 1.65, 3.70, 0.22, 0.25, 16, NAVY, True, PP_ALIGN.CENTER)
    rect(slide, 6.05, 5.18, 8.98, 0.88, LIGHT, LIGHT, radius=True)
    text(slide, "统一对象 ReportPoint", 6.28, 5.33, 2.20, 0.25, 11, NAVY_DARK, True)
    text(slide, "point_id = hash(source + report + row + name + description)；用于跨报告稳定关联和结果留痕。", 8.38, 5.30, 6.35, 0.44, 10.5, INK)
    footer_band(slide, ("解析", "Word/PDF/表格"), ("抽取", "结构化功能点"), ("上下文", "项目 + 阶段 + 来源"))


def build_slide_3(slide):
    header(slide, "候选筛选与语义判定", "确定性算法缩小范围，模型完成语义复核")
    rect(slide, 0.55, 2.15, 4.95, 4.25, WHITE, LINE)
    rect(slide, 5.72, 2.15, 9.73, 4.25, WHITE, LINE)
    panel_title(slide, "候选筛选规则", 0.78, 2.34, 4.4)
    text(slide, "duplicate_similarity_score", 0.78, 2.76, 4.2, 0.30, 15, NAVY_DARK, True)
    text(slide, "= 0.55 × 向量余弦\n  + 0.30 × 关键词重合（上限 1）\n  + 0.15 × 同义词命中", 0.78, 3.18, 4.2, 0.85, 12, INK, True)
    pill(slide, "候选阈值 ≥ 0.50", 0.82, 4.30, 1.55)
    pill(slide, "严格重复 ≥ 0.62", 2.48, 4.30, 1.55, PALE_ORANGE, ORANGE)
    panel_title(slide, "降级策略", 0.78, 4.95, 4.4)
    bullet_list(slide, ["未配置 API Key 或调用失败时走本地规则", "严格重复还需关键词重合度 ≥ 0.28", "相似度 ≥ 0.50 或重合度 ≥ 0.22 标为高度相似"], 0.82, 5.37, 4.2, 0.29, 10.0)
    panel_title(slide, "语义复核输出", 6.02, 2.34, 8.9)
    labels = [
        ("重复", "目标对象、业务流程、核心能力基本相同", "高", PALE_ORANGE, ORANGE),
        ("高度相似", "同一业务域有重叠，但存在范围/层级差异", "中", PALE_BLUE, NAVY),
        ("无关", "业务目标或核心能力不同，过滤不入问题清单", "通过", PALE_GREEN, RGBColor(0x2B, 0x7A, 0x4B)),
    ]
    for i, (lab, desc, risk, fill, accent) in enumerate(labels):
        y = 2.92 + i * 0.85
        rect(slide, 6.02, y, 8.90, 0.66, fill, fill, radius=True)
        rect(slide, 6.22, y + 0.16, 1.20, 0.34, accent, accent, radius=True)
        text(slide, lab, 6.22, y + 0.17, 1.20, 0.27, 10.5, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0.0)
        text(slide, desc, 7.62, y + 0.12, 5.55, 0.30, 11, INK, True)
        text(slide, risk, 13.85, y + 0.16, 0.76, 0.25, 10.5, accent, True, PP_ALIGN.RIGHT)
    rect(slide, 6.02, 5.67, 8.90, 0.48, WHITE, LINE, radius=True)
    text(slide, "提示词约束：只输出 JSON；reason / suggestion 各控制在 80 字以内。", 6.24, 5.77, 8.45, 0.25, 10.5, MUTED)
    footer_band(slide, ("筛选", "高召回候选集"), ("复核", "低温度 JSON 输出"), ("可靠性", "模型失败可降级"))


def build_slide_4(slide):
    header(slide, "内部与跨报告双范围比对", "同一项目内既看本期重复，也看历史边界")
    rect(slide, 0.55, 2.15, 4.95, 4.25, WHITE, LINE)
    rect(slide, 5.72, 2.15, 9.73, 4.25, WHITE, LINE)
    panel_title(slide, "本报告内部重复", 0.78, 2.34, 4.4)
    text(slide, "当前报告功能点两两组合\n\n候选：i < j 且 score ≥ 0.50\n复核：判断同一报告内是否重复申报、功能边界重叠或仅为合理拆分。", 0.78, 2.80, 4.25, 1.25, 12, INK, True)
    panel_title(slide, "关联报告重复", 0.78, 4.45, 4.4)
    text(slide, "当前功能点 × 往期功能点\n\n上下文带入报告名称和阶段，重点区分本期重复申报、边界不清与合理阶段延续。", 0.78, 4.91, 4.25, 0.95, 12, INK, True)
    panel_title(slide, "比对矩阵", 6.02, 2.34, 8.9)
    text(slide, "当前报告", 9.02, 2.76, 1.55, 0.30, 12.5, NAVY_DARK, True, PP_ALIGN.CENTER)
    text(slide, "内部", 6.48, 3.45, 0.85, 0.30, 12, NAVY_DARK, True, PP_ALIGN.CENTER)
    text(slide, "往期报告", 9.02, 5.27, 1.55, 0.30, 12.5, NAVY_DARK, True, PP_ALIGN.CENTER)
    text(slide, "跨报告", 13.36, 3.45, 0.85, 0.30, 12, NAVY_DARK, True, PP_ALIGN.CENTER)
    rect(slide, 8.14, 3.18, 3.15, 1.02, PALE_BLUE, NAVY, radius=True)
    text(slide, "当前功能点", 8.35, 3.40, 2.73, 0.30, 15, NAVY_DARK, True, PP_ALIGN.CENTER)
    rect(slide, 12.15, 3.18, 2.25, 1.02, PALE_ORANGE, ORANGE, radius=True)
    text(slide, "疑似重复\n来源明确", 12.35, 3.32, 1.85, 0.48, 13, ORANGE, True, PP_ALIGN.CENTER)
    rect(slide, 8.14, 4.48, 3.15, 1.02, PALE_GREEN, RGBColor(0x2B, 0x7A, 0x4B), radius=True)
    text(slide, "往期功能点", 8.35, 4.70, 2.73, 0.30, 15, RGBColor(0x2B, 0x7A, 0x4B), True, PP_ALIGN.CENTER)
    line(slide, 9.72, 4.20, 9.72, 4.48, NAVY, 1.4)
    line(slide, 11.30, 3.69, 12.15, 3.69, ORANGE, 1.4)
    text(slide, "结果分组：本报告内部重复 / 与往期报告重复", 6.30, 5.88, 8.20, 0.28, 11, MUTED, True, PP_ALIGN.CENTER)
    footer_band(slide, ("内部", "N×N 两两比对"), ("跨报告", "N×M 关联比对"), ("判定", "重复 / 相似 / 无关"))


def build_slide_5(slide):
    header(slide, "结果输出与前端闭环", "从候选对到可追溯审查意见")
    rect(slide, 0.55, 2.15, 4.95, 4.25, WHITE, LINE)
    rect(slide, 5.72, 2.15, 9.73, 4.25, WHITE, LINE)
    panel_title(slide, "输出字段", 0.78, 2.34, 4.4)
    bullet_list(slide, ["重复项标题、来源报告、来源章节/阶段", "相似度、标签、风险等级、判定原因", "证据摘要、修改建议、模型名称", "comparison_type：internal / cross_report"], 0.82, 2.80, 4.2, 0.40, 10.8)
    panel_title(slide, "持久化", 0.78, 4.72, 4.4)
    text(slide, "CheckResult：severity / score / result_label / reason / suggestion / model_name / reference_data", 0.78, 5.16, 4.2, 0.66, 11, INK, True)
    panel_title(slide, "接口与展示闭环", 6.02, 2.34, 8.9)
    steps = [
        ("API 入口", "POST /api/evaluate/duplicate/internal\nPOST /api/evaluate/duplicate/compare"),
        ("结果归一化", "统一 findings、summary、internal_pairs、cross_pairs"),
        ("任务存储", "项目状态 evaluated；旧重复结果按子类型替换"),
        ("WPF 展示", "选择重复建设检查；可选上传往期文件；结果分组展示"),
        ("复核与导出", "按风险、来源和建议定位原文，形成审查留痕"),
    ]
    for i, (t, b) in enumerate(steps, 1):
        y = 2.88 + (i - 1) * 0.62
        step_card(slide, i, t, b, 6.02, y, 8.9, fill=WHITE if i % 2 else LIGHT, number_fill=ORANGE if i == 4 else NAVY)
    footer_band(slide, ("入库", "结果可追溯"), ("展示", "内部 / 跨报告分组"), ("使用", "复核、整改、导出"))


def build_slide_6(slide):
    header(slide, "实现状态与验证结论", "当前代码已形成完整验证链路")
    rect(slide, 0.55, 2.15, 4.95, 4.25, WHITE, LINE)
    rect(slide, 5.72, 2.15, 9.73, 4.25, WHITE, LINE)
    panel_title(slide, "已实现", 0.78, 2.34, 4.4)
    bullet_list(slide, ["当前报告内部功能点检查", "当前报告与一个或多个往期报告比对", "Word/PDF 解析后功能点抽取", "DeepSeek 复核 + 本地规则降级", "后端接口、数据库、WPF 入口已接通"], 0.82, 2.80, 4.2, 0.37, 10.6)
    panel_title(slide, "验证样例", 0.78, 5.02, 4.4)
    text(slide, "current_feasibility_report.docx\nhistory_phase1_feasibility_report.docx", 0.78, 5.44, 4.20, 0.55, 10.8, INK, True)
    panel_title(slide, "验证结果", 6.02, 2.34, 8.9)
    metrics = [("HTTP", "200", "接口成功"), ("当前功能点", "14", "解析完成"), ("往期功能点", "13", "解析完成"), ("跨报告重复项", "13", "已返回"), ("模型", "deepseek", "语义复核")]
    for i, (label, value, note) in enumerate(metrics):
        x = 6.05 + (i % 3) * 2.98
        y = 2.95 + (i // 3) * 1.02
        rect(slide, x, y, 2.60, 0.76, PALE_BLUE if i != 3 else PALE_ORANGE, LINE, radius=True)
        text(slide, label, x + 0.12, y + 0.10, 2.36, 0.20, 9.4, MUTED, True, PP_ALIGN.CENTER)
        text(slide, value, x + 0.12, y + 0.29, 2.36, 0.28, 17, ORANGE if i == 3 else NAVY_DARK, True, PP_ALIGN.CENTER)
        text(slide, note, x + 0.12, y + 0.58, 2.36, 0.14, 8.5, MUTED, False, PP_ALIGN.CENTER)
    rect(slide, 6.05, 5.18, 8.90, 0.82, LIGHT, LIGHT, radius=True)
    text(slide, "结论", 6.28, 5.36, 0.72, 0.25, 12, NAVY_DARK, True)
    text(slide, "Word 报告解析 → 功能点抽取 → 跨报告重复判断 → 后端返回 → 前端调用入口，已完成并验证通过。", 7.05, 5.30, 7.55, 0.42, 11, INK, True)
    footer_band(slide, ("状态", "已完成并验证"), ("证据", "HTTP 200 / 14 / 13 / 13"), ("下一步", "缓存、分章、结果分组优化"))


def main():
    prs = Presentation(str(TEMPLATE))
    slide = prs.slides[0]
    # Deliver the requested single-page technical route overview.
    builders = [build_slide_1]
    for index, builder in enumerate(builders):
        if index:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
        clear_slide(slide)
        builder(slide)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    print(f"created {OUTPUT} with {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
