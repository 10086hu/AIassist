from pathlib import Path

from docx import Document


OUTPUT = Path(__file__).resolve().parents[1] / "samples" / "price_reasonableness_failure_demo.docx"


def main() -> None:
    document = Document()
    document.add_heading("价格合理性超标检测样例", level=1)
    document.add_paragraph("本项目为上海市市级新建项目，按等保二级建设。以下金额仅用于验证价格合理性规则。")

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.cell(0, 0).text = "投资估算总表（单位：万元）"
    table.cell(0, 1).text = ""
    table.cell(0, 2).text = ""

    rows = [
        ("项目总投资", "2000", "计费总基数"),
        ("一、系统建设费", "1000", "直接建设费"),
        ("1.1应用软件开发", "600", "软件测试费计费基数"),
        ("1.2硬件购置", "300", "软硬件购置费组成"),
        ("1.3产品软件", "100", "软硬件购置费组成"),
        ("咨询费", "300", "故意超过标准上限"),
        ("工程监理费", "200", "故意超过标准上限"),
        ("软件测试费", "100", "故意超过标准上限"),
        ("系统集成费", "100", "故意超过6%上限"),
        ("安全测评费", "150", "故意超过标准上限"),
        ("等级保护测评费", "150", "与安全测评费重复申报"),
        ("密码应用测评费", "150", "故意超过标准上限"),
    ]
    for name, amount, remark in rows:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = amount
        cells[2].text = remark

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
