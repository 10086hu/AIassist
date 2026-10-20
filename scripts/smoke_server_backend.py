"""Exercise a deployed backend from a client, including database/file persistence.

Run with backend/.venv/bin/python. --keep retains synthetic data so it can be
verified after restarting the services with --verify-record <record.json>.
"""
from __future__ import annotations

import argparse
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import time
from uuid import uuid4

from docx import Document
import requests


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://10.60.150.179:8000")
    parser.add_argument("--record", type=Path)
    parser.add_argument("--verify-record", type=Path)
    parser.add_argument("--keep", action="store_true")
    args = parser.parse_args()
    client = requests.Session()
    client.trust_env = False
    token = os.environ.get("AIASSIST_ACCESS_TOKEN")
    if not token:
        parser.error("先登录，并通过 AIASSIST_ACCESS_TOKEN 环境变量提供 Bearer 令牌")
    client.headers["Authorization"] = "Bearer " + token
    base = args.base_url.rstrip("/")

    def request(method: str, endpoint: str, **kwargs):
        response = client.request(method, base + endpoint, timeout=60, **kwargs)
        if response.status_code != 200:
            raise RuntimeError(f"{method} {endpoint}: HTTP {response.status_code}: {response.text[:500]}")
        return response

    checks = {}
    project_id = None
    try:
        assert request("GET", "/api/health").json() == {"status": "ok"}
        checks["health"] = "passed"
        if args.verify_record:
            record = json.loads(args.verify_record.read_text())
            assert base == record["base_url"], "The recorded data belongs to another server"
            project_id = record["project_id"]
            assert request("GET", f"/api/projects/{project_id}").json()["id"] == project_id
            report = request("GET", f"/api/completeness/reports/{record['report_id']}?project_id={project_id}").json()
            source = request("GET", f"/api/completeness/reports/{report['id']}/sources/0?project_id={project_id}")
            assert hashlib.sha256(source.content).hexdigest() == record["source_sha256"]
            checks["persistence_after_restart"] = "passed"
            print(json.dumps({"base_url": base, "checks": checks}, ensure_ascii=False, indent=2))
            return

        assert "/api/projects" in request("GET", "/openapi.json").json()["paths"]
        checks["openapi"] = "passed"
        for module in ("basis", "function_correspondence", "data_reasonableness", "resource", "security", "price", "price_reference", "sensitive_word"):
            response = client.get(base + "/api/evaluate/rules", params={"module": module}, timeout=30)
            if module in {"price", "price_reference"} and response.status_code == 400 and "价格核价标准配置缺失" in response.json().get("detail", ""):
                checks["rules_" + module] = {"status": "requires_configuration", "reason": response.json()["detail"]}
                continue
            response.raise_for_status()
            rules = response.json()
            assert rules["rules"], module
            checks["rules_" + module] = {"source": rules.get("source"), "count": len(rules["rules"])}
            missing = [r["standard_error"] for r in rules["rules"] if r.get("standard_status") == "missing"]
            if missing:
                checks["rules_" + module].update(status="requires_configuration", reason=missing[0])
        project = request("POST", "/api/projects", json={"name": "部署连通性测试-" + uuid4().hex[:8], "department": "自动化测试"}).json()
        project_id = project["id"]
        assert any(p["id"] == project_id for p in request("GET", "/api/projects").json())
        checks["project_write_read"] = "passed"

        document = Document()
        for title, body in (
            ("立项依据与背景", "本项目根据政务服务改革要求建设，解决群众跨部门办事的材料重复提交问题。"),
            ("现状与需求分析", "现有业务系统分散，需要统一申报、材料复用和跨部门审批。"),
            ("技术方案", "本次开发统一申报、数据交换及业务审批三个模块，为群众提供在线办理服务。"),
            ("投资估算", "本项目总投资120万元，其中软件开发费100万元，实施服务费20万元。"),
        ):
            document.add_heading(title, level=1)
            document.add_paragraph(body)
        document.add_paragraph("文件存储容量测试：" + os.urandom(80000).hex())
        table = document.add_table(rows=1, cols=4)
        for cell, value in zip(table.rows[0].cells, ("来源", "资源名称", "数量", "单位")):
            cell.text = value
        for values in (("计算资源清单", "云服务器", "2", "台"), ("PaaS服务清单", "服务器操作系统", "1", "套")):
            for cell, value in zip(table.add_row().cells, values):
                cell.text = value
        buffer = BytesIO()
        document.save(buffer)
        raw = buffer.getvalue()
        assert len(raw) > 65535
        data = {"project_id": project_id, "request_id": str(uuid4())}
        filename = "部署测试报告.docx"
        report = request("POST", "/api/completeness/reports", data=data, files={"file": (filename, raw)}).json()
        repeated = request("POST", "/api/completeness/reports", data=data, files={"file": (filename, raw)}).json()
        assert repeated["id"] == report["id"]
        assert request("GET", "/api/completeness/reports", params={"project_id": project_id}).json()["total"] == 1
        source = request("GET", f"/api/completeness/reports/{report['id']}/sources/0?project_id={project_id}")
        assert source.content == raw
        assert "完整性检查报告" in request("GET", f"/api/completeness/reports/{report['id']}/export?project_id={project_id}").text
        checks["document_upload_download"] = {"status": "passed", "bytes": len(raw)}
        checks["idempotent_upload"] = "passed"
        checks["report_export"] = "passed"

        task = request("POST", "/api/evaluate/tasks", data={"project_id": project_id, "modules": "resource", "rule_source": "local", "use_llm": "false"}, files={"file": (filename, raw)}).json()
        deadline = time.monotonic() + 60
        while time.monotonic() < deadline:
            state = request("GET", "/api/evaluate/tasks/" + task["task_id"]).json()
            if state["status"] in ("completed", "failed"):
                break
            time.sleep(0.5)
        assert state["status"] == "completed", state
        assert state["result_ids"], state
        request("GET", "/api/evaluate/results/" + state["result_ids"][0])
        checks["async_review_task"] = "passed"
        checks["project_tree"] = "passed" if any(p["id"] == project_id for p in request("GET", "/api/projects/tree").json()["items"]) else "failed"
        assert checks["project_tree"] == "passed"
        llm = request("GET", "/api/evaluate/llm/status").json()
        record = {"base_url": base, "project_id": project_id, "report_id": report["id"], "source_sha256": hashlib.sha256(raw).hexdigest(), "checks": checks, "llm": llm}
        if args.record:
            args.record.parent.mkdir(parents=True, exist_ok=True)
            args.record.write_text(json.dumps(record, ensure_ascii=False, indent=2) + "\n")
        print(json.dumps(record, ensure_ascii=False, indent=2))
    finally:
        if project_id and not args.keep:
            request("DELETE", "/api/projects/" + project_id)


if __name__ == "__main__":
    main()
