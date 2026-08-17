from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "samples" / "duplicate_rich_test"
CURRENT_FILE = OUT_DIR / "current_phase2_rich_feasibility_report.docx"
HISTORY_FILE = OUT_DIR / "history_phase1_rich_feasibility_report.docx"
README_FILE = OUT_DIR / "README.md"


def setup_document() -> Document:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[style_name].font.name = "Microsoft YaHei"
    return document


def heading(document: Document, text: str, level: int) -> None:
    document.add_heading(text, level=level)


def para(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.first_line_indent = Pt(21)
    paragraph.paragraph_format.line_spacing = 1.25


def bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text, style="List Bullet")
    paragraph.paragraph_format.line_spacing = 1.2


def feature_table(document: Document, rows: list[tuple[str, str, str]]) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "功能点"
    header[1].text = "建设内容"
    header[2].text = "说明"
    for name, content, note in rows:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = content
        cells[2].text = note


def add_cover(document: Document, title: str, stage: str) -> None:
    title_para = document.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(title)
    run.bold = True
    run.font.size = Pt(20)
    document.add_paragraph()
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"阶段：{stage}\n测试用途：重复建设检查接口联调样例").font.size = Pt(12)
    document.add_page_break()


def build_current_report() -> None:
    document = setup_document()
    add_cover(document, "市域营商服务一网通办平台（二期）可行性研究报告", "当期（二期）")

    heading(document, "一、项目概况", 1)
    para(document, "本项目拟建设市域营商服务一网通办平台（二期），围绕企业开办、项目申报、政策兑现、诉求办理和运行监管等场景，完善跨部门协同服务能力。服务对象包括市级业务部门、区级经办人员、园区服务专员和企业用户。")
    para(document, "一期项目已形成统一门户、基础身份认证、事项受理、数据交换和基础统计能力。二期拟在既有基础上扩展业务协同深度，但本测试样例中故意保留若干与一期高度相似的能力，用于验证跨报告重复建设识别。")
    para(document, "为验证当前报告内部重复建设识别，本报告还在不同章节中以不同名称重复描述了若干相同能力，例如统一账号管理与身份服务、数据共享交换与数据协同服务总线。")

    heading(document, "二、现状问题与建设必要性", 1)
    bullet(document, "企业服务入口分散，不同业务系统仍存在重复登录、重复授权和账号状态不一致问题。")
    bullet(document, "跨部门数据共享依赖人工协调，接口发布、接口订阅、交换监控和异常告警能力需要进一步体系化。")
    bullet(document, "项目申报材料预审主要依赖人工核对，规则命中、材料缺项和表单一致性校验效率不高。")
    bullet(document, "运行态势监测指标分散，企业办件量、政策兑现量、诉求响应时长和接口调用量缺少统一看板。")

    heading(document, "三、建设目标", 1)
    para(document, "建设统一身份认证与用户接入服务，提供统一账号管理、单点登录、权限校验、会话管理和登录审计能力，支撑多部门业务系统统一接入。")
    para(document, "建设跨部门数据共享交换能力，提供跨部门数据交换、接口发布、接口订阅、交换监控和异常告警能力，提升事项协同办理效率。")
    para(document, "建设智能材料预审和政策精准推送能力，支撑企业申报材料自动核验、政策匹配、办理进度跟踪和闭环反馈。")
    para(document, "建设运行态势监测能力，形成业务办理量、数据交换量、系统运行状态、异常告警趋势和服务满意度综合分析。")

    heading(document, "四、建设内容与功能设计", 1)
    heading(document, "4.1 统一身份认证中心", 2)
    para(document, "建设统一身份认证中心，提供账号统一管理、单点登录、权限校验、会话管理、登录审计和组织机构同步能力。该模块负责企业用户、经办人员和部门管理员的身份认证，支持与现有政务门户账号体系对接。")
    para(document, "功能点包括统一账号管理、单点登录、权限校验、会话管理、登录审计、密码策略管理和账号生命周期管理。")

    heading(document, "4.2 用户接入与身份服务", 2)
    para(document, "建设用户接入与身份服务模块，提供统一账号管理、单点登录、权限校验、会话管理和登录审计能力，为移动端、PC 端和部门业务系统提供统一身份服务。")
    para(document, "本模块与 4.1 统一身份认证中心在账号统一管理、单点登录、权限校验、会话管理和登录审计方面存在高度重合，保留该描述用于验证当期报告内部重复建设检查。")

    heading(document, "4.3 跨部门数据共享交换平台", 2)
    para(document, "建设跨部门数据共享交换平台，提供跨部门数据交换、接口发布、接口订阅、交换监控、异常告警和数据目录维护能力。平台支持企业登记、税务、社保、信用和项目审批等数据按需共享。")
    para(document, "功能点包括跨部门数据交换、接口发布、接口订阅、交换监控、异常告警、数据目录维护和交换日志审计。")

    heading(document, "4.4 数据协同服务总线", 2)
    para(document, "建设数据协同服务总线，提供跨部门数据交换、接口发布、接口订阅、交换监控、异常告警和数据目录维护能力，实现业务系统之间的数据协同和接口治理。")
    para(document, "该模块与 4.3 跨部门数据共享交换平台在跨部门数据交换、接口发布、接口订阅、交换监控和异常告警方面存在明显重复，用于验证当期报告内部重复功能点识别。")

    heading(document, "4.5 企业诉求闭环办理", 2)
    para(document, "建设企业诉求闭环办理模块，提供诉求登记、工单分派、部门协同、办理反馈、超期预警和满意度评价能力。企业可通过统一门户提交问题，系统按照事项类型和属地自动分派。")
    para(document, "功能点包括诉求登记、工单流转、部门协同办理、超期预警、办理反馈和服务评价。")

    heading(document, "4.6 智能材料预审", 2)
    para(document, "建设智能材料预审模块，提供申报材料完整性校验、表单字段一致性校验、证照复用校验、规则命中提示和缺项补正建议能力。")
    para(document, "系统依据事项材料清单、字段规则和历史办件经验，对企业上传材料进行自动检查，减少人工初审工作量。")

    heading(document, "4.7 政策精准推送", 2)
    para(document, "建设政策精准推送模块，提供企业画像、政策标签、匹配规则、推送订阅和触达反馈能力。系统根据行业、规模、区域、信用等级和项目类型自动匹配政策。")

    heading(document, "4.8 运行态势监测与驾驶舱", 2)
    para(document, "建设运行态势监测与驾驶舱，提供业务办理量统计、数据交换量统计、系统运行状态监测、异常告警趋势分析、服务满意度分析和部门绩效排行。")
    para(document, "该能力与一期建设的运行监测看板、交换监控和基础统计能力存在延续关系，可用于验证当期与往期报告之间的重复或高度相似功能。")

    heading(document, "4.9 移动端事项办理", 2)
    para(document, "建设移动端事项办理能力，提供事项查询、掌上申报、材料上传、进度提醒、消息订阅和移动端评价能力，提升企业使用便利性。")

    heading(document, "五、主要功能点清单", 1)
    para(document, "为便于自动抽取，本章节以文字方式再次列出主要功能点：统一账号管理、单点登录、权限校验、会话管理、登录审计、跨部门数据交换、接口发布、接口订阅、交换监控、异常告警、诉求登记、工单流转、智能材料预审、政策精准推送、运行态势监测、移动端事项办理。")
    feature_table(
        document,
        [
            ("统一身份认证中心", "统一账号管理、单点登录、权限校验、会话管理、登录审计", "与用户接入与身份服务形成当期内部重复"),
            ("用户接入与身份服务", "统一账号管理、单点登录、权限校验、会话管理、登录审计", "与统一身份认证中心高度相似"),
            ("跨部门数据共享交换平台", "跨部门数据交换、接口发布、接口订阅、交换监控、异常告警", "与数据协同服务总线形成当期内部重复"),
            ("数据协同服务总线", "跨部门数据交换、接口发布、接口订阅、交换监控、异常告警", "与一期数据交换平台形成跨期重复"),
            ("运行态势监测与驾驶舱", "业务办理量、数据交换量、系统运行状态、异常告警趋势分析", "与一期运行监测看板相似"),
        ],
    )

    heading(document, "六、实施计划与投资估算", 1)
    para(document, "项目建设周期拟定为 12 个月，分为需求深化、系统设计、开发集成、试运行和验收五个阶段。投资内容包括软件开发、系统集成、测试部署、数据治理和运维培训。")
    para(document, "本测试样例不用于真实投资测算，仅用于重复建设检查模块联调验证。")

    document.save(CURRENT_FILE)


def build_history_report() -> None:
    document = setup_document()
    add_cover(document, "市域营商服务一网通办平台（一期）可行性研究报告", "往期（一期）")

    heading(document, "一、项目概况", 1)
    para(document, "一期项目面向市级业务部门和区级窗口单位，建设营商服务统一门户、基础身份认证、事项受理、数据交换、运行监测和基础报表能力。")
    para(document, "项目重点解决多系统重复登录、数据接口分散、事项办理过程不可视和部门协同效率不足等问题，为后续二期扩展提供基础平台。")

    heading(document, "二、建设目标", 1)
    para(document, "建设统一门户与身份认证基础能力，实现统一账号管理、单点登录、权限校验、会话管理和登录审计。")
    para(document, "建设数据交换平台能力，实现跨部门数据交换、接口发布、接口订阅、交换监控和异常告警。")
    para(document, "建设运行监测看板能力，展示业务办理数量、数据交换数量、系统运行状态和异常告警趋势。")

    heading(document, "三、建设内容与功能设计", 1)
    heading(document, "3.1 统一门户与身份认证", 2)
    para(document, "建设统一门户与身份认证模块，提供统一账号管理、单点登录、权限校验、会话管理和登录审计能力。系统支持部门经办人员、企业用户和管理员统一登录，并记录登录审计信息。")
    para(document, "功能点包括账号统一管理、单点登录、权限校验、用户会话管理、登录审计和组织机构同步。")

    heading(document, "3.2 数据交换平台", 2)
    para(document, "建设数据交换平台，提供跨部门数据交换、接口发布、接口订阅、交换监控、异常告警和交换日志审计能力。平台支撑企业登记、信用信息和项目审批数据共享。")
    para(document, "功能点包括跨部门数据交换、接口发布、接口订阅、交换监控、异常告警、交换日志审计和基础数据目录维护。")

    heading(document, "3.3 事项受理与进度跟踪", 2)
    para(document, "建设事项受理与进度跟踪模块，提供事项查询、在线申报、材料上传、办理进度查询和消息提醒能力。")

    heading(document, "3.4 企业诉求办理", 2)
    para(document, "建设企业诉求办理模块，提供诉求登记、工单分派、部门协同、办理反馈和服务评价能力。企业可在线提交问题，部门按流程办理并反馈结果。")

    heading(document, "3.5 运行监测看板", 2)
    para(document, "建设运行监测看板，提供业务办理量统计、数据交换量统计、系统运行状态监测、异常告警趋势分析和部门办理效率统计。")
    para(document, "该模块为管理人员提供基础运行态势展示，支持按部门、区县、事项类型和时间维度查看统计结果。")

    heading(document, "3.6 基础报表统计", 2)
    para(document, "建设基础报表统计模块，提供业务数量统计、数据交换统计、常用报表导出和部门办理情况汇总。")

    heading(document, "四、主要功能点清单", 1)
    para(document, "一期报告主要功能点包括统一账号管理、单点登录、权限校验、会话管理、登录审计、跨部门数据交换、接口发布、接口订阅、交换监控、异常告警、事项受理、进度跟踪、诉求登记、工单分派、运行监测看板和基础报表统计。")
    feature_table(
        document,
        [
            ("统一门户与身份认证", "统一账号管理、单点登录、权限校验、会话管理、登录审计", "预期与二期统一身份认证能力形成跨期重复"),
            ("数据交换平台", "跨部门数据交换、接口发布、接口订阅、交换监控、异常告警", "预期与二期数据共享交换能力形成跨期重复"),
            ("企业诉求办理", "诉求登记、工单分派、部门协同、办理反馈、服务评价", "预期与二期企业诉求闭环办理高度相似"),
            ("运行监测看板", "业务办理量、数据交换量、系统运行状态、异常告警趋势分析", "预期与二期运行态势监测高度相似"),
        ],
    )

    heading(document, "五、实施情况", 1)
    para(document, "一期项目已完成统一门户、身份认证、数据交换、事项受理、企业诉求办理、运行监测看板和基础报表统计建设，并完成试运行。")
    para(document, "本测试样例中的往期内容用于验证二期报告与历史报告之间的重复建设识别，不代表真实项目验收材料。")

    document.save(HISTORY_FILE)


def write_readme() -> None:
    README_FILE.write_text(
        """# 丰富版重复建设检查测试样例

本目录包含两份用于重复建设检查的 Word 样例：

- `current_phase2_rich_feasibility_report.docx`：当期（二期）可研报告
- `history_phase1_rich_feasibility_report.docx`：往期（一期）可研报告

预埋的当期内部重复：

- `统一身份认证中心` 与 `用户接入与身份服务`
- `跨部门数据共享交换平台` 与 `数据协同服务总线`

预埋的当期与往期重复：

- 当期 `统一身份认证中心/用户接入与身份服务` 与往期 `统一门户与身份认证`
- 当期 `跨部门数据共享交换平台/数据协同服务总线` 与往期 `数据交换平台`
- 当期 `运行态势监测与驾驶舱` 与往期 `运行监测看板`
- 当期 `企业诉求闭环办理` 与往期 `企业诉求办理`

使用方式：在前端选择当期报告，勾选“重复建设检查”，选择“有往期文件”，再选择往期报告进行测试。
""",
        encoding="utf-8",
    )


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_current_report()
    build_history_report()
    write_readme()
    print(CURRENT_FILE)
    print(HISTORY_FILE)
    print(README_FILE)


if __name__ == "__main__":
    main()
