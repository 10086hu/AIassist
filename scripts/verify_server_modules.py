"""Live nine-module smoke check. Creates one isolated project; never records credentials."""
import argparse
import io
import json
import secrets
import time
from datetime import datetime, timezone
from pathlib import Path

import requests
from docx import Document

MODULES = ['basis', 'duplicate', 'function_correspondence', 'data_reasonableness',
           'security', 'resource', 'sensitive_word', 'price', 'price_reference']


def sample():
    doc = Document()
    doc.add_heading('服务器联调测试可研报告（合成材料）', 0)
    sections = {
        '一、项目概况与建设依据': ['依据《数字化转型规划》建设政务业务系统。项目总投资100万元。'],
        '二、建设内容': ['建设用户管理、数据查询、统计报表模块。总投资120万元。'],
        '三、业务功能分析': ['用户管理支持用户新增、修改和删除；数据查询支持条件查询。'],
        '四、系统功能需求分析': ['用户管理支持用户新增、修改和删除；数据查询支持条件查询。'],
        '4.7 安全需求分析': ['按等保三级建设，数据分类分级，采用访问控制、审计和密码应用。'],
        '五、应用功能设计': ['用户管理实现账号增删改查；统计报表实现汇总。'],
        '六、数据填报与资源申请': ['用户量1000人，并发用户2000人，活跃用户比例200%。数据量100GB，存储量10GB。'],
        '七、安全建设内容': ['部署防火墙、日志审计与备份，使用Windows服务器、Oracle数据库。'],
        '八、项目成效考核目标': ['系统可用性99.9%，完成用户管理与报表功能。'],
    }
    for title, paras in sections.items():
        doc.add_heading(title, 1)
        for para in paras:
            doc.add_paragraph(para)
    for title, rows in [
        ('功能点清单', [('模块名称', '功能点名称', '功能描述', '工作量（人月）', '单价（万元）', '金额（万元）'),
                      ('用户管理', '用户维护', '用户新增、修改、删除', '6', '3.5', '21'),
                      ('用户管理', '账号维护', '用户新增、修改、删除', '6', '3.5', '21')]),
        ('资源申请清单', [('来源', '资源名称', '数量', '单位'),
                        ('计算资源清单', '云服务器', '2', '台'),
                        ('PaaS服务清单', '服务器操作系统', '1', '套')]),
        ('投资估算明细表', [('费用名称', '金额（万元）'), ('软件开发费', '60'), ('咨询设计费', '10'),
                          ('监理费', '10'), ('软件测试费', '10'), ('系统集成费', '10'), ('合计', '120')]),
        ('软硬件产品清单', [('产品名称', '品牌', '型号', '数量', '单价（元）'),
                          ('服务器', '联调样本', 'TEST-01', '2', '80000')]),
    ]:
        doc.add_heading(title, 1)
        table = doc.add_table(rows=0, cols=len(rows[0]))
        for row in rows:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--base', default='http://10.60.150.179:8000/api')
    parser.add_argument('--record', default='output/server-nine-modules-20261023.json')
    args = parser.parse_args()
    client = requests.Session()
    client.trust_env = False
    report = {'base': args.base, 'started_at': datetime.now(timezone.utc).isoformat(),
              'scope': 'Synthetic DOCX API smoke test, not business acceptance', 'modules': []}
    project_id = None

    def call(method, path, **kw):
        r = client.request(method, args.base.rstrip('/') + path, timeout=90, **kw)
        r.raise_for_status()
        return r.json() if r.content else None

    try:
        report['health'] = call('GET', '/health')
        username = 'desktop_check_' + secrets.token_hex(5)
        password = secrets.token_urlsafe(24)
        user = call('POST', '/auth/register', json={'username': username, 'password': password,
                                                 'display_name': '桌面九模块联调测试'})
        report['test_user_id'] = user.get('id')
        login = call('POST', '/auth/login', json={'username': username, 'password': password})
        client.headers['Authorization'] = 'Bearer ' + login['access_token']
        report['auth_me_ok'] = call('GET', '/auth/me')['id'] == user['id']
        report['llm_status'] = call('GET', '/evaluate/llm/status')
        report['llm_ping'] = call('POST', '/evaluate/llm/ping', json={'text': '请返回简短JSON确认连接。'})
        project = call('POST', '/projects', json={'name': '桌面九模块联调_' + secrets.token_hex(4), 'department': '合成样本测试'})
        project_id = project['id']
        payload = sample()
        Path('output').mkdir(exist_ok=True)
        Path('output/server-nine-modules-sample.docx').write_bytes(payload)
        for module in MODULES:
            entry = {'module': module}
            report['modules'].append(entry)
            try:
                task = call('POST', '/evaluate/tasks', data={'project_id': project_id,
                            'modules': module, 'rule_source': 'api', 'use_llm': 'true'},
                            files={'file': ('server-check.docx', payload, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
                entry['task_id'] = task['task_id']
                deadline = time.monotonic() + 600
                while True:
                    state = call('GET', '/evaluate/tasks/' + task['task_id'])
                    if state['status'] in {'completed', 'partial', 'failed'}:
                        break
                    if time.monotonic() > deadline:
                        raise TimeoutError('Task still running after 600 seconds')
                    time.sleep(2)
                entry['task'] = state
                entry['results'] = [call('GET', '/evaluate/results/' + rid) for rid in state.get('result_ids', [])]
                entry['execution_ok'] = state['status'] == 'completed' and bool(entry['results']) and all(
                    d.get('result', {}).get('summary', {}).get('llm_error_count', 0) == 0
                    for d in entry['results'])
                entry['sources'] = []
                for detail in entry['results']:
                    for source in detail.get('source_documents', []):
                        for key in ('download_url', 'preview_url'):
                            url = source.get(key)
                            if url:
                                from urllib.parse import urljoin, urlparse
                                url = urljoin(args.base, url)
                                if urlparse(url).netloc != urlparse(args.base).netloc:
                                    raise ValueError('Unexpected source host')
                                response = client.get(url, timeout=30)
                                entry['sources'].append({'kind': key, 'status': response.status_code, 'bytes': len(response.content)})
                                response.raise_for_status()
            except Exception as exc:
                entry['execution_ok'] = False
                entry['error'] = str(exc)
            print(module, entry.get('execution_ok'), entry.get('task', {}).get('error') or entry.get('error', ''), flush=True)
            Path(args.record).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        tree = call('GET', '/projects/tree')
        projects = tree if isinstance(tree, list) else tree.get('items', [])
        report['history_ok'] = any(p.get('id') == project_id for p in projects)
    finally:
        if project_id:
            try:
                call('DELETE', '/projects/' + project_id)
                report['test_project_deleted'] = True
            except Exception as exc:
                report['cleanup_error'] = str(exc)
        if 'Authorization' in client.headers:
            try:
                call('POST', '/auth/logout')
                report['logged_out'] = True
            except Exception as exc:
                report['logout_error'] = str(exc)
        report['finished_at'] = datetime.now(timezone.utc).isoformat()
        Path(args.record).parent.mkdir(parents=True, exist_ok=True)
        Path(args.record).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        client.close()


if __name__ == '__main__':
    main()
